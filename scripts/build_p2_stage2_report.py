#!/usr/bin/env python3
"""Build P2's editable Chinese report from completed, saved evidence only.

Thin adapter of build_p4_paper_report.Report and its native OMML helpers.
This entry point never imports mechanics/observers/controllers, starts a run,
repairs a state, or constructs a synthetic experimental time series. At the
current stage it accepts independent-dynamics qualification bundles; later
closed-loop evidence needs its own explicit schema, rather than guessed plots.
Requires python-docx==1.2.0, numpy and matplotlib.
"""
from __future__ import annotations

import argparse
import gzip
import hashlib
import json
from dataclasses import dataclass
from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib import font_manager
import numpy as np

from build_p4_paper_report import Report, brackets, fraction, matrix, sub, sup
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

ROOT = Path(__file__).resolve().parents[1]
TITLE = "第二篇P2物体位姿与内部作用力控制论文复现汇总"
PROTOCOL = ROOT / "configs/p2_closed_loop_protocol_v1.json"
PDF_SHA = "0188b6be5f9029b40c99b377bcd1e2885ec6e6b56acc49276d6fe1fc72cc0346"
PROTOCOL_SHA = "423f58ca7cfe13861586bb486dcfd6d19c5e3d82b69e40c3d403c90d7c871df3"
MATERIALS = [
    "reports/P2_ORIGINAL_PAPER_AUDIT.md",
    "reports/P2_CONTACT_GEOMETRY_DERIVATION_RENDERED.md",
    "reports/P2_CLOSED_LOOP_FROZEN_PROTOCOL.md",
    "configs/p2_closed_loop_protocol_v1.json",
    "reports/P2_CONTACT_GEOMETRY_POST_TASK.md",
    "references/private_sources/P2_original.pdf",
]


def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def read_json(path):
    return json.loads(path.read_text(encoding="utf-8"))


def finite_array(value, shape, name):
    array = np.asarray(value, dtype=float)
    if array.shape != shape or not np.isfinite(array).all():
        raise ValueError(f"Invalid saved field {name}: {array.shape}")
    return array


@dataclass
class Bundle:
    path: Path
    metadata: dict
    summary: dict
    termination: dict | None
    traces: dict
    hashes: dict


def load_bundle(path):
    """Validate persisted identities; never infer completed control from a gate."""
    path = path.resolve(strict=True)
    metadata = read_json(path / "metadata.json")
    summary = read_json(path / "summary.json")
    hashes = read_json(path / "artifact_sha256.json")
    if metadata.get("kind") != "INDEPENDENT_DYNAMICS_QUALIFICATION_NO_CONTROL":
        raise ValueError("Unsupported evidence schema; no guessed closed-loop curves")
    if summary.get("closed_loop_runs") != 0:
        raise ValueError("Qualification bundle must not claim closed-loop trials")
    for name, expected in hashes.items():
        file = (path / name).resolve(strict=True)
        if file.parent != path or sha(file) != expected:
            raise ValueError(f"Saved artifact identity mismatch: {name}")
    for name in ("metadata.json", "summary.json", "cases.jsonl", "config_snapshot.json"):
        if name not in hashes:
            raise ValueError(f"Artifact index does not bind {name}")
    if sha(path / "config_snapshot.json") != PROTOCOL_SHA:
        raise ValueError("Evidence used a different frozen protocol")
    if metadata.get("protocol_sha256") != PROTOCOL_SHA:
        raise ValueError("Metadata protocol SHA mismatch")
    cases = [json.loads(line) for line in (path / "cases.jsonl").read_text().splitlines() if line.strip()]
    if cases != summary.get("cases"):
        raise ValueError("Case ledger differs from summary")
    if len(cases) != summary.get("executed_checks") or sum(bool(c["pass"]) for c in cases) != summary.get("passed"):
        raise ValueError("Executed/passed counts differ from saved cases")
    termination_path = path / "termination.json"
    termination = read_json(termination_path) if termination_path.is_file() else None
    if termination is not None and "termination.json" not in hashes:
        raise ValueError("Termination is not bound by artifact index")
    traces = {}
    for file in sorted(path.glob("*.jsonl.gz")):
        if file.name not in hashes:
            raise ValueError(f"Unindexed dynamics trace: {file.name}")
        with gzip.open(file, "rt", encoding="utf-8") as stream:
            rows = [json.loads(line) for line in stream if line.strip()]
        previous = -1.0
        for row in rows:
            t = float(row["time_s"])
            if not np.isfinite(t) or t <= previous:
                raise ValueError(f"Nonmonotonic saved clock: {file.name}")
            previous = t
            finite_array(row["state"]["O"], (4, 4), "O")
            finite_array(row["state"]["V"], (6,), "V")
            for key in ("tau", "hK_N_Nm", "hB_N_Nm", "internal_N_Nm"):
                finite_array(row[key], (2, 6), key)
            for key in ("external_N_Nm",):
                finite_array(row[key], (6,), key)
            for key in ("force_balance_N", "moment_balance_Nm"):
                finite_array(row[key], (3,), key)
        if rows:
            traces[file.name.removesuffix(".jsonl.gz")] = rows
    if summary.get("status") == "FAIL_STOP" and termination is None:
        raise ValueError("Stopped run has no persisted termination")
    return Bundle(path, metadata, summary, termination, traces, hashes)


def indexed_audit(path):
    """Read an already-completed mathematical/artifact audit without execution."""
    path=path.resolve(strict=True)
    hashes=read_json(path/'artifact_sha256.json')
    for name,expected in hashes.items():
        file=(path/name).resolve(strict=True)
        if file.parent!=path or sha(file)!=expected:
            raise ValueError(f'Audit artifact identity mismatch: {name}')
    return path,hashes


def load_prefix_audit(path,bundles):
    path,hashes=indexed_audit(path)
    for name in ('actual_prefix_metrics.json','artifact_and_preservation_audit.json','control_case_ledger.json'):
        if name not in hashes: raise ValueError(f'Prefix audit does not bind {name}')
    origin=read_json(path/'artifact_and_preservation_audit.json')
    if origin.get('kind')!='READ_ONLY_ARTIFACT_AND_PRESERVATION_AUDIT' or not origin.get('all_checks_pass'):
        raise ValueError('Prefix audit identity/checks are not accepted')
    source_run=Path(origin['source_run'])
    if not source_run.is_absolute(): source_run=ROOT/source_run
    matching=[b for b in bundles if b.path==source_run.resolve()]
    if len(matching)!=1: raise ValueError('Prefix audit does not identify one supplied dynamics run')
    b=matching[0]
    if origin['source_metadata_sha256']!=sha(b.path/'metadata.json') or origin['source_run_index_sha256']!=sha(b.path/'artifact_sha256.json'):
        raise ValueError('Prefix audit source identity mismatch')
    stats=read_json(path/'actual_prefix_metrics.json')
    rows=b.traces['transport']; poses=np.array([r['state']['O'] for r in rows])
    internal=np.array([r['internal_N_Nm'] for r in rows])
    raw=np.array([r['hK_N_Nm'] for r in rows])+np.array([r['hB_N_Nm'] for r in rows])
    exact={'final_logged_world_displacement_m':poses[-1,:3,3]-poses[0,:3,3],
           'final_logged_contact_force_components_N':raw[-1,:,:3],
           'final_logged_contact_moment_components_Nm':raw[-1,:,3:],
           'max_internal_force_norm_each_arm_N':np.max(np.linalg.norm(internal[:,:,:3],axis=2),axis=0),
           'max_internal_moment_norm_each_arm_Nm':np.max(np.linalg.norm(internal[:,:,3:],axis=2),axis=0)}
    for key,value in exact.items():
        if not np.allclose(value,np.asarray(stats[key]),rtol=1e-12,atol=1e-16):
            raise ValueError(f'Prefix metric does not match saved samples: {key}')
    if stats['sample_count']!=len(rows) or stats['last_log_time_s']!=rows[-1]['time_s']:
        raise ValueError('Prefix metric sample identity mismatch')
    for key in ('fixed_five_second_terminal_metrics','half_step_metrics','main_experiment_metrics'):
        if stats[key] is not None: raise ValueError('Unexecuted experiment has unexpected result')
    ledger=read_json(path/'control_case_ledger.json')
    if ledger['executed_primary']!=0 or ledger['executed_control_cases']!=0 or any(c['status']!='NOT_EXECUTED' for c in ledger['cases']):
        raise ValueError('Delivery ledger contains a different control execution scope')
    return {'path':str(path.relative_to(ROOT)),'hashes':hashes,'metrics':stats,'ledger':ledger}


def load_numerical_audit(path,bundles):
    path,hashes=indexed_audit(path)
    for name in ('metadata.json','summary.json','fixtures.jsonl','config_snapshot.json'):
        if name not in hashes: raise ValueError(f'Numerical audit does not bind {name}')
    metadata=read_json(path/'metadata.json'); summary=read_json(path/'summary.json')
    if metadata.get('kind')!='LOCAL_NUMERICAL_RISK_DIAGNOSIS_NO_TRAJECTORY' or metadata.get('status')!='COMPLETE_DIAGNOSIS':
        raise ValueError('Unsupported or incomplete numerical diagnosis')
    if metadata['protocol_sha256']!=PROTOCOL_SHA or sha(path/'config_snapshot.json')!=PROTOCOL_SHA:
        raise ValueError('Numerical diagnosis used another protocol')
    for name,expected in metadata['source_sha256'].items():
        if sha(ROOT/name)!=expected: raise ValueError(f'Numerical audit source mismatch: {name}')
    for name in ('baselines/p2_pose_force/mechanics.py','baselines/p2_pose_force/contact_geometry.py'):
        if any(b.metadata['source_sha256'][name]!=metadata['source_sha256'][name] for b in bundles):
            raise ValueError('Numerical diagnosis and real trajectory use different mechanics')
    if summary['status']!='COMPLETE_DIAGNOSIS' or summary['new_ODE_trajectories']!=0 or summary['closed_loop_runs']!=0:
        raise ValueError('Numerical audit has a different execution scope')
    if summary['actual_C_failure_claim'] or summary['full_filter_divergence_claim'] or summary['dynamics_gate_reexecuted']:
        raise ValueError('Numerical audit claims unexecuted trajectory/filter evidence')
    fixtures=[json.loads(line) for line in (path/'fixtures.jsonl').read_text().splitlines() if line.strip()]
    if len(fixtures)!=4 or summary['completed_fixtures']!=len(fixtures) or metadata['completed_fixtures']!=len(fixtures):
        raise ValueError('Numerical fixture count mismatch')
    if [x['name'] for x in fixtures]!=metadata['fixtures'] or [x['name'] for x in fixtures]!=[x['name'] for x in summary['fixtures']]:
        raise ValueError('Numerical fixture identities differ')
    for fixture,entry in zip(fixtures,summary['fixtures']):
        finite_array(fixture['continuous_right_field_Jacobian_42'],(42,42),'continuous FD42')
        if fixture['minimum_continuous_real_part_per_s']!=entry['fastest_continuous_real_per_s']:
            raise ValueError('Continuous spectrum summary differs from raw diagnosis')
        if len(fixture['steps'])!=2: raise ValueError('Expected main/half isolated steps')
        for index,step in enumerate(fixture['steps']):
            finite_array(step['FD_right_step_Jacobian_42'],(42,42),'step FD42')
            if step['step_s']!=metadata['midpoint_steps_s'][index] or not step['isolated_one_step_is_not_a_trajectory']:
                raise ValueError('Isolated-step identity mismatch')
            if step['FD_step_spectrum']['spectral_radius']!=entry['step_spectral_radii'][index]:
                raise ValueError('One-step spectrum summary differs from raw diagnosis')
        if not fixture['observer_full_heartbeat_Euler_without_updates_risk']['actual_IESEKF_or_measurement_updates_not_run']:
            raise ValueError('Euler risk scope is not the unexecuted observer prediction')
    return {'path':str(path.relative_to(ROOT)),'hashes':hashes,'metadata':metadata,'summary':summary,'fixtures':fixtures}


class P2Report(Report):
    def __init__(self):
        super().__init__()
        section = self.doc.sections[0]
        section.header.paragraphs[0].text = "P2 物体位姿与内部作用力控制  |  数学模型与实际动力学证据"
        for run in section.header.paragraphs[0].runs:
            run.font.size = Pt(8)
        footer = section.footer.paragraphs[0]
        for run in footer.runs:
            if "科研 Benchmark" in run.text:
                run.text = " 页  ·  P2完整复现：部分完成"
        properties = self.doc.core_properties
        properties.title = TITLE
        properties.subject = "SE(3)软接触、42维IESEKF与原位姿/内力控制理论；独立动力学真实资格及限制"
        properties.keywords = "P2, SE(3), IESEKF, 位姿控制, 内部作用力, 接触几何, 独立动力学"

    def page(self,heading):
        # A separate empty page-break paragraph can spill to a blank PDF page.
        paragraph=self.doc.add_heading(heading,1)
        paragraph.paragraph_format.page_break_before=True

    def table(self,*args,**kwargs):
        table=super().table(*args,**kwargs)
        # The inherited report helper appends a spacing paragraph after tables.
        # Remove that known empty paragraph before a following forced heading.
        paragraph=self.doc.paragraphs[-1]
        if not paragraph.text:
            paragraph._p.getparent().remove(paragraph._p)
        return table

    def picture(self, path, caption, width=6.25):
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.keep_with_next = True
        p.add_run().add_picture(str(path), width=Inches(width))
        p = self.doc.add_paragraph(caption, "Caption")
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER


def setup_plots():
    font = Path("/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc")
    if not font.is_file():
        raise RuntimeError("Chinese font missing; refuse missing-glyph figures")
    font_manager.fontManager.addfont(str(font))
    plt.rcParams.update({"font.family": font_manager.FontProperties(fname=str(font)).get_name(), "axes.unicode_minus": False,
                         "font.size": 9, "figure.dpi": 160, "savefig.dpi": 180})


def save_figure(fig, path):
    fig.savefig(path, bbox_inches="tight")
    plt.close(fig)


def flow_asset(assets):
    fig, ax = plt.subplots(figsize=(10, 4))
    ax.set_xlim(0, 10); ax.set_ylim(0, 4); ax.axis("off")
    nodes = [(1.1, 2.6, "视觉／触觉／EE位姿\n原式(17)"),
             (3.6, 2.6, "42维右扰动IESEKF\n估计 O、V、C、G"),
             (6.1, 2.6, "位姿(25)＋内力(20/21)\n仅读取估计状态"),
             (8.6, 2.6, "H选择／约化运动学(26)\n异步发布与施加"),
             (5.9, .85, "独立物体动力学(16)\n接触力(9)–(12)"),
             (2.3, .85, "物体／EE真实状态\n仅观测生成与评分使用")]
    for x, y, text in nodes:
        ax.text(x, y, text, ha="center", va="center", fontsize=9,
                bbox={"boxstyle": "round,pad=.5", "facecolor": "#edf3f7", "edgecolor": "#547087"})
    for start, end in [((2.05,2.6),(2.6,2.6)), ((4.65,2.6),(4.95,2.6)),
                       ((7.15,2.6),(7.45,2.6)), ((8.6,2),(6.9,.9)),
                       ((4.9,.85),(3.5,.85)), ((1.25,1.2),(1.1,2))]:
        ax.annotate("", xy=end, xytext=start, arrowprops={"arrowstyle":"->", "color":"#314e64"})
    ax.text(5, 3.75, "原论文闭环结构；本图说明计算关系，实施状态由第二部分实际记录决定", ha="center", fontsize=10)
    ax.text(5, .05, "轨迹与接触模式由规划层给定；地面真值不进入正式控制器", ha="center", fontsize=9)
    path = assets / "01_p2_control_flow.png"; save_figure(fig, path); return path


def dynamics_assets(bundle, assets, index):
    figures = []
    for name, rows in bundle.traces.items():
        times = np.array([x["time_s"] for x in rows])
        poses = np.array([x["state"]["O"] for x in rows])
        twists = np.array([x["state"]["V"] for x in rows])
        hK = np.array([x["hK_N_Nm"] for x in rows])
        hB = np.array([x["hB_N_Nm"] for x in rows])
        balance = np.array([x["force_balance_N"] for x in rows])
        mbalance = np.array([x["moment_balance_Nm"] for x in rows])
        fig, axes = plt.subplots(2, 2, figsize=(10, 6), sharex=True)
        for component, label in enumerate("xyz"):
            axes[0,0].plot(times, 1000*(poses[:,component,3]-poses[0,component,3]), label=label)
            axes[0,1].plot(times, twists[:,component], label=label)
            axes[1,0].plot(times, balance[:,component], label=label)
            axes[1,1].plot(times, mbalance[:,component], label=label)
        for ax, title, unit in zip(axes.ravel(), ["COM相对初态位移", "COM body线速度", "接触＋重力的分量残差", "COM力矩平衡分量残差"],
                                   ["位移 / mm", "速度 / (m/s)", "力 / N", "力矩 / (N·m)"]):
            ax.set_title(title); ax.set_ylabel(unit); ax.grid(alpha=.25); ax.legend(fontsize=8)
        for ax in axes[1]: ax.set_xlabel("保存记录中的仿真时间 / s")
        fig.suptitle(f"{bundle.path.name} / {name}：固定EE、无估计器、无反馈控制")
        fig.tight_layout()
        file = assets / f"{index:02d}_{bundle.path.name}_{name}_response.png"
        save_figure(fig,file)
        figures.append((file, f"图{index} 实际无控制动力学响应；资格程序标识{name}，不含运输控制。{len(rows)}个保存样本，时间{times[0]:.6g}–{times[-1]:.6g}s。横轴未延伸至未执行时刻，不能称5s资格已完成。"))
        fig, axes = plt.subplots(2, 1, figsize=(10, 5.4), sharex=True)
        for r in range(2):
            axes[0].plot(times, np.linalg.norm((hK+hB)[:,r,:3],axis=1), label=f"侧{r+1}总接触力范数")
            axes[1].plot(times, np.linalg.norm((hK+hB)[:,r,3:],axis=1), label=f"侧{r+1}总接触力矩范数")
        axes[0].set_ylabel("力 / N"); axes[1].set_ylabel("力矩 / (N·m)")
        axes[1].set_xlabel("保存记录中的仿真时间 / s")
        for ax in axes: ax.grid(alpha=.25); ax.legend(fontsize=8)
        fig.suptitle(f"{name}：实际hK＋hB，分别以N和N·m表示")
        fig.tight_layout()
        file = assets / f"{index:02d}_{bundle.path.name}_{name}_contact.png"
        save_figure(fig,file)
        figures.append((file,f"图{index+1} 两侧真实总接触力及力矩；包含支撑重力的外部作用，不能称为内力控制收敛曲线。时间范围只覆盖实际保存前缀。"))
        internal=np.array([x['internal_N_Nm'] for x in rows])
        fig,axes=plt.subplots(2,1,figsize=(10,5.4),sharex=True)
        for side,style in ((0,'-'),(1,'--')):
            axes[0].plot(times,np.linalg.norm(internal[:,side,:3],axis=1),style,label=f'侧{side+1}投影内力范数')
            axes[1].plot(times,np.linalg.norm(internal[:,side,3:],axis=1),style,label=f'侧{side+1}投影内力矩范数')
        axes[0].set_ylabel('投影内力 / N'); axes[1].set_ylabel('投影内力矩 / (N·m)')
        axes[1].set_xlabel('保存记录中的仿真时间 / s')
        for ax in axes: ax.grid(alpha=.25); ax.legend(fontsize=8)
        fig.suptitle(f'{name}：对称固定EE下的投影数值残差；内力反馈未开启')
        fig.tight_layout(); file=assets/f'{index+2:02d}_{bundle.path.name}_{name}_internal.png'; save_figure(fig,file)
        figures.append((file,f'图{index+2} 实际保存hI的两侧force与moment范数，分别以N、N·m表示。约10⁻¹³N量级来自对称无控制平衡及投影数值误差；不能当作内力反馈收敛、抗标定误差或联合控制有效的证据。'))
        index += 3
    return figures, index


def theory(r, cfg, assets, review=False):
    r.doc.add_heading(TITLE, 0)
    r.doc.add_paragraph("阶段性审查 · 数值动力学资格与数学风险 · 完整P2部分完成" if review else "数值动力学阶段结果 · 完整P2复现：部分完成", "Subtitle")
    r.p("Marco De Simone、Marco Costanzo、Giuseppe De Maria、Ciro Natale：Cooperative Object Transport and Assembly: Pose/Force Control by Visual-Tactile Feedback。IEEE Robotics and Automation Letters，2026；DOI：10.1109/LRA.2026.3665445。阅读全文为2026年2月接受的八页作者版本。[1]")
    r.h("1 研究问题与方法组成")
    r.p("双臂共同搬运时，相对标定误差与执行时延会使两端对同一物体提出不一致的运动要求。只跟踪物体位姿可能积累内部应力，只消除内部作用又不能保证物体到达目标。P2以SE(3)软接触描述这种运动与力的耦合，利用视觉和触觉融合估计物体、两臂相对变换和接触状态，再将位姿与内力反馈叠加到各末端速度。")
    r.p("我将问题分为接触几何、物体动力学、流形估计与联合控制四层。接触模型决定末端运动怎样改变垫片形变；动力学使接触wrench真正推动物体；估计器从观测恢复控制所需状态；控制器依据估计量计算末端输入。下述理论与原文实验现象属于论文方法解释；本机取得的证据和未执行部分在第二部分分别列出。")
    r.picture(flow_asset(assets), "图1 原论文视觉触觉位姿／内力闭环的原创计算结构图。[1，图4、式(17)–(26)]")
    r.page("第一部分  原论文数学模型与核心方法")
    r.h("2 SE(3)坐标、右扰动与物理单位")
    r.p("我记Tᵃ_b为将b坐标映到a的变换，twist按[v,ω]、wrench按[f,m]排列。位姿的六维Log由ρ和θ组成，单位为m和rad；ρ包含旋转与平移的耦合，不等于直接相减世界位置。body速度与其对应的力必须在同一原点和坐标轴配对。")
    r.eq(["X ⊞ ξ = X Exp(ξ),   X ⊟ Y = Log(",sup("Y","−1"),"X)"],"原文1/2")
    r.eq(["T = ",brackets(matrix([["R","t"],["0","1"]])),",   τ = Log(T) = ",brackets(matrix([["V(θ)⁻¹t"],["θ"]]))],"原文3–5")
    r.eq([sub("Ad","T")," = ",brackets(matrix([["R","[t]×R"],["0","R"]]))],"原文6")
    r.p("完整Ad的平移块负责作用点变化。力变换使用其逆转置，而非直接沿用twist变换；虚功hᵀV必须保持一致。由此得到的几何一致性与有限形变力学的客观性仍是两个问题。")
    r.h("2.1 软接触的变换链")
    r.p("O=Tᵇ_o表示物体COM位姿；Eᵣ为机器人原生base内末端位姿；Gᵣ=Tᵒ_gᵣ为抓取坐标；B₁=Tᵇ_b₁；C=Tᵇ²_b₁将robot1 base坐标变为robot2 base坐标。两端必须先统一base，再形成Dᵣ=Tᵉʳ_gᵣ。")
    r.eq([sub("D","1")," = ",sup(sub("E","1"),"−1"),sup(sub("B","1"),"−1"),"O",sub("G","1"),",   ",sub("D","2")," = ",sup(sub("E","2"),"−1"),"C",sup(sub("B","1"),"−1"),"O",sub("G","2")],"原文8")
    r.eq([sub("τ","r")," = Log(",sub("D","r"),"),   ",sub("h","K,r")," = −",sub("K","r"),sub("τ","r")],"原文7/9")
    r.p("hK是机器人作用于物体的弹性wrench，负号与作用方向对应。前三刚度单位N/m，后三为N·m/rad；小垫片形变和刚性物体是模型前提，式(9)不能自动外推为摩擦锥或大变形接触。")
    r.page("3 接触选择、阻尼与双侧wrench映射")
    r.eq([sub("V","rel,r")," = ",sub("R̄","r"),sub("Ad","Gᵣ⁻¹"),sub("V","o")," − ",sub("U","r"),",   ",sub("R̄","r")," = diag(",sub("R","Dᵣ"),",",sub("R","Dᵣ"),")"],"原文10")
    r.eq([sub("h","B,r")," = −",sub("B","r"),sub("V","rel,r"),",   ",sub("h","r")," = ",sub("h","K,r")," + ",sub("h","B,r")],"原文11")
    r.table(["模式","原选择矩阵Hᵣ","物理自由方向"],[["tight","I₆","无"],["pivot","[I₅,0]","EE z轴转动，不传递Mz"],["sliding","[0₄×₂,I₄]","EE x/y滑动，不传递Fx/Fy"]],[2.0,5.0,9.4])
    r.p("H按传感器所在EE原点与轴定义，不按世界轴定义。自由方向会改变抓取坐标G，不应将pivot或sliding建成双刚性weld；模式切换也不能无记录地抹掉已有载荷。")
    r.eq([sub("A","r")," = ",sub("G","r"),sup(sub("D","r"),"−1"),",   H = diag(",sub("H","1"),",",sub("H","2"),")"],"原文12的坐标")
    r.eq(["W = [",sup(sub("Ad","A₁⁻¹"),"T"),",",sup(sub("Ad","A₂⁻¹"),"T"),"]",sup("H","T"),",   ",sub("h","E")," = WHh"],"原文12")
    r.p("W的行数为6、列数为H的行数m。两tight/pivot/sliding分别使m为12/10/8；不同模式的秩须实际检查。逆转置伴随同时旋转力和力矩，并补充接触点到COM的力臂。转置配对恒等式只是代数自洽，独立力学核验还需直接计算R f及R m＋p×R f。")
    r.h("3.1 保留原阻尼近似的含义")
    r.p("原式(10)保留R̄而非完整Ad(D)。有限相对平移与旋转共同存在时，即便物体和EE作同一刚体运动，D保持不变，原阻尼输入仍可能非零。因此本次几何修订没有自动认证完整粘弹性模型的客观耗散或被动性，也没有悄悄替换原阻尼项。")
    r.page("4 印刷式(15)与我采用的接触几何偏离")
    r.eq([sub("V","Gᵣ,printed")," = +",sub("H̄","r"),sup(sub("H̄","r"),"T"),sub("Ad","Dᵣ⁻¹"),sub("V","rel,r")],"原文15：正号保留")
    r.p("原文称自由方向没有垫片形变，但印刷正号在最简单滑动例中与该描述不一致：当O、E、G初值均为I、EE静止且物体沿自由x方向以0.01m/s移动时，4ms内O和印刷G同向各移动40μm，D产生80μm形变。这个例子是公式一致性审计，不能当作运输实验。仅改负号只解决该局部例子；有限形变还必须处理投影坐标、作用点和Log微分。")
    r.p("我采用已批准的p2.spatial_free_log_constraint.v1，明确标为[DEVIATION／已批准]。它替换式(15)中的G演化，不宣称是作者原实现或已核实勘误。原力模型(9)–(12)、内力(20)/(21)与位姿(25)保持。")
    r.eq(["D = ",sup("E","−1"),"OG,   a = ",sub("Ad","G⁻¹"),sub("V","o")," − ",sub("Ad","D⁻¹"),"U,   τ̇ = ",sup(sub("J","r"),"−1"),"(τ)(a + ",sub("V","G"),")"],"D1/D4：精确几何")
    r.eq(["HN = 0,   ",sup("N","T"),"N = I,   ",sub("N","g")," = ",sub("Ad","D⁻¹"),"N"],"D6/D7")
    r.eq([sub("A","c")," = ",sup("N","T"),sup(sub("J","r"),"−1"),"(τ)",sub("N","g"),",   ",sub("V","G")," = −",sub("N","g"),sup(sub("A","c"),"−1"),sup("N","T"),sup(sub("J","r"),"−1"),"(τ)a"],"D9：已批准模型")
    r.eq([sup("N","T"),"τ̇ = 0,   ",sub("Ad","D"),sub("V","G")," ∈ null(H),   (",sup("H","T"),"Hh)",sup(" ","T"),sub("Ad","D"),sub("V","G")," = 0"],"D11/D12")
    r.p("第一条约束保持自由Log分量的初值，并不使非零预载自动归零；第二条要求G的物理自由twist位于EE原点的null(H)，使其对传递wrench不做虚功。tight没有自由方向，直接取VG=0。标准pivot/sliding在主Log旋转域内的小矩阵可逆；实现使用直接solve及奇异性检查，拒绝伪逆兜底、重锚或每步形变投影。")
    r.h("4.1 可逆域与模式切换限制")
    r.eq(["sym(",sub("A","c"),") ⪰ β(φ)I,   β(φ) = ",fraction("φ","2"),"cot(",fraction("φ","2"),"),   φ = ‖θ‖ < π"],"D14/D15")
    r.p("原形变退出界逐分量为|θᵢ|≤0.15rad，所以φ≤√3×0.15rad，β≥0.99436866。这个解析域界不授权扩大软垫有效域。68项历史资格只检验有限输入的运动学、自由基、坐标与切换负例，不能替代独立动力学或闭环性能。")
    r.p("在固定事件时刻，新增自由分量和恢复／释放分量的Log形变须≤1e−6m/rad，改变分量切换前未经过H选择的总force须≤0.006N、moment须≤2e−6N·m。失败保存MODE_SWITCH_PRELOAD_INVALID，不能等待卸载后改时间或清零抓点。本模型没有原文0.3N低抓力与摩擦释放动力学，也没有证明混合切换全局稳定。")
    r.page("5 独立COM刚体动力学")
    r.eq([sup("O","−1"),"Ȯ = ",sup(sub("V","o"),"∧"),",   M = diag(mI₃,",sub("I","COM"),")"],"原文15/16")
    r.eq([sub("V̇","o")," = ",sup("M","−1"),"[",sub("h","E")," − ",sub("B","a"),sub("V","o")," − ",brackets(matrix([["[ω]×","0"],["[v]×","[ω]×"]])),"M",sub("V","o"),"] + ",brackets(matrix([["Rᵀg"],["0"]]))],"原文16")
    r.p("物体运动由接触合wrench、重力、空气阻尼及body坐标惯性项共同产生；不能从desired pose直接赋值，也不能只画式(24)的指数衰减。原文将g称为gravity force，但该项位于M逆外，SI量纲必须解释为重力加速度。本模型取[0,0,−9.81]m/s²，等价接触平衡项为mRᵀg。")
    r.p("A的质量0.684kg、尺寸280×90×40mm来自原图5(d)。原文没有公开COM惯量，故我用均质箱Ixx=m(ly²＋lz²)/12及其循环式补齐，标为[ADAPTATION]。Cartesian速度端点承接原式(13)，省略原机器人关节和位置伺服；它是独立最小数值模型，不能据此宣称真实机器人可执行。")
    r.h("5.1 数值积分与初始状态")
    r.p("我将O、E和G按右乘Exp更新，物体速度与这些位姿通过Lie中点耦合积分。独立plant接口只接收已施加速度及接触模式，不接收desired pose、观测、估计器或控制器对象。锁定EE的5s重力settle及半步比较是实现资格，必须从积分得到真实终态，不使用解析平衡解重置物体。")
    r.page("6 42维视觉触觉IESEKF")
    r.eq(["X = (O,",sub("V","o"),",",sub("E","1"),",",sub("E","2"),",C,",sub("G","1"),",",sub("G","2"),"),   dim(TX)=42"],"原文14")
    r.p("这是六个SE(3)与一个ℝ⁶，42维切空间顺序固定为O、Vo、E1、E2、C、G1、G2。状态的42维协方差不等于一组姿态数组的欧氏协方差。右误差按X⊞δx注入，六个SE(3)通过各自右Log求差，速度块作向量差。")
    r.eq(["Z = (",sub("E","1"),",",sub("E","2"),",",sub("h","K,1"),",",sub("h","K,2"),",",sub("O","vis,1"),"×i,",sub("O","vis,2"),"×j),   dim(Z)=24+6i+6j"],"原文17")
    r.p("EE pose、触觉弹性wrench以及各可见marker转换到COM的pose组成异步观测。触觉通道对应hK，动力学与控制用hK＋hB；不能把包含阻尼的总力偷偷改成式(17)的弹性力。直接向COM提供模拟带噪pose省略图像渲染和ArUco，属于明确的感知适配。")
    r.eq([sub("X","k+1")," = ",sub("X","k")," ⊞ ΔT(f(",sub("X","k"),",",sub("U","k"),") + ",sub("w","k"),"),   Z = g(X) ⊞ v"],"原文18")
    r.h("6.1 预测、迭代更新与坐标重置")
    r.p("预测先在流形上传播均值，再以误差状态Jacobian传播P。对新观测，以当前迭代状态重算观测残差与线性化，求误差增量；每次注入保持SE(3)，最终协方差须转换到注入后的切空间。必须有迭代更新和协方差重置，普通向量EKF或只平均视觉位姿不能算完整IESEKF。原文引用[18]的框架，不提供本篇所有解析Jacobian或可直接运行代码。")
    r.eq([sup("P","−")," = FPFᵀ + LQLᵀ,   K = PHᵀ(HPHᵀ+R)⁻¹,   X ← X ⊞ δx"],"误差状态实现结构，非原文新增定理")
    r.p("上述实现结构中的F和H分别是误差状态转移与观测Jacobian，H在此处不是接触选择矩阵。观测残差的pose块由流形box-minus构造、力块作向量差；迭代中不同切空间的先验误差与P须同步搬运，最后注入后的P也须重置。只保证旋转矩阵正交还不等于完成这些滤波步骤。")
    r.p("原式(18)的w位于ΔT括号内，其字面离散噪声映射给出ΔT²尺度。不能无说明改成ΔTQ。预测250Hz，EE两侧200/50Hz、触觉100Hz、视觉15fps；每个新样本有唯一身份、采样与到达时间，缓存帧不能被重复消费。过去时延观测只能按预声明因果处理，不能读取未来数据。")
    r.p("C在运输前估计、运输时固定；仅QC=0不能阻止Kalman测量更新改变C，需要同时冻结均值及对应更新通道。估计器使用历史已发布且经H选择的held input作为u_model，不能由GT速度或未来命令补偿阻尼。GT仅用于评分和明确登记的诊断，正式内力与位姿控制只读估计状态。")
    r.h("6.2 原表I的协方差块")
    r.table(["块","原表I对角元素","使用含义"],[
        ["Q_Vo","(1,1,1,10,10,10)×10⁻⁶","速度状态噪声"],
        ["Q_O、Q_E1、Q_E2","(1,1,1,10,10,10)×10⁻⁸","pose切空间状态噪声"],
        ["Q_C","(1,1,1,10,10,10)×10⁻⁷","抓取前标定；运输时固定"],
        ["Q_G1、Q_G2","(1,1,1,10,10,10)×10⁻⁹","抓点切空间状态噪声"],
        ["R_E1、R_E2","(2,2,2,50,50,50)×10⁻⁹","新EE pose观测"],
        ["R_hK1、R_hK2","(25,25,2.25,10⁻³,10⁻³,10⁻⁵)×10⁻³","弹性wrench观测"],
        ["每个视觉R_O","(1,1,1,10,10,10)×10⁻¹⁰","每个新COM pose观测"],
    ],[4.3,8,4.1],size=8)
    r.p("这些是原文Q/R块的对角值，不是已乘采样步长的离散协方差。pose观测R的前三／后三按m²和rad²解释，wrench观测按N²和(N·m)²分开；Q须随式(18)的噪声映射传播。自定初始误差、直接COM视觉和样本时延设计必须单独说明，不能称原作者完整传感器过程已复现。")
    r.page("7 内部作用力分解及原控制律")
    r.eq([sub("h","I,t")," = (",sub("I","m")," − ",sup("W","†"),"W)Hh,   ",sub("h","I")," = ",sup("H","T"),sub("h","I,t")],"原文20")
    r.p("内力位于W零空间，不对物体合wrench做贡献。先在m维传递子空间投影，再提升回12维EE wrench；先投影12维再选择可能改变物理含义。目标是双臂零内部wrench，并不是把两指握力也降为零。原低层grasp controller仍承担保持物体的责任。")
    r.eq([sub("U","F,r")," = −",sub("K","F,r"),sub("h","I,r")],"原文21")
    r.h("7.1 原式(22)的有限形变争议与局部解释")
    r.eq([sub("ḣ","I,r")," = ",sub("K","r"),sup(sub("J","r"),"−1"),"(",sup(sub("K","r"),"−1"),sub("h","I,r"),")",sub("Ad","Dᵣ⁻¹"),sub("U","F,r")],"原文22：印刷形式")
    r.eq([sub("ḣ","K")," = K",sup(sub("J","r"),"−1"),"(−",sup("K","−1"),sub("h","K"),")",sub("Ad","D⁻¹"),"U"],"由式9直接微分的固定O/G情形")
    r.p("由D(t)=Exp(−tU)D(0)直接微分，Jr逆的参数与印刷式(22)不同；等价左平凡化也会改变伴随位置。这个差异涉及Jacobian和伴随约定，不能仅断言某一个符号已由作者勘误。本式只核对固定O/G的弹性环节，不声称给出时变W/H、阻尼及重力下投影内力的完整导数。")
    r.eq([sub("ḣ","I")," ≃ −K",sub("K","F"),sub("h","I"),",   ",sub("T","i")," = ",fraction("1",[sub("K","i"),sub("K","F,i")])],"原文24：近零局部线性化")
    r.p("原局部论证假设完美估计、物体不动且位姿贡献为零。式(24)说明近零增益时间常数，不能替代移动物体的独立动力学、全局稳定性、时延闭环或任意接触切换证明。本次原控制(21)保持，式(22)/(24)只用于理论解释，绝不用它们直接生成实验力曲线。")
    r.page("8 位姿前馈／反馈与联合末端输入")
    r.eq(["e = Log(",sup("O","−1"),sub("O","d"),"),   ",sub("U","P,r")," = ",sub("Ad","Aᵣ⁻¹"),"(",sub("V","o,d")," + ",sub("K","P,r"),"e)"],"原文25")
    r.p("误差和前馈均在当前物体坐标系表达。若五次参考产生desired-body twist，必须先以Ad(O⁻¹Od)转到当前body；随后用Aᵣ逆将这两项一起变为EE body输入。把世界位置差、Euler角差或不同坐标的速度直接相加不满足原式。")
    r.eq([sub("U","r")," = ",sub("U","F,r")," + ",sub("U","P,r"),",   ",sub("J","t,r")," = ",sub("H","r"),sub("J","r")],"原文19/26")
    r.eq([sub("q̇","r")," = ",sup(sub("J","t,r"),"†"),sub("H","r"),sub("U","r")," + (I − ",sup(sub("J","t,r"),"†"),sub("J","t,r"),")",sub("q̇","0,r")],"原文26")
    r.p("最小模型取J=I、零空间速度0，所以Uapplied=HᵀHU。它保留接触模式带来的约化输入，仍不具有真实iiwa/SIA5F或FR3的关节限位、动力学和位置执行器。命令产生、模型held input与实际施加时间应分开记录，才能判断异步时延造成的相对运动。")
    r.h("8.1 四组消融及原文实验现象")
    r.table(["原组","反馈设置","需要区分的物理含义"],[["exp1","KF=0，KP正常","关闭内力反馈；位姿反馈／前馈保留"],["exp2","KF正常，KP=0","只关闭位姿增益，前馈仍保留"],["exp3","KF/KP正常","联合控制"],["exp4","同exp3；仅估计模型K失配","trueplant仍用名义K，不同时改变两边"]],[2,5,9.4])
    r.p("原文在7s开启内力、13s开启位姿并开始运输，76s到达最后路点且移除robot2视觉，95s关闭内力。五次插值从每段当时actual pose出发，生成新段时估计跟踪误差会归零，因此曲线归零不等于真实最终定位误差为零。完整作者轨迹和抓取点数值未公开，本协议补齐的段界与路点属于适配。")
    r.p("原图6报告exp2约10mm／4.5°，exp1多数运输阶段约1.5N不必要内部力；exp3兼顾位姿和内力，exp4瞬态变差而稳态相近。卡尺平移结果低于0.1mm分辨率，是删失界而非精确0；角误差0.42°（exp2除外）。高力矩例exp3末端减约80%，但95s关闭前仍在下降，不能称已稳态归零。这些数字是原论文结果，不是本机实测。")
    r.h("8.2 原表II参数与判据")
    par = cfg["original_parameters"]
    r.table(["参数","对角元素","出处／单位"],[[key,", ".join(str(v) for v in par[key]),unit] for key,unit in [("K_nominal","原表II；N/m、N·m/rad"),("B_nominal","原表II；N·s/m、N·m·s/rad"),("B_air","原表II；按对应分量"),("KF1","原表II；速度／wrench"),("KF2","原表II；速度／wrench"),("KP","原表II；s⁻¹"),("exp4_model_K1","原表II实验失配；K单位"),("exp4_model_K2","原表II实验失配；K单位")]],[3.4,8,5],size=8)
    r.p("主试验预声明为exp1–4各五个配对seed，代表例独立统计。只有在独立动力学、滤波、控制因果性和全动力学半步资格之后，才可比较85–95s GT目标位置／角度RMS与13–76s真内力RMS。力与力矩分别以N和N·m报告，原图6未明确stack或单侧范数，不能用不同定义硬对齐幅值。")


def case_description(case):
    """Readable component-wise results, each retaining its physical unit."""
    e = case["evidence"]
    maxabs = lambda name: float(np.max(np.abs(np.asarray(e[name],dtype=float))))
    if case["name"].startswith("wrench_physical_"):
        singular = np.asarray(e["W_singular_values"],dtype=float)
        return (f"映射max|Δf|={maxabs('force_mapping_error_N'):.3g}N；"
                f"max|Δm|={maxabs('moment_mapping_error_Nm'):.3g}N·m；"
                f"内力合力残差max={maxabs('internal_force_projection_N'):.3g}N；"
                f"合力矩残差max={maxabs('internal_moment_projection_Nm'):.3g}N·m；"
                f"W最小奇异值={singular[-1]:.6g}")
    if case["name"] == "free_fall_without_contact":
        return (f"单tick位置误差max={maxabs('position_error_m'):.3g}m；"
                f"速度误差max={maxabs('velocity_error_m_s'):.3g}m/s")
    if case["name"] == "body_newton_euler_and_kinetic_energy":
        return (f"线加速度差max={maxabs('linear_acceleration_error_m_s2'):.3g}m/s²；"
                f"角加速度差max={maxabs('angular_acceleration_error_rad_s2'):.3g}rad/s²；"
                f"动能导数={e['kinetic_energy_derivative_W']:.3g}W")
    return json.dumps(e,ensure_ascii=False,separators=(",",":"))


def numerical_risk(r,audit):
    r.page('10.1 固定数学点的显式步长与预测风险')
    meta=audit['metadata']; summary=audit['summary']; fixtures=audit['fixtures']
    r.p(f"我另在4个固定数学点检查完整42维连续field和孤立one-step map的中央差分。诊断科学来源为{meta['source_commit'][:12]}，差分步为{meta['fd_step']:.6g}，状态和原力／几何模型不变，结果状态为{summary['status']}。这是局部数学风险诊断，不是新的ODE轨迹、滤波资格或控制消融，未对失败的真实plant补积分或设为平衡。")
    r.p('A、C各有一个无形变名义点和一个解析重力平衡诊断点。名义点虽然速度为0，重力加速度尚未平衡；解析点是在独立诊断副本上构造，不是5s积分终态，不能用于替代真实settle。C物体实际动力学资格没有执行，其风险不得写成已运行C失败。')
    r.eq(['μ_mid(h,λ) = 1 + hλ + ½',sup('(hλ)','2'),',   μ_Euler(Δ,λ) = 1 + Δλ'], '局部线性预测乘子；非控制试验')
    r.table(['固定诊断点','最快连续实λ / s⁻¹','中点快支倍率 h=1/3000s','半步快支倍率 h=1/6000s','4ms Euler无更新谱半径'],[
        [x['name'],f"{x['minimum_continuous_real_part_per_s']:.9g}",
         f"{x['steps'][0]['fastest_real_branch_polynomial_multiplier']['modulus']:.9g}",
         f"{x['steps'][1]['fastest_real_branch_polynomial_multiplier']['modulus']:.9g}",
         f"{x['observer_full_heartbeat_Euler_without_updates_risk']['spectrum']['spectral_radius']:.9g}"] for x in fixtures
    ],[4.5,3.0,3.0,3.0,2.9],size=8)
    r.p('表中的中点倍率来自连续Jacobian的线性多项式。该多项式在静止平衡点对应局部离散线性化；名义点f₀非零，所以不能把它当作完整非线性one-step map的精确Jacobian。下面单列实际保存的one-step map差分谱作独立核对。')
    r.table(['固定诊断点','主步FD谱半径','主步分离非中性越界数','半步FD谱半径','半步分离非中性越界数'],[
        [x['name'],f"{x['steps'][0]['FD_step_spectrum']['spectral_radius']:.12g}",
         x['steps'][0]['FD_step_spectrum']['outside_unit_circle_separated_from_neutral_count'],
         f"{x['steps'][1]['FD_step_spectrum']['spectral_radius']:.12g}",
         x['steps'][1]['FD_step_spectrum']['outside_unit_circle_separated_from_neutral_count']] for x in fixtures
    ],[4.5,3.0,3.0,3.0,2.9],size=8)
    r.p('A快连续模态约−2183.07s⁻¹；C约−9335.81s⁻¹。C主步的孤立map存在约2.73014的非中性倍率，显示主步显式中点的局部数值风险；半步快支约0.65455。完整状态含EE、C和抓点等中性方向，约1＋2.9×10⁻¹¹的谱半径与分离的2.73模式须区分，不能把近1浮点值统称全状态发散。near-one标签沿用原algebra_atol，仅标注谱，不新增PASS／FAIL验收阈值。')
    r.p('4ms完整Euler无测量更新预测在A／C的快模态风险也超过单位倍率；这只说明原式(18)预测解释需要审查。IESEKF尚未实现、未执行、未取得资格，实际异步观测、误差注入和协方差更新没有运行，不能由此声称完整滤波已经发散。局部谱也不证明更小步长可让未来完整控制通过。')
    r.p('真实首门失败仍是120s壁钟内未达到5s终态；本数学诊断没有改变这个原因。性能等价优化、物理substep或积分／预测方法调整需分别论证，任何冻结数值解释的改变不能以本表直接替换协议。原质量、刚度、阻尼、增益、时刻、判据和已批准G模型均保持。')
    r.p('数学诊断原始完整矩阵、复特征值和谱标签保存在'+audit['path']+'；本报告只读该目录绑定的文件。',small=True)
    r.table(['诊断科学文件','运行时保存SHA256'],list(meta['source_sha256'].items()),[7,9.4],size=7)


def results(r, cfg, bundles, assets, prefix_audit=None, numerical_audit=None):
    r.page("第二部分  本机独立动力学与实际复现证据")
    r.h("9 复现平台、真实执行范围和证据层级")
    r.p("我使用原物块A、原表I／II和已冻结接触演化构建独立COM动力学；实际已执行内容以以下run的metadata、逐case、原始积分状态及终止记录为准。几何资格、无控制动力学资格与滤波／控制试验属于不同证据层级，不能互换。")
    r.table(["层次","已取得／尚缺证据"],[["此前接触几何","68/68数学资格已取得；只涉及有限运动学输入"],["本輪独立动力学","下表列出实际保存case和轨迹；失败保留"],["视觉触觉42维IESEKF","本构造器接收的无控制资格run未执行滤波"],["式20/21/25联合控制与主20","闭环run为0；没有四组消融曲线或性能统计"],["真实机器人／装配","Cartesian端点不能代表机器人硬件或真实摩擦接触"]],[4.4,12])
    r.table(["本机run","环境与科学来源","实际状态"],[[b.path.name,f"Python {b.metadata.get('python')}；NumPy {b.metadata.get('numpy')}；source {b.metadata.get('source_commit','未记录')[:12]}",f"{b.summary['status']}；通过{b.summary['passed']}/已记录{b.summary['executed_checks']}项；闭环{b.summary['closed_loop_runs']}；壁钟{b.summary.get('wall_s',0):.4g}s"] for b in bundles],[5,7.3,4.1],size=8)
    for b in bundles:
        r.h("9.1 " + b.path.name)
        r.p("实际命令：" + " ".join(str(x) for x in b.metadata.get("command",[])),small=True)
        r.p("冻结protocol SHA256："+b.metadata["protocol_sha256"],small=True)
        r.table(["实际case","判定","实测证据"],[[case["name"],"通过" if case["pass"] else "未通过",case_description(case)] for case in b.summary["cases"]],[5.4,1.2,9.8],size=8)
        if b.termination:
            t=b.termination
            r.p(f"该run实际停止原因为{t['reason']}；活动资格为{t.get('active_qualification')}；最后积分tick={t.get('last_tick')}、dt={t.get('dt_s')}s。记录保留最后积分状态与异常堆栈；未重置、重启或用解析平衡状态替换。")
            if t['reason'] == 'DYNAMICS_WALL_LIMIT':
                r.p(f"单case壁钟达到120s限制时，最后积分时刻为{t['last_tick']*t['dt_s']:.9g}s，尚未到规定5s终态。11项已完成代数／单步检查通过不代表动力学资格全部通过；固定5ssettle未完成，半步及其他初态资格未执行。主20与另10个已登记控制／诊断run均为0，滤波与控制资格未取得；保留的2次小修额度没有因报告重绘而消耗。")
        else:
            r.p("该资格run没有终止记录；其通过状态仅覆盖已保存检查，不能扩大为滤波或联合控制已完成。")
    r.page("10 从真实保存日志重绘的无控制动力学响应")
    if prefix_audit:
        s=prefix_audit['metrics']
        down=-1000*s['final_logged_world_displacement_m'][2]
        f=s['final_logged_contact_force_components_N']
        hi=s['max_internal_force_norm_each_arm_N']; mi=s['max_internal_moment_norm_each_arm_Nm']
        r.p(f"独立交付审计重新读取{len(bundles)}份指定run中的固定EE前缀。末个保存样本t={s['last_log_time_s']:.9g}s，物体COM相对初态下沉{down:.9g}mm；两侧EE接触Fx分别为{f[0][0]:.9g}N和{f[1][0]:.9g}N，按本夹具轴方向映射后支撑重力。它是弹性／阻尼支撑响应，不是内力反馈产生的控制力。")
        r.p(f"整个保存前缀每侧投影内力范数最大值为{hi[0]:.6g}N／{hi[1]:.6g}N，内力矩最大值为{mi[0]:.6g}N·m／{mi[1]:.6g}N·m。两侧理想对称夹具使内部应力接近数值零；不能用这个无扰动、无控制例替代非零bias、时延或模型失配下的消融。")
        r.p('该审计的5s终态、半步和主试验指标均为null；主20与其他10run均标NOT_EXECUTED。数值零不等于已达到原论文控制性能。',small=True)
    index=2; any_trace=False
    for b in bundles:
        figures,index=dynamics_assets(b,assets,index)
        if not figures:
            r.p(f"{b.path.name}没有保存有效积分时间样本；不生成占位响应图。")
        for path,caption in figures:
            any_trace=True; r.picture(path,caption)
        for name,rows in b.traces.items():
            last=rows[-1]
            v=np.array(last['state']['V']); f=np.array(last['force_balance_N']); m=np.array(last['moment_balance_Nm'])
            r.p(f"{name}实际保存到t={last['time_s']:.9g}s；末个保存样本max|v|={np.max(np.abs(v[:3])):.9g}m/s、max|ω|={np.max(np.abs(v[3:])):.9g}rad/s、max|力平衡残差|={np.max(np.abs(f)):.9g}N、max|力矩平衡残差|={np.max(np.abs(m)):.9g}N·m。该样本与终止last_integrated_state可能相隔未写入日志的若干tick；二者不能混称最终样本。")
    if any_trace:
        r.p("以上位姿和力均来自独立plant在固定EE、零控制输入下的实际积分。阻尼瞬态或力平衡残差下降不能证明式(21)内力反馈有效，也不能证明物体跟踪参考、估计收敛或对时延稳健。曲线长度只到实际落盘记录，未为补齐原95s运输而延长。")
    if numerical_audit: numerical_risk(r,numerical_audit)
    r.page("11 与原文实验的对应及未完成部分")
    r.table(["原文实验／结果","本机实际证据","可比性"],[["图6 exp1–4运输／约1.5N、10mm等","未执行主20；只有前置无控制资格","不可作数值性能对比"],["图7高力矩／约80%降低","无控制响应不能算该对照","尚未复现"],["图8五次装配／卡尺<0.1mm","无机器人、支承接触及装配试验","尚未复现；<0.1mm不等于0"],["图9低握力滑动／G2变化","几何资格不包含0.3N抓力或摩擦释放","尚未复现"],["视频IESEKF与普通EKF对比","滤波未执行；没有估计精度数据","原文正文也未给量化统计"]],[5.3,7.4,3.7],size=8)
    stopped=any(b.summary.get('status')!='PASS' for b in bundles)
    if stopped:
        r.p("我已取得接触运动学资格和部分独立动力学证据，但本轮动力学资格没有全部通过。当前只取得无控制动力学响应；42维滤波、主20运输及高力矩／时延／切换控制实验未执行。P2完整复现为部分完成，原文联合控制的定性／定量现象尚未由本机验证。")
    else:
        r.p("本轮独立动力学已取得已保存资格证据，滤波和联合控制仍未执行；P2完整复现为部分完成。不能从前置资格通过推断原图6／7／9可重现。")
    r.p("若修复需要改变原力模型、冻结参数、形变有效域、时刻或性能阈值，就会改变科学问题，须另行裁决。针对明确程序错误的小修与原失败分别保留；不将不同修正或重复run混成主20独立样本，不为凑原文曲线调参。")
    r.h("11.1 复现入口与保存记录")
    r.p("科学入口：baselines/p2_pose_force/contact_geometry.py、mechanics.py；资格入口：scripts/qualify_p2_dynamics.py；固定条件：configs/p2_closed_loop_protocol_v1.json、reports/P2_STAGE2_DYNAMICS_QUALIFICATION_SPEC.md。本报告图由scripts/build_p2_stage2_report.py只读上列已保存run重绘；重绘本身不是新的科学实验。")
    for b in bundles:
        r.p("源文件哈希（运行时metadata，而非把当前工作区当作当时源码）：",small=True)
        r.table(["科学文件","保存SHA256"],list(b.metadata.get("source_sha256",{}).items()),[7,9.4],size=7)
    r.p("原论文PDF及逐页全文仅本地阅读；公开交付保留引文、私有源SHA、原创数学推导、源码与真实日志图，不上传原PDF。")


def application(r):
    r.page("第三部分  应用于当前双FR3课题的差距与待研究问题")
    r.p("原P2使用传感平行夹爪、软垫及可控握力；我的双FR3／Cube／受限车厢采用侧吸接触。两种接触的可传递wrench、柔性、释放机制和低层控制不同，不能把原pivot/sliding的H直接套入吸盘。此处讨论应用差距，不把应用改进或Isaac迁移算成本轮算法结果。")
    r.table(["差距","后续必须取得的证据"],[["接触拓扑与有限柔性","吸盘拉脱、剪切、转动约束与柔性辨识，明确H与作用点；不能默认原软垫K/B"],["6D力及视觉触觉观测","可靠wrench来源、传感器反力符号／变换、COM／TCP外参；模拟hK不是硬件观测"],["标定／多率／执行时延","真实相对标定、时间戳与采样／到达延迟，关节位置伺服响应及因果控制"],["FR3约化运动学","真实7DoF Jacobian、关节限位／零空间与碰撞约束；J=I资格不证明可执行"],["重力保持、接触释放和装配","真空保持能力、支承接触、物体接触与摩擦；自由Log初值保持不等于卸载"],["Isaac物理及科研指标","另行授权的可见执行、实测日志、误差定义及独立验证；数值ODE不能代替物理结果"]],[4.6,11.8])
    r.p("我首先需要完成原论文核心估计与联合控制的独立复现，确认其在声明的最小模型中能产生可重复真实响应，再分别研究接触映射、感知和执行条件。当前尚无证据表明任何FR3应用改进有效，也不据本轮几何／动力学检查宣称连续无碰撞、真实力安全或完整装配成功。")
    r.h("参考文献及本机证据")
    r.p("[1] De Simone M., Costanzo M., De Maria G., Natale C. Cooperative Object Transport and Assembly: Pose/Force Control by Visual-Tactile Feedback. IEEE RA-L, 2026. DOI:10.1109/LRA.2026.3665445；本地八页accepted作者版。")
    r.p("[2] He D., Xu W., Zhang F. Symbolic Representation and Toolkit Development of Iterated Error-State Extended Kalman Filters on Manifolds. IEEE TIE, 70(12):12533–12544, 2023；为原文[18]，本报告不声称已复现其工具全部功能。")
    r.p("[3] Solà J., Deray J., Atchuthan D. A micro Lie theory for state estimation in robotics. arXiv:1812.01537；为原文[19]，Jacobian／伴随约定同时见本项目已验收推导。")
    r.p("[4] 本机P2原文审计、接触推导、冻结协议、资格记录以及本报告第二部分列出的实际run；科学来源与原始数据SHA分别保留。")


def main():
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--run-dir",type=Path,action="append",required=True,
                        help="completed saved independent-dynamics bundle; repeat to preserve failures")
    parser.add_argument("--output-docx",type=Path,default=ROOT / "reports" / (TITLE+"_动力学资格阶段.docx"))
    parser.add_argument("--assets-dir",type=Path,default=ROOT / "reports/p2_stage2_report_assets")
    parser.add_argument('--delivery-audit-dir',type=Path,help='completed read-only preservation/prefix/30-case ledger audit')
    parser.add_argument('--numerical-audit-dir',type=Path,help='completed fixed-point numerical diagnosis; no trajectory/filter execution')
    parser.add_argument("--audit-only",action="store_true",help="validate evidence without writing report/assets")
    args=parser.parse_args()
    if sha(PROTOCOL)!=PROTOCOL_SHA: raise ValueError("Frozen protocol changed")
    cfg=read_json(PROTOCOL)
    if sha(ROOT / MATERIALS[-1])!=PDF_SHA: raise ValueError("Private P2 source changed")
    bundles=[load_bundle(p) for p in args.run_dir]
    if len({b.path for b in bundles})!=len(bundles): raise ValueError("Duplicate run is not an independent sample")
    prefix_audit=load_prefix_audit(args.delivery_audit_dir,bundles) if args.delivery_audit_dir else None
    numerical_audit=load_numerical_audit(args.numerical_audit_dir,bundles) if args.numerical_audit_dir else None
    identity={"title":TITLE,"materials_sha256":{name:sha(ROOT/name) for name in MATERIALS},
              'report_builder_sha256':sha(Path(__file__)),
              'inherited_report_helper_sha256':sha(ROOT/'scripts/build_p4_paper_report.py'),
              "runs":[{"path":str(b.path.relative_to(ROOT)),"metadata_sha256":sha(b.path/'metadata.json'),
                       "summary":{k:v for k,v in b.summary.items() if k!='cases'},
                       "saved_traces":{name:{"samples":len(rows),"last_saved_time_s":rows[-1]['time_s']} for name,rows in b.traces.items()}}
                      for b in bundles],"closed_loop_evidence_used":False,"experiment_execution_by_report":False,
              'prefix_audit':prefix_audit,
              'numerical_audit':None if numerical_audit is None else {k:v for k,v in numerical_audit.items() if k!='fixtures'}}
    if args.audit_only:
        print(json.dumps(identity,ensure_ascii=False,indent=2)); return
    assets=args.assets_dir.resolve(); output=args.output_docx.resolve()
    if assets.exists(): raise FileExistsError("Preserve prior report assets; select a new assets directory")
    if output.exists(): raise FileExistsError("Preserve prior Word report; select a new output path")
    assets.mkdir(parents=True); output.parent.mkdir(parents=True,exist_ok=True)
    setup_plots(); r=P2Report()
    theory(r,cfg,assets,review=numerical_audit is not None)
    results(r,cfg,bundles,assets,prefix_audit,numerical_audit); application(r)
    r.doc.save(output)
    identity['output_docx_sha256']=sha(output)
    identity['assets_sha256']={p.name:sha(p) for p in sorted(assets.glob('*.png'))}
    (assets/'report_data.json').write_text(json.dumps(identity,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
    print(json.dumps({"output":str(output),"assets":str(assets),"figures":len(identity['assets_sha256']),
                      "status":"部分完成","closed_loop_runs":0},ensure_ascii=False))


if __name__ == "__main__":
    main()
