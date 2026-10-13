#!/usr/bin/env python3
"""Chinese first-person scientific report, solely from saved comparison data.

Reuses existing report layout/OMML helpers without changing the historical report.
No planner, model, IK, collision-query or simulator imports.
"""
import hashlib
import json
from pathlib import Path
import sys
import zipfile

ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT/"scripts"))
import build_p4_paper_report as word
from docx.shared import Pt

ASSETS=ROOT/"reports/p4_three_assets"
word.ASSETS=ASSETS


def fmt(v):return "不可用" if v is None else f"{v:.6g}"


def main():
    data=json.loads((ASSETS/"report_data.json").read_text())
    qa=json.loads((ROOT/"reports/P4_THREE_METHOD_EVIDENCE_QA.json").read_text())
    assert data["raw_artifacts_verified"]==63 and qa["statistical_trials"]==60
    for name,sha in data["source_sha256"].items():
        assert hashlib.sha256((ROOT/name).read_bytes()).hexdigest()==sha,name
    private=Path("/home/ubuntu2004/lmy/papers_and_reference_all/references/private_sources")
    for n in ["P4_original.pdf","DEFORM_report_example.docx"]:
        assert hashlib.sha256((private/n).read_bytes()).hexdigest()==word.EXPECTED[n]
    r=word.Report()
    r.doc.sections[0].header.paragraphs[0].text="P4 闭链规划  |  RSS / PCS / NR 同条件复现实验"
    r.doc.core_properties.title="第一篇P4闭链运动规划论文复现汇总（三方法对比）"
    r.doc.core_properties.subject="三种核心算法、20R同类重建实验、原论文对照与双臂应用差距"
    r.doc.core_properties.keywords="P4,RSS,PCS,NR,闭链,20R,同条件对比"
    r.doc.add_heading("第一篇P4闭链运动规划\n论文复现汇总",0)
    r.doc.add_paragraph("RSS、PCS、NR 三方法统一实验",style="Subtitle")
    r.p("我研究了全驱动转动关节闭链的采样规划问题，并在统一20R平面机构上实现RSS、PCS和NR。三种方法共用闭链约束、关节边界、碰撞查询、双向树和局部连接分辨率；区别在于样本生成、逆解分支处理和约束投影。")
    r.p("论文：Avishai Sintov、Andy Borum、Timothy Bretl，Motion Planning of Fully-actuated Closed Kinematic Chains with Revolute Joints: A Comparative Analysis，IEEE RA-L 3(4):2886–2893，2018；DOI 10.1109/LRA.2018.2846806。")
    r.p("我完成了60次固定条件规划和3次不计入统计分母的重复检查。RSS 0/20，PCS 20/20，NR 20/20。登记奇异构型确实连接了RSS的两个分支，但随机搜索没有生成通过几何检查的奇异目标，20次均达到5s计算上限。PCS与NR的重复数据一致，RSS的超时截断数据不完全一致。")
    r.p("这些数据来自原论文环境I同类的独立几何重建，不是缺少完整输入的原环境精确副本，也不构成六种原环境的完整复现。我的结果保留了RSS失败与重复一致性缺口，不能据此给三方法作普遍排名。")
    r.p("实验日期：2026年10月13日。初始协议6d64b26，执行源码3d51ca0；原23页NR阶段报告和全部历史失败另行保留。")

    r.page("第一部分：原论文数学模型与三种核心算法")
    r.h("1 闭链规划问题与共同搜索框架")
    r.p("原论文把系统写成n个全驱动转动关节、n−1个移动刚性杆和固定基杆。‘被动链’仅是计算闭合构型时指定的子链，并非未驱动机构。平面与空间闭链分别产生3个和6个独立约束。[1，II-A] 我将闭合条件、碰撞和关节界分开，只有同时满足三者的点才可进入树。")
    r.eq([word.sub("C","cl")," = {q ∈ C : C(q)=0},   ",word.sub("C","o")," = ",word.sub("C","free")," ∩ ",word.sub("C","cl")],"原文1、2")
    r.eq(["α:[0,1]→",word.sub("C","o"),",   α(0)=",word.sub("q","s"),",   α(1)=",word.sub("q","g")],"规划要求")
    r.p("RSS、PCS、NR是约束处理策略，RRT、Bi-RRT、SBL是搜索框架。原文比较两层组合，不能将RSS与‘双向树’看成互斥算法。我采用一个双向树壳，通过策略回调完成采样、投影和边资格检查。")
    r.fig("01_three_strategy_flow","图1 三策略共用搜索与RBS流程。原创理论示意；不是原论文图的实验数据。",width=6.2)

    r.page("2 主动—被动链与解析逆解")
    r.p("我按II-B.1将闭链拆为主动链与非冗余被动链。固定主动关节后，要求被动链匹配目标位姿。我的平面被动链有三个转动关节，能同时控制平面x、y和姿态角；其两个肘部逆解分支可用中间关节正负区分。")
    r.eq([word.sub("T","p")," = ",word.sup(word.sub("T","prefix"),"−1"),word.sub("T","target"),word.sup(word.sub("T","suffix"),"−1")],"本次3R闭合")
    r.eq([word.sub("x","w"),"=x−",word.sub("l","3"),"cosφ,   ",word.sub("y","w"),"=y−",word.sub("l","3"),"sinφ"],"解析IK-1")
    r.eq(["D=",word.fraction([word.sup(word.sub("x","w"),"2"),"+",word.sup(word.sub("y","w"),"2"),"−",word.sup(word.sub("l","1"),"2"),"−",word.sup(word.sub("l","2"),"2")],["2",word.sub("l","1"),word.sub("l","2")]),",   ",word.sub("θ","2"),"=±acosD"],"解析IK-2")
    r.eq([word.sub("θ","1"),"=atan2(",word.sub("y","w"),",",word.sub("x","w"),")−atan2(",word.sub("l","2"),"sin",word.sub("θ","2"),",",word.sub("l","1"),"+",word.sub("l","2"),"cos",word.sub("θ","2"),")"],"解析IK-3")
    r.eq([word.sub("θ","3"),"=φ−",word.sub("θ","1"),"−",word.sub("θ","2")],"解析IK-4")
    r.p("|D|>1时几何不可达，不能用迭代重启或NR兜底把APC失败改成成功。我仅在预登记的1e−12编码误差范围处理余弦边界，之后仍独立检查原闭链和关节界。主动关节保持不变，两个分支分别接受验证。")
    r.p("作者代码提供ABB IRB120的八分支解析逆解；我的3R几何由平面余弦定理独立实现，没有复制未经明确授权的IK或网格。3R的两分支和ABB的八分支数量不同，不能直接套用空间实验可见性数值。")

    r.page("3 RSS：人工奇异集合与分支桥接")
    r.p("III-A指出，固定被动链的APC局部连接只能沿同一逆解分支。其被动约束Jacobian不满秩时出现执行器奇异；因为所有关节仍可驱动，这不是必然的物理不可运动状态。RSS以概率γ显式采样这些奇异点，使相邻分支可以接入同一桥节点。")
    r.eq([word.sub("D","p"),"={q∈",word.sub("C","cl")," : det(",word.sub("J","p"),")=0}"],"原文III-A定义")
    r.eq(["det(",word.sub("J","p"),")=",word.sub("l","1"),word.sub("l","2"),"sin",word.sub("θ","2"),"=0"],"本次3R奇异条件")
    r.p("我固定被动关节[16,17,18]（零起点编号），显式设置中间角q17=0或π。为得到闭合样本，我仅在奇异生成器中解析求解另一个不重叠三关节子链；普通RSS扩展和局部连接始终使用原固定被动链。它不是把PCS换名，也没有调用NR。")
    r.p("原文选择γ=0.2。我的每次奇异请求最多尝试32次，每个候选需同时满足精确奇异条件、闭链、边界和碰撞。随机请求失败保留，不暗中改为普通目标。只有奇异状态和两段真实RBS都合格，才能说该具体分支桥存在。")
    r.eq(["μ(V(q))≤μ(",word.sub("M","k"),")=p μ(",word.sub("C","cl"),"),  p=1/ξ"],"原文4，等测度假设")
    r.p("此可见性上界要求各分支等测度；实际碰撞和关节界会削弱可见性。投影后的样本也未必在流形上均匀，因此我不把γ或分支数直接转换为规划成功率。")

    r.page("4 PCS：改变被动划分以扩大可见集合")
    r.p("III-B的关键是不同被动划分具有不同人工奇异集合。某点对划分i奇异时，对划分j可以仍然规则。PCS尝试在两端共有逆解分支的划分上连接，而不是强行穿过当前划分的奇异。")
    r.eq([word.sub("V","m"),"(q) ⊆ ∪",word.sub(word.sup("M","(i)"),"kᵢ")],"原文5")
    r.eq([word.sub("P","v"),"=1−∏(1−",word.sub("p","i"),"),   μ(",word.sub("V","m"),")≤",word.sub("P","v"),"μ(",word.sub("C","cl"),")"],"原文6、9，独立划分")
    r.p("上式是独立被动链、等测度假设下的宽松上界。存在多逆解重叠子链时，原文式7–8使用条件概率修正。平面连续3R划分的共链不产生额外多分支，而一般空间链不能照搬该独立性。")
    r.p("我选择p0=[0,1,2]和p16=[16,17,18]两个真实、不同的被动划分。扩展先随机选择一个，再在其解析闭合失败时尝试另一划分一次；RBS检查两端是否共有该划分的逆解分支。每条合格边保留partition、branch和全部中点资格。")
    r.p("在我的起终点，p16分支从−1变为+1，但p0仍有共同−1分支。因此PCS可以通过p0绕过p16的人工分支限制。‘观察两个划分’不等于解路径必然使用两个划分；最终采用了什么，仍需看实际边记录。")

    r.page("5 NR、递归二分和双向随机树")
    r.p("II-B.2及III-C使用全链数值投影。每轮从约束残差和Jacobian求最小范数关节修正，重复至收敛或迭代上限。原文以KDL的Newton/SVD实现，并非普通无约束RRT轨迹的事后修补。")
    r.eq([word.sub("q","k+1"),"=",word.sub("q","k"),"−",word.sup(word.sub("J","c"),"†"),"(",word.sub("q","k"),") C(",word.sub("q","k"),")"],"NR更新")
    r.eq([word.sub("J","c"),"=UΣ",word.sup("V","T"),",   ",word.sup(word.sub("J","c"),"†"),"=V",word.sup("Σ","†"),word.sup("U","T")],"SVD伪逆")
    r.p("我复用此前已验证的Newton–SVD数学结构，按20变量/3约束做维度适配。旧FR3代码的14变量/6残差和满行秩6护栏没有被伪装为通用接口。新实现记录数值秩，对截断奇异值取零权重；rank0或停滞明确失败，不设置阻尼、线搜索或随机重启。")
    r.eq([word.sub("q","mid"),"=(",word.sub("q","a"),"+",word.sub("q","b"),")/2"],"RBS中点")
    r.p("II-C的RBS把中点投影并检查，再对左右子段递归。三方法共享同一递归壳，只有投影和分支资格不同。我的叶分辨率为L2≤0.05rad，所有叶端点均有闭链/碰撞记录，但有限点检查不等于连续扫掠证明。")
    r.p("双树交替，一侧朝采样目标扩展一步，另一侧朝新节点最多连接25步。每条插入树的边先通过RBS；进入另一树既有节点范围后保留精确目标，不用不匹配的随机划分改变目标。没有目标偏置、捷径或路径平滑。")

    r.page("6 原论文的实验与已有结论")
    r.p("原文测试六环境：I为20维平面闭链；II–III为ABB双6R搬杆；IV–V为Comau双6R搬杆/箱；VI为三Comau多闭环。C++/OMPL、KDL和PQP运行于i7-6700HQ/8GB。空间Table I每组合500次，1000s内无解用横线标记。作者各法单独调优最大扩展距离。")
    r.table(["原环境 / Bi-RRT","RSS均值 s","PCS均值 s","NR均值 s"],[["II",7.27,.22,.13],["III",23.10,1.06,.92],["IV",175.5,27.18,38.84],["V",330.6,3.81,11.76],["VI",537.2,3.7,36.63]],size=9)
    r.p("以上仅为原论文Table I摘录，不是我的实验组。Fig.3表明，平面PCS观察更多被动划分可改善局部可见性并减少运行时间；空间环境中NR的可见性优势与PCS的解析投影成本优势相互竞争，不能脱离障碍密度和机构结构判断优劣。")
    r.table(["原文Table II / 环境II","RSS","PCS","NR"],[["可见份额 %",.85,1.31,23.74],["调优最大距离 rad",.6,1,2.6],["平均采样数",5544,532,151],["失败样本 %",83.64,20.48,16.17],["采样占耗时 %",3.03,19.05,27.80],["局部连接均值 ms",4.97,.32,1.40],["局部检查每秒",102,1639,508],["局部连接占耗时 %",50.73,53.82,71.58],["树节点均值",476,408,13],["路径节点均值",26,29,7]],size=8.4)
    r.p("可见性测试使用50000个点对，不是50000次规划。Fig.4图注写环境I、正文及Table II指环境II，我保留该引用差异；我的20R数据不与该空间表混为同一组。")

    r.page("第二部分：三种算法的实际复现实验、图表与数据比较")
    r.h("7 统一模型和测试协议")
    r.p("我采用原环境I同类20R机构，因为作者固定源码没有公开该平面实验的完整输入。它保留20关节、19移动杆、固定基杆和两个解析3R划分，但尺寸、障碍和端点是独立事前声明的重建，不是原图2a的坐标复刻。")
    r.eq(["C(q)=",word.brackets(word.matrix([["∑ lᵢcos(∑ⱼ≤ᵢqⱼ)−1"],["∑ lᵢsin(∑ⱼ≤ᵢqⱼ)"],["wrap(∑ⱼqⱼ)"]]))],"本次平面约束")
    r.table(["设置","固定值","来源/差异"],[["拓扑/尺度","20R；19杆l=sin(π/38)m；A(0,0)/B(1,0)","论文I同类，几何独立重建"],["末端/边界","第20关节零长度法兰；所有q∈[−π,π]","正确20R拓扑；未声称20划分"],["被动划分","RSS p16；PCS p0/p16","两个解析分支，真实不同关节集"],["胶囊杆半径","0.002m；仅相邻实体杆排除碰撞","独立共用几何"],["E0圆障碍","(.45,.72),r=.045m","独立固定输入"],["E1新增圆障碍","(.5,.18),r=.065；(.72,.6),r=.05m","不能随结果移动"],["随机/统计","PCG64；5个seed；每条件5次","7、23、41、101、20261013"],["预算","outer 200/800；2000nodes；5s","60次+3个资格重复"],["步长/RBS",".6rad；叶.05rad；depth20；4096中点","三方法完全相同"],["闭链/NR","位置及角度各1e−8；100更新","独立工程精度，不是正式Benchmark"],["后端/机器","NumPy1.24.4，GEOS3.13.1，i7-10700F","非原C++/KDL/PQP机器性能"],["赋时/执行","没有时间轨迹，没有机器人执行","耗时仅规划计算"]],size=8.3)

    r.page("8 登记构型、数学核验与碰撞资格")
    r.p("我用上半圆19条等弦构造起点，并把最后3R切换为另一解析分支构造终点。两环境端点均通过原登记边界、闭链和真实GEOS距离检查，最小几何净距约0.0782973m。17项数学/接口测试包括Jacobian差分、两解析分支、碰撞负例、RSS分支限制、PCS划分选择和NR数值更新。")
    r.fig("02_registered_model_geometries","图2 实际登记起点、终点和奇异桥构型的杆位置。坐标从保存q绘制；不是原论文场景截图。",width=6.1)
    r.p("我没有用近限位关节代替奇异性。奇异fixture的q17=0，被动Jacobian秩2，det=4.06472×10⁻¹⁹，最小奇异值1.19786×10⁻¹⁷；闭链位置误差8.34977×10⁻¹⁶m。该点左右两段RSS局部连接都通过，证据只说明这一个离散分支桥存在。")
    r.p("旧14维NR与FCL结果没有进入三法成功率分母。三方法本轮都按真实GEOS胶囊/圆距离查询，并排除接触，未扩大碰撞许可关系。")

    r.page("9 三方法成功率与预算")
    r.table(["环境/外层预算","方法","成功/次数","全部调用均时 s","均值节点","成功路径长 rad"],[[g["environment"].replace("E0_sparse","E0").replace("E1_clutter","E1")+"/"+str(g["budget"]),g["method"],f'{g["success"]}/{g["trials"]}',fmt(g["wall_all_trials_mean"]),fmt(g["nodes_mean"]),fmt(g["path_length_success_mean"])] for g in data["summary"]["groups"]],size=8.2)
    r.fig("03_success_rate_by_budget_env","图3 本次复现实测成功率。每柱5个固定seed；同seed在不同预算不是独立新随机样本。",width=6.15)
    r.p("60次统计组内，RSS20次全部TIME_LIMIT，PCS与NR各20次成功。失败RSS没有有效路径长，表格留不可用而非填0。5s是右截断的计算预算，不能将其解释成RSS实际求解时间。")
    r.p("增加外层预算200→800没有解除RSS的5s瓶颈。PCS/NR早于200次完成，因此两预算常重复同一路径；这些重复条件不能人为扩大置信度。")

    r.page("10 计算时间、搜索规模与投影成本")
    r.fig("04_wall_time_all_trials","图4 全部真实调用时间，含失败标记。RSS为右截断数据，不用成功样本均值掩盖失败。",width=6.2)
    r.fig("05_nodes_projection_geometry_queries","图5 实际节点、投影请求和GEOS查询工作量。不同计数定义不能互换为作者采样数。",width=6.2)
    r.p("我的NR每次仅需1个外层目标和3条合格树边，均值5个树节点；PCS均值32.6个节点。对于这些端点与统一全关节距离，NR更快、解路径更短，但这不是对原论文全部环境的性能优势证明。原文主动坐标度量、单法调优距离、Python数值成本及m的数量均不同。")
    r.p("samples统计外层目标请求，包含RSS失败奇异请求；projection_calls不含奇异生成器的IK，ik_calls包含已闭合直接接受请求，ik_solver_calls才是实际解析求解数。我的表和原Table II不能据字段名称直接等同比较。")

    r.page("11 局部连接拒绝与RSS随机覆盖不足")
    r.fig("06_local_rejection_counts","图6 三策略实际局部连接失败原因。没有把投影失败改称碰撞，时间截断独立显示。",width=6.2)
    rss=data["rss_sampling"]
    r.p(f'RSS在20次统计调用中请求奇异目标{rss["requests"]}次，构造候选{rss["construction_attempts"]}个；通过全部闭链/碰撞资格的奇异样本为{rss["qualified_singular_samples"]}，合格奇异端点树边为{rss["qualified_singular_endpoint_edges"]}。固定桥存在，但随机采样没有覆盖到可用桥，所以两个常规分支的树无法合并。')
    r.p("我只能把这一现象归于当前独立重建、全维均匀主动变量、严格自碰撞和有限预算下的可用奇异样本覆盖不足，不能推断机构无解或原论文错误。折叠点+π在当前未改变的欧氏角坐标中还存在跨−π边界问题；没有改距离度量来掩盖它。")
    r.p("三轮修正分别处理Git中文路径解析、已合法精确奇异解的数值身份、NumPy标量JSON输出。没有改障碍、端点、种子、门限或实验判据。第1轮17项中的分支合并失败和第2轮部分实验ABORTED都保存；预算已用满，我没有继续修改RSS来逼近论文结果。")
    r.table(["版本","真实状态","保留结果"],[["round0","启动文件路径错误；0规划调用","提交2ab2348及修正1记录"],["round1","16/17测试通过；无比较规划","原测试FAIL/奇异证据"],["round2","17测试通过；RSS1次完整超时，PCS保存中断","完整RSS与不完整PCS gzip/ABORTED"],["round3","17测试通过；60统计+3重复","全量raw、summary与INCOMPLETE结果"]],size=8.5)

    r.page("12 合格路径、关节曲线与逐点几何")
    r.fig("07_first_qualified_paths","图7 每法按seed升序选首条成功路径的真实杆运动。RSS无合格规划路径，不用单位桥接冒充。",width=6.15)
    r.table(["选例：E0/200/seed7","PCS","NR"],[[label,fmt(data["selected_paths"]["PCS"]["trial"]["path_metrics"][key]),fmt(data["selected_paths"]["NR"]["trial"]["path_metrics"][key])] for label,key in [("离散点数","point_count"),("解树边数","edge_count"),("累计L2路径长 rad","joint_path_length_rad"),("最大位置残差 m","maximum_position_residual_m"),("最大角残差 rad","maximum_rotation_residual_rad"),("最小关节余量 rad","minimum_joint_margin_rad"),("最小平面净距 m","minimum_clearance_m"),("最大相邻L2 rad","maximum_adjacent_l2_rad")]],size=8.5)
    r.p("选例规则不按最好/最快路径挑选。全部40条成功路径都有独立逐点重检；表中只是固定选例，完整数据仍保存。空间曲线横轴是路径点号，不是物理时间；当前没有速度、加速度或真实执行稳定性结论。")
    r.page("13 路径约束、划分来源与重复一致性")
    r.fig("08_selected_path_joint_residual_clearance","图8 保存路径的20关节、闭链残差和真实GEOS净距。有限点检验，不声称连续扫掠安全。",width=6.2)
    r.fig("09_rss_fixture_bridge_and_pcs_partitions","图9 RSS登记桥接的两段局部证据，以及PCS实际搜索边/解边划分。两类证据不能混用。",width=6.2)
    r.p(f'PCS合格搜索边p16={data["pcs_tree_partitions"].get("16",0)}、p0={data["pcs_tree_partitions"].get("0",0)}；但本次最终解路径的全部{sum(data["pcs_path_partitions"].values())}条边均用p0。因此我证实了两个真实划分进入搜索及其分支资格，没有把最终解描述为一定经历p16↔p0切换。')
    r.p("PCS和NR的seed7重复去除计算计时后完整数据一致；RSS两次均TIME_LIMIT，但达到截止时刻前完成的工作量不同，未满足原预声明的完整数据一致性。这个重复缺口保留，最终结果为部分完成，不因‘原因可解释’而改判。")

    r.page("14 与原论文对应比较及尚未覆盖项")
    r.fig("10_termination_and_qualified_path_length","图10 全部终止状态与成功条件下的路径长。缺失RSS路径不画零长度。",width=6.2)
    r.table(["维度","原论文","我的实际复现","可比资格"],[["平面机构","20R；原几何未完整公开","同拓扑20R独立重建","机制可比，数值不等同"],["PCS划分数量","Fig3 m=1,2,3,7,20","m=2，p0/p16","未复现m增长曲线"],["搜索/后端","OMPL/KDL/PQP，单法调优距离","共用Python/NumPy/GEOS，统一.6rad","不能直接比原秒数"],["空间实验","ABB/Comau，500次/组合","没有运行空间模型","TableI/II仅原文参照"],["RSS","显式奇异偏置γ=.2","单位桥通过，统计组0/20","未取得全局合格路径"],["PCS/NR","环境相关的成本/可见性差异","同条件各20/20；NR此设置更快","非普遍排名"],["重复性","原统计500/点对50000","5seed/条件+3资格重复","RSS全payload重复未通过"],["安全","离散碰撞与局部连接","真实平面净距/离散点资格","不是连续/物理证明"]],size=8.2)
    r.p("E0和E1中，PCS/NR对同seed得到相同路径和节点记录，新增障碍没有改变已找到的解。这说明E1没有形成足够强的路径绕障压力。我不能把它当作原文高密度IV–VI环境，更不能据此否定原文‘复杂障碍下PCS有成本优势’的观察。")
    r.p("我的三策略已各自真实运行并形成统一统计，但原论文六环境、m增长实验、完整可见性曲线、单策略最优距离和RSS全局解尚未复制。结论是同条件重建比较已取得数据，整体复现仍部分完成；未对失败作参数补偿。")

    r.page("15 代码、原始数据与可重复计算")
    r.table(["文件/版本","内容"],[["baselines/p4_comparison/model.py","20R FK/Jacobian、两解析IK、GEOS"],["strategies.py / planner.py","RSS/PCS/NR、RBS、共同双向树"],["configs/engineering/p4_three_methods_v1.json","冻结条件，基线d5b5096/协议6d64b26"],["results/20261013_P4_three_methods_round3","60试验+3重复gzip，source hash、tests/summary"],["round1/round2","原失败与部分ABORTED，未覆盖"],["scripts/build_p4_three_assets.py","保存数据→10幅中文图/派生CSV"],["scripts/check_p4_three_evidence.py","只读artifact/hash/path资格审计"],["scripts/build_p4_three_word.py","可编辑中文Word，仅存量出图"],["旧P4_论文复现汇总.docx/PDF","原NR历史23页报告原样保留"]],size=8.5)
    r.p("复建图表：python scripts/build_p4_three_assets.py --run-id 20261013_P4_three_methods_round3；随后运行build_p4_three_word.py。不重新计算规划。")
    r.p("独立重跑科学套件时，run_p4_three_methods.py必须指定新的run-id，已有目录会拒绝覆盖。三轮修正后当前结果保留；没有隐式重启和成功后替换原失败。")
    r.p("文献：[1] Sintov et al. 2018，DOI 10.1109/LRA.2018.2846806；[2] CKCplanning固定909bd72fff778aca8682c96f9eb86ae4eb75557b，https://github.com/avishais/CKCplanning；[3] 本次实际源提交3d51ca0及上述结果；[4] Shapely/GEOS Cartesian distance官方文档，https://shapely.readthedocs.io/en/stable/reference/shapely.distance.html。")

    r.page("最后一章：应用于我当前双机械臂研究场景的差距和待解决问题")
    r.p("我的课题是双FR3侧吸盘共同搬运Cube，再进入受限车厢。平面三法比较帮助我理解被动链分支和数值投影，但它没有验证该空间机构的RSS或PCS，也没有证明规划路径可直接执行。")
    r.table(["差距","已有依据","仍待研究"],[["空间解析被动链","平面3R两解可核验；作者ABB6R八解","FR3七自由度子链划分及多分支IK，不能随机命名"],["奇异点覆盖","平面固定桥通过、随机RSS无可用桥","空间真实执行器奇异定义与生成资格，不能近限位代替"],["闭链/工具","旧FR3 NR C/Jc数学已接受","固定抓持与工具坐标必须独立核对"],["环境/碰撞","本次GEOS二维胶囊，旧FR3 FCL历史另存","空间网格、吸盘、车厢与动态连续安全"],["时间轨迹/限位","平面只有空间路径；旧B0余量.236039°","赋时、实际跟踪界、起点建立，不能拿计算时间当运动时间"],["持件内力/执行","尚无本次物理执行","弹性抓持、内部力和实际测量属于后续控制研究"]],size=8.5)
    r.p("我保留旧NR全部成功/失败，不把它们并入此次三法分母。原TASK03历史0.05μm精度未通过、TASK06原障碍协议部分完成、B0执行条件不足仍然成立。新同类模型不会追溯改变这些结论。")
    r.p("这一差距分析没有提前改变三方法以适应车厢，也没有提出已验证有效的算法小改。后续空间应用需要单独建立模型与测量资格；当前不运行Isaac，也不自动进入下一篇论文。")

    for paragraph in r.doc.paragraphs:
        for forbidden in ["编写目的","根据用户要求","本报告旨在","按照导师要求","ChatGPT","Codex"]:
            assert forbidden not in paragraph.text
    target=ROOT/"reports/P4_三方法复现汇总.docx"
    r.doc.save(target)
    with zipfile.ZipFile(target) as z:
        xml=z.read("word/document.xml").decode()
        checks={"editable_equations":xml.count("<m:oMath>"),"figures":xml.count("<w:drawing>"),
            "docx_sha256":hashlib.sha256(target.read_bytes()).hexdigest(),"raw_sources_verified":len(data["source_sha256"]),
            "statistical_trials":60,"rss_success":0,"pcs_success":20,"nr_success":20,
            "full_protocol":"PARTIAL","new_planning_calls":0,"historical_docx_overwritten":False}
    (ASSETS/"word_build_check.json").write_text(json.dumps(checks,ensure_ascii=False,indent=2)+"\n")
    print(json.dumps(checks,ensure_ascii=False))


if __name__=="__main__":main()
