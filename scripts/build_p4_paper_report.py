#!/usr/bin/env python3
"""Editable Chinese Word report from already-saved P4 evidence only.

Requirements: python-docx==1.2.0, numpy, matplotlib (asset generator).
No scientific modules, simulator, ROS, subprocess or command publishers.
"""
from __future__ import annotations

import argparse
import hashlib
import json
from collections import Counter
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
ASSETS = ROOT / "reports/p4_report_assets"
SOURCES = Path("/home/ubuntu2004/lmy/papers_and_reference_all/references/private_sources")
EXPECTED = {
    "P1_original.pdf": "8b4589ce07f932bf1294c7a39b30320a8791f630a4bcd2cd71782b3a2a56de3c",
    "P2_original.pdf": "0188b6be5f9029b40c99b377bcd1e2885ec6e6b56acc49276d6fe1fc72cc0346",
    "P3_original.pdf": "d16235aa3af25e8e8a52299abb8d1e47c5f9190b3e766f2450cd963687139bd1",
    "P4_original.pdf": "193710e5a9d4284cf62b30ff189a1df038a9d5c5ce3f1545525dc2f78bea08c6",
    "P5_original.pdf": "6f367be45e0c925df64ec5e4ab3046e1de8aa96627c65ddb2fd71720d9eedc02",
    "DEFORM_report_example.docx": "d036bd716f89498d379d3f7bdd734542f04e6d5019ed3870a27eeabfa4a64c2d",
}


def text_run(value):
    r = OxmlElement("m:r")
    t = OxmlElement("m:t"); t.text = value
    r.append(t)
    return r


def expr(parent, value):
    if isinstance(value, list):
        for part in value: expr(parent, part)
    elif isinstance(value, str): parent.append(text_run(value))
    else: parent.append(value)


def sub(base, index):
    el = OxmlElement("m:sSub")
    e = OxmlElement("m:e"); expr(e, base); el.append(e)
    s = OxmlElement("m:sub"); expr(s, index); el.append(s)
    return el


def sup(base, exponent):
    el = OxmlElement("m:sSup")
    e = OxmlElement("m:e"); expr(e, base); el.append(e)
    s = OxmlElement("m:sup"); expr(s, exponent); el.append(s)
    return el


def fraction(numerator, denominator):
    el = OxmlElement("m:f")
    for tag, value in [("m:num", numerator), ("m:den", denominator)]:
        e = OxmlElement(tag); expr(e, value); el.append(e)
    return el


def matrix(rows):
    el = OxmlElement("m:m")
    for row in rows:
        r = OxmlElement("m:mr")
        for cell in row:
            e = OxmlElement("m:e"); expr(e, cell); r.append(e)
        el.append(r)
    return el


def brackets(value):
    el = OxmlElement("m:d")
    pr = OxmlElement("m:dPr")
    for tag, val in [("m:begChr", "["), ("m:endChr", "]")]:
        e = OxmlElement(tag); e.set(qn("m:val"), val); pr.append(e)
    el.append(pr); e = OxmlElement("m:e"); expr(e, value); el.append(e)
    return el


class Report:
    def __init__(self):
        self.doc = Document()
        sec = self.doc.sections[0]
        sec.page_width = Cm(21); sec.page_height = Cm(29.7)
        sec.top_margin = Cm(1.8); sec.bottom_margin = Cm(1.8)
        sec.left_margin = Cm(2); sec.right_margin = Cm(2)
        sec.header_distance = Cm(.7); sec.footer_distance = Cm(.7)
        styles = self.doc.styles
        for name in ["Normal", "Title", "Subtitle", "Heading 1", "Heading 2", "Caption"]:
            styles[name].font.name = "Noto Serif CJK SC"
            styles[name].element.rPr.rFonts.set(qn("w:eastAsia"), "Noto Serif CJK SC")
        styles["Normal"].font.size = Pt(10.5)
        styles["Normal"].paragraph_format.line_spacing = 1.12
        styles["Normal"].paragraph_format.space_after = Pt(5)
        for name, size in [("Title",24),("Heading 1",16),("Heading 2",12),("Subtitle",12),("Caption",9)]:
            styles[name].font.size = Pt(size)
            styles[name].font.color.rgb = RGBColor.from_string("203F55")
        for name in ["Heading 1", "Heading 2"]:
            styles[name].paragraph_format.space_before = Pt(4)
            styles[name].paragraph_format.space_after = Pt(7)
        h = sec.header.paragraphs[0]; h.text = "P4 闭链运动规划论文复现汇总  |  已有离线证据整理"
        h.runs[0].font.size = Pt(8)
        f = sec.footer.paragraphs[0]; f.alignment = WD_ALIGN_PARAGRAPH.CENTER
        f.add_run("第 ")
        field = OxmlElement("w:fldSimple"); field.set(qn("w:instr"), "PAGE"); f._p.append(field)
        f.add_run(" 页  ·  科研 Benchmark：草案 / 未定稿")
        for run in f.runs: run.font.size = Pt(8)
        self.doc.core_properties.title = "第一篇P4闭链运动规划论文复现汇总"
        self.doc.core_properties.subject = "数学模型与已有离线仿真实验；NR路线；RSS/PCS尚未复现"
        self.doc.core_properties.author = "双机械臂协调操作复现项目"
        self.doc.core_properties.keywords = "P4, 闭链, FR3, NR, RBS, RRTConnect, 离线复现"

    def page(self, heading):
        self.doc.add_page_break()
        self.doc.add_heading(heading, 1)

    def p(self, value, small=False):
        p = self.doc.add_paragraph(value)
        if small:
            for r in p.runs: r.font.size = Pt(9)
        return p

    def h(self, value): self.doc.add_heading(value, 2)

    def eq(self, value, label):
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(3); p.paragraph_format.space_after = Pt(7)
        math = OxmlElement("m:oMath"); expr(math, value); p._p.append(math)
        p.add_run("    (" + label + ")").font.size = Pt(9)

    def table(self, headers, rows, widths=None, size=9):
        t = self.doc.add_table(rows=1, cols=len(headers)); t.style = "Table Grid"
        t.autofit = False
        if widths:
            for c, w in zip(t.columns,widths): c.width = Cm(w)
        for c, value in zip(t.rows[0].cells, headers): c.text = str(value)
        pr = t.rows[0]._tr.get_or_add_trPr(); pr.append(OxmlElement("w:tblHeader"))
        for row in rows:
            for c, value in zip(t.add_row().cells, row): c.text = str(value)
        for i,row in enumerate(t.rows):
            pr = row._tr.get_or_add_trPr(); pr.append(OxmlElement("w:cantSplit"))
            for c in row.cells:
                for p in c.paragraphs:
                    p.paragraph_format.space_after = Pt(3)
                    p.paragraph_format.line_spacing = 1.04
                    for r in p.runs: r.font.size = Pt(size); r.bold = i==0
                if i==0:
                    shade = OxmlElement("w:shd"); shade.set(qn("w:fill"),"E8EFF4")
                    c._tc.get_or_add_tcPr().append(shade)
        self.doc.add_paragraph().paragraph_format.space_after = Pt(0)
        return t

    def fig(self, stem, caption, width=6.45):
        p = self.doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.keep_with_next = True
        p.add_run().add_picture(str(ASSETS/(stem+".png")),width=Inches(width))
        p = self.doc.add_paragraph(caption, "Caption"); p.alignment = WD_ALIGN_PARAGRAPH.CENTER


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--private-sources",type=Path,default=SOURCES)
    args = parser.parse_args()
    verified = {}
    for name, expected in EXPECTED.items():
        p = args.private_sources / name
        if not p.is_file(): raise FileNotFoundError(f"Missing private source: {p}")
        sha = hashlib.sha256(p.read_bytes()).hexdigest()
        if sha != expected: raise ValueError(f"Private source hash mismatch: {name}")
        verified[name] = sha
    data = json.loads((ASSETS/"report_data.json").read_text())
    m3 = json.loads((ROOT/"results/20261009_TASK03_closure02/model_metrics.json").read_text())
    m4 = json.loads((ROOT/"results/20261009_TASK04_projection03/projection_metrics.json").read_text())
    m6 = json.loads((ROOT/"results/20261012_TASK06_constrained_search02/planning_metrics.json").read_text())
    b0 = data["b0"]
    r = Report()
    d = r.doc
    d.add_heading("第一篇P4闭链运动规划\n论文复现汇总",0)
    d.add_paragraph("数学模型与仿真复现实验", "Subtitle")
    r.p("整理日期：2026年10月12日；规划基线 e32d19d；独立分支 report-p4-paper-first。")
    r.p("论文：Motion Planning of Fully-actuated Closed Kinematic Chains with Revolute Joints: A Comparative Analysis。")
    r.p("作者：Avishai Sintov、Andy Borum、Timothy Bretl。IEEE Robotics and Automation Letters，2018，3(4)：2886–2893。DOI：10.1109/LRA.2018.2846806。[1]")
    r.h("报告结论与范围")
    r.p("我将原文的牛顿投影（NR）、递归二分（RBS）与约束双向树组合，映射到真实双FR3模型，并依据已保存日志核对其数学与离散规划结果。本文不是三方法完整性能复现：RSS、PCS及三策略对比尚未实现，也没有本P4路径的Isaac物理执行成功证据。")
    r.table(["已完成部分","证据与限制"],[
        ["闭链数学与雅可比","435组真实模型导数核验通过；原293条构型精度仍未通过"],
        ["NR投影与RBS","194/194邻域投影、293/293历史诊断；10个真实FCL连接正例"],
        ["约束双树规划","18调用 / 9重复对，14成功；原障碍专项未完成，TASK06整体部分完成"],
        ["既有场景离线适配","1次同源名义规划：96节点、94条解树边、1869密集点"],
        ["执行资格","左J4最小余量0.236039°；B0执行准备条件不满足，不能下发"],
    ],[3.5,13])
    r.p("本次只做文献阅读、存量数据分析及报告生成：新规划、IK、投影、FCL查询、Isaac/ROS/机器人实验均为0。历史日志、原阈值、抓取变换、场景、ACM与benchmark_v1.yaml保持不变。所有图横轴使用真实记录的迭代次数或路径点编号，不伪造物理时间。")
    r.p("五份原始论文及DEFORM范例均已找到并通过SHA256核对；本次实际阅读全文为P4八页和DEFORM范例。资料仅本地阅读，公开仓库不上传原PDF和范例。正文采用“数学模型 / 仿真复现实验”两部分；报告待用户审查，不自动开始P2。",small=True)

    r.page("第一部分  数学模型")
    r.h("1 研究问题：闭合、可连接与无碰撞不是同一件事")
    r.p("两个末端与物体刚性连接后，独立改变一只臂的关节一般会破坏闭合关系。合法构型位于高维关节空间中的低维约束集合，普通随机采样几乎不会恰好满足等式。P4讨论的核心问题是：如何将采样和局部连接限制在闭链可行集合，并在限位与障碍条件下搜索路径。[1，II-A，p.2]")
    r.eq([sub("C","cl")," = {q ∈ C : C(q) = 0},   ",sub("C","o")," = ",sub("C","free")," ∩ ",sub("C","cl")],"原文(1)/(2)，记号改为q")
    r.p("Cfree同时考虑关节边界和碰撞。C(q)=0只说明几何闭合，不自动说明无碰撞；一条连续路径还需从给定起点连接到终点。原文默认抓取刚性、机械臂全驱动、环境与状态已知；它并未验证吸盘力学、内力分配或真实跟踪误差。")
    r.h("1.1 三种约束处理策略与三种规划器")
    r.table(["策略","原理与研究差异","本次状态"],[
        ["RSS：随机奇异构型采样","主动/被动链APC；不同IK分支可能被人工驱动奇异集隔开，通过奇异构型连接分支。","未实现"],
        ["PCS：被动链切换","不同被动链对应不同奇异集；切换计算分组扩大局部可见集，依赖可用闭式IK。","未实现"],
        ["NR：牛顿数值投影","对完整关节构型作约束雅可比伪逆修正；不要求每个被动链都有闭式IK。","主要实现路线"],
    ],[3.2,10.5,2.8])
    r.p("上述是三种约束处理策略；RRT、Bi-RRT、SBL是原文比较的三种采样规划器，二者不可混称。本次只组成NR路线的受约束双向树，没有复现RRT/SBL及三策略×三规划器完整矩阵。[1，III、IV]")
    r.fig("01_nr_planning_flow","图1 本次NR规划组合的原创流程示意。依据原文II-B/C、IV及已固定作者增长顺序整理；非新实验。")

    r.page("2 原文可见性分析与双FR3系统模型")
    r.h("2.1 RSS/PCS为什么与NR不同")
    r.p("APC将关节分为主动部分和被动部分，先采样主动关节，再由闭式IK补全被动链。这里“被动”是求解分组，不表示机械臂关节没有电机。单一被动链的不同IK分支可能形成不相连的局部域；RSS显式采样人工奇异构型，PCS改用另一被动链绕开这一分组的奇异集。[1，II-B.1、III-A/B，pp.2–5]")
    r.eq(["V(q) = {q′ ∈ ",sub("C","o")," : valid_edge(q,q′)}"],"原文(3)")
    r.p("可见性是“在指定局部连接器下可以直接连接”的集合，不是视野或传感器覆盖。原文对独立被动链的分支模型给出不可共域概率的乘积；依赖链另有条件概率修正，不能把独立性假设直接用在任意冗余FR3分组。")
    r.eq([sub("P","m")," = ",sup(sub("∏","k=1"),"m"),"(1 − ",sub("p","k"),"),   ",sub("p","k")," = 1/",sub("ξ","k"),",   ",sub("P","v")," = 1 − ",sub("P","m")],"原文(6)/(9)独立链含义")
    r.p("ξk是相应被动链的IK分支数。式(6)的计数采用独立被动链及分支等测度近似；分支不等测度或相关链时不能直接套用，障碍与关节限位还会另外限制可见性。它不是我们的NR成功率公式。本项目尚未实现被动链枚举、奇异集识别或这组可见性统计。NR更易通用化，代价是数值投影开销与投影后分布的不均匀性。[1，III-C、IV-B]")
    r.h("2.2 坐标系与固定抓取方向")
    r.p("我采用Tᴬ_B将B坐标变换到A；q=[qL1…qL7,qR1…qR7]，角度为rad、平移为m。FK沿world→base→FR3关节→link8法兰→侧吸TCP，严格复用原模型，不将法兰或双TCP中点当作杯面TCP。")
    r.eq([sub("F","i"),"(",sub("q","i"),") = ",sup(sub("T","TCP_i"),"W"),",   ",sub("G","i")," = ",sup(sub("T","O"),"TCP_i"),",   ",sub("T","O,i")," = ",sub("F","i"),"(",sub("q","i"),")",sub("G","i")],"R1：模型适配")
    r.p("旧配置H_i=Tᴼ_TCP_i是TCP在物体系中的位姿，将TCP坐标映到O；G_i=H_i⁻¹则将O坐标映到TCP。左右臂各自预测同一物体位姿；固定抓取不能在每个路径点重拟合，否则会掩盖闭链误差。[2，坐标约定]")
    r.fig("02_frames","图2 双FK与固定抓取的原创坐标链示意。只说明变换方向，不是机械臂形状或实际状态截图。")

    r.page("3 闭链约束与约束雅可比")
    r.h("3.1 对原文通用约束的SE(3)具体化")
    r.p("从左右预测位姿Ti=[Ri,pi]建立三维位置残差和三维旋转残差。以下R编号公式是双FR3适配，原文只规定一般C(q)=0，没有印出这些具体6×14公式。[1，II-A/B；2]")
    r.eq(["C(q) = ",brackets(matrix([[[sub("p","L")," − ",sub("p","R")]],["φ"]])),",   φ = Log(",sup(sub("R","R"),"T"),sub("R","L"),")∨"],"R2")
    r.p("位置行在世界系，单位m；旋转行在右臂预测物体系，单位rad。旋转采用SO(3)对数：四元数q与−q表示同一旋转，不能直接相减，也不能用欧拉角差替代。π附近对数分支不唯一，雅可比拒绝θ≥π−1e−8的图表奇点，这不是正式Benchmark姿态门限。")
    r.h("3.2 从原生TCP几何雅可比得到物体雅可比")
    r.p("原生世界几何雅可比J_i=[V_i;Ω_i]以TCP原点为参考。固定抓取的平移形成力臂r_i，使预测物体原点的速度随腕部转动发生变化。")
    r.eq([sub("r","i")," = ",sub("R","W,TCP_i"),sub("g","i"),",   ",sub("B","i")," = ",sub("V","i")," − ",sub(brackets(sub("r","i")),"×"),sub("Ω","i")],"R3")
    r.eq([sub("J","c")," = ",brackets(matrix([
        [sub("B","L"),["−",sub("B","R")]],
        [[sup(sub("J","l"),"−1"),"(φ)",sup(sub("R","R"),"T"),sub("Ω","L")],
         ["−",sup(sub("J","l"),"−1"),"(φ)",sup(sub("R","R"),"T"),sub("Ω","R")]]]))," ∈ ",sup("ℝ","6×14")],"R4")
    r.eq([sup(sub("J","l"),"−1"),"(φ) = I − ½",sub(brackets("φ"),"×")," + a(θ)",sup(sub(brackets("φ"),"×"),"2"),",   θ=‖φ‖"],"R5")
    r.eq(["a(θ) = ",fraction(["1 − (θ/2)cot(θ/2)"],sup("θ","2"))," ≈ 1/12 + θ²/720 + θ⁴/30240"],"R6：小角度展开")
    r.p("B_i的负叉乘来自Ω×r=−[r]×Ω；旋转导数由SO(3)左雅可比逆乘右物体系内的相对角速度。MoveIt组雅可比先将线速度/角速度两个块旋转到world，不使用空间螺旋的平移伴随项，避免混淆雅可比定义。")
    r.h("3.3 如何验证，而不是只看公式")
    r.eq([sub("J","FD,j")," = ",fraction(["C(q+h",sub("e","j"),") − C(q−h",sub("e","j"),")"],"2h")],"R7：中央差分")
    r.p("已有实验分别使用三组h及五点差分核对解析导数；真实FR3模型和固定抓取保持不变。数学一致G*只在一个构型上构造一次作为测试夹具，不能替代历史抓取、物理校准或传感器实测。第9节保留原历史精度失败与导数通过的不同结论。")

    r.page("4 Newton–SVD约束流形投影")
    r.p("对不闭合的候选构型，复用R2、R4计算残差和导数，以最小范数关节修正减少闭链误差。此式来自原文NR思想；具体FR3残差、停止条件与数值护栏是本项目工程协议。[1，II-B.2，IV-A；3]")
    r.eq([sub("q","k+1")," = ",sub("q","k")," − ",sup(sub("J","c"),"†"),"(",sub("q","k"),")C(",sub("q","k"),")"],"R8")
    r.eq([sub("J","c")," = UΣ",sup("V","T"),",   ",sup(sub("J","c"),"†")," = V",sup("Σ","†"),sup("U","T")],"R9")
    r.p("SVD奇异值σ超过max(1e−12,1e−10σmax)才取倒数，其余置零。位置/旋转未作行加权，伪逆求解并不意味着可用混合量纲范数判断物理成功；每次收敛分别核对位置和旋转。")
    r.eq(["‖",sub("e","p"),"‖ ≤ ",sub("ε","p"),",   ‖",sub("e","R"),"‖ ≤ ",sub("ε","R"),",   ",sub("ε","p"),"=1e−8 m,  ",sub("ε","R"),"=1e−8 rad"],"R10：预声明工程标准")
    r.table(["步骤/条件","已有实现行为","性质"],[
        ["初值已经闭合","直接收敛；不无谓执行SVD修正","数值工程"],
        ["需修正时rank<6","拒绝，不宣称一般奇异构型可通行","保守差异"],
        ["整步更新超过原生限位","拒绝，不截断、不wrap、不clamp","保守工程策略"],
        ["更新与迭代上限","最多40次；步长≤1e−14 rad判断停滞","预声明预算"],
        ["阻尼/回溯/尺度","未使用；原文提及可选backtracking","选择与原文允许项区分"],
        ["记录","初/末残差、q改变量、每步奇异值/秩、耗时和原因","轻量可追溯日志"],
    ],[4,9.2,3.3])
    r.p("投影只解决闭合与原生限位，不查询FCL。即使数值收敛，也可能碰撞、发生不允许的关节分支变化，或无法建立真实吸附；这些由局部连接与后续资格分别检查。")
    r.h("4.1 数值收敛与原论文完整性的边界")
    r.p("原文所选各研究起终点之间的路径需要穿越奇异构型；当前rank护栏和局部性条件实际缩小可接受集合，因此本次不是一般奇异通道性能的复现。NR路线可组合成立与论文所有场景/结论成立是两种证据，不互相替代。原始0.05 μm历史精度失败也不能被后续投影新q覆盖。")

    r.page("5 递归二分RBS局部连接")
    r.p("关节空间两端闭合，不保证直线中间闭合。RBS先在线段中点投影，再对两半递归；每个实际检查点分别验证残差、限位和原生FCL。它构成规划树的边资格，而不是最后对普通MoveIt轨迹作修补。[1，II-C，p.3；4]")
    r.eq([sub("q","mid,raw")," = ½(",sub("q","a"),"+",sub("q","b"),"),   ",sub("q","mid")," = Project(",sub("q","mid,raw"),")"],"R11")
    r.eq(["Connect(",sub("q","a"),",",sub("q","b"),") ⇒ Connect(",sub("q","a"),",",sub("q","mid"),") ∧ Connect(",sub("q","mid"),",",sub("q","b"),")"],"R12：递归关系")
    r.p("伪代码：先核验精确端点；若间距达到叶分辨率则接受叶；否则投影中点并核验资格、投影改变量；随后连接左半、右半；任一失败拒绝整条连接，保存明确原因和从起点顺序可达的已核验前缀。不把散落的检查点全部称为可达路径。")
    r.table(["参数","TASK05预声明值","意义"],[
        ["叶关节L2分辨率","0.010 rad","递归停止分辨率，不是笛卡尔毫米"],
        ["相邻L2 / Linf","各≤0.010 rad","防止离散关节跳变"],
        ["投影中点修正L2","≤0.020 rad","防止连接离开局部区域"],
        ["递归深度 / 投影次数","12 / 1023","有界计算"],
        ["最多输出点","1025","局部连接预算"],
        ["位置 / 旋转闭链","≤1e−8 m / ≤1e−8 rad","沿用NR工程精度"],
        ["碰撞与许可","原PlanningScene/FCL、原SRDF/ACM","不扩大ACM以取得通过"],
    ],[4.5,5,7])
    r.h("5.1 真实碰撞后端与合法接触")
    r.p("完整14关节RobotState由已有MoveIt模型提供，自碰撞（含臂间）与机器人—环境碰撞分开保存。刚体Cube以原生AttachedBody随左TCP运动，右TCP由闭链约束保持一致；这只是离线几何表达，不是PhysX双吸盘附件。只保留原SRDF已有许可，持件touch-link列表为空；不新增杯面/Cube豁免。")
    r.p("离散FCL通过说明检查点没有不许可碰撞，不提供连续扫掠证明。现有后端没有查询原生最近距离，所有报告将它标为UNAVAILABLE_NATIVE_DISTANCE_NOT_QUERIED，不能把空接触列表解释为正净空。真实环境碰撞负例已覆盖；未单独构造真实自碰撞/臂间碰撞负例的证据也必须明确。")

    r.page("6 受约束双向随机树与复现忠实性")
    r.p("我使用两棵树分别以已合格的q_start、q_goal为根。每轮交替主树，在14维原生关节边界内均匀采样，查找欧氏最近邻，执行有界steer→NR→RBS/FCL；另一树朝主树末节点重复连接。精确汇合后回溯父边、反向终点树的密集边并只删除重复汇合点。[1，IV；5、8]")
    r.eq([sub("q","raw")," = ",sub("q","near")," + min(1,η/‖",sub("q","rand"),"−",sub("q","near"),"‖)(",sub("q","rand"),"−",sub("q","near"),")"],"R13：steer")
    r.p("固定作者CKCplanning提交909bd72的NR增长顺序给出重复EXTEND/CONNECT及最终精确目标处理；本项目上限101/501来自作者含端点循环，不是原文印刷算法参数。每条新增边调用原TASK05连接器；只验证树节点或将普通RRTConnect轨迹末端修正都不满足此实现。")
    r.table(["类别","本项目对应","必须披露的边界"],[
        ["原论文方法 [ORIGINAL]","NR + RBS + 双向采样树","只对应NR路线；未复现三法矩阵"],
        ["双FR3适配 [ADAPTATION]","14维FK、具体SE(3)残差、MoveIt/FCL","不同于论文ABB/Comau、PQP"],
        ["软件组织 [ENGINEERING]","mt19937_64、稳定最近邻、完整边资格与预算","不是新的控制器/执行链"],
        ["保守差异 [DEVIATION]","rank<6拒绝、局部投影/相邻步门禁","不能宣称论文的一般奇异通过结论"],
        ["工程实验 [EXPERIMENTAL]","私有障碍、小样本固定种子与Task26场景输入","不是原论文六场景坐标复刻"],
    ],[3.8,6.7,6])
    r.h("6.1 参数及终止规则")
    r.p("原始推进η=0.20 rad；候选投影改变量≤0.25 rad；接受扩展≤0.45 rad；进度>1e−12 rad；总节点≤600、外迭代≤800、规划计算预算20 s。RBS/NR内部也有独立预算，回调返回后再检查总时间，不声称硬抢占任意回调。最近邻线性扫描、等距选先插入节点，随机u=(engine()>>11)×2⁻⁵³。")
    r.p("没有goal bias、捷径、平滑、路径优化或直接返回快路。相同输入/seed对比实际路径、终止状态及搜索计数，wall-time不要求相同。预算耗尽只说明此次搜索未在预算内得到解，不证明任务几何不可行。原作者数值参数没有盲目继承：其RBS分辨率、KDL精度和角度wrap与本工程协议不同。")

    r.page("第二部分  仿真复现实验")
    r.h("7 实验平台与来源：本轮没有新增实验")
    r.p("本文“仿真复现实验”指已保存的运动学/离线规划计算，不等同于Isaac物理机器人执行。我从固定结果JSON/CSV/gzip读取数值，只重新出图与制作报告，没有调用FK、IK、NR、规划或FCL。真实模型测试与纯合成单元样例在下面分开列出。")
    r.table(["项目","原论文","本项目已有记录"],[
        ["模型/维度","20维平面；双ABB/Comau及三Comau场景","真实双FR3，14维，固定L型侧吸TCP"],
        ["软件","C++、OMPL、KDL/SVD、PQP","C++/Eigen；MoveIt core 2.5.9 FK/FCL；g++11.4.0/Python3.10.12已有记录"],
        ["计算机","i7-6700HQ、Ubuntu、8GB RAM","原科学日志未完整登记可比CPU/内存，不据当前机器补认证"],
        ["约束精度","原文/作者实现参数，非我们的统一阈值","TASK03历史位置5e−8m；TASK04–06位置/旋转各1e−8，均工程协议"],
        ["种子/次数","Table I每方法/规划器500次；可见性5万对","TASK06三个seed×三场景×两重复；IS-A只seed7一次"],
        ["物理时间/采样周期","规划分析，非真实FR3闭环控制","没有执行时间；simulation_time/physics_step为空；不是50%或100%运动"],
        ["来源一致性","原文工业机械臂/场景","旧Task26固定631b1f6/批准修正5ed0c96；P4原数据不覆盖"],
    ],[3.2,6.4,6.9])
    r.h("7.1 既有协议与证据等级")
    r.table(["阶段","真实验证范围","原始结果保持"],[
        ["TASK03","真实FK、解析雅可比/有限差分；历史抓取诊断","数学初步验收；原历史精度未通过"],
        ["TASK04","固定数学构型邻域与独立历史投影诊断","仅离线投影数学初步验收"],
        ["TASK05","原模型RBS与真实FCL离散状态","仅离散局部连接初步验收"],
        ["TASK06","真双树、固定种子与真实FCL","核心软件初步验收；整体部分完成"],
        ["IS-A / B0","同源名义场景离线规划 / 固定路径执行前审查","A初步验收；执行准备不满足"],
    ],[3.2,8.4,4.9])
    r.p("所有成功阈值属于此前锁定的工程测试标准，不是正式Benchmark。TASK02-D原生物理时间/步/reset身份测量认证仍暂缓；旧两Cube工程成功不能算入P4路径物理成功率。",small=True)

    r.page("8 闭链数学与真实FR3雅可比验证")
    r.fig("03_closure_and_jacobian","图3 原历史293条构型的闭合残差与逐状态有限差分误差。[2] 图中门限保持原值，位置精度未通过。")
    r.table(["指标","保存值","解释"],[
        ["真实雅可比配置","435＝405原抓取/扰动＋29代数夹具＋1坐标变换","不是435次机器人实验"],
        ["中央差分最差元素误差","3.2806663e−10（门限2e−7）","三步长核验通过，按对应行量纲解释"],
        ["五点差分最差元素误差","1.5442696e−10，h=3e−6 rad","真实模型导数通过"],
        ["历史位置最大残差","0.435027826 μm，state196","超过0.05 μm；219/293超限"],
        ["历史旋转最大残差","1.2266396e−5 rad","小于原2e−5 rad"],
    ],[4.5,6.4,5.6])
    r.p("数学实现验收与历史构型精度验收不矛盾：前者检查导数是否正确，后者检查给定q和原抓取是否满足预声明精度。293条记录仅292个唯一q；该残差与历史IK/数值编码的近似精度相符，本报告不作未经验证的原因归因，也不因此修改抓取、阈值或历史结果。图下半仅展示293条原状态，435组总覆盖由保存聚合字段给出。[2]")

    r.page("9 NR投影收敛：真实逐步残差")
    r.p("固定数学一致抓取G*只构造一次，在97个确定性输入上各重复两次，检查97对q输出最大差为0。非零扰动使用单关节正/负0.001、0.01、0.03 rad和已锁定耦合方向；零扰动不需要更新。曲线在收敛处结束，不补零、不延长。[3]")
    r.fig("04_projection_convergence","图4 96条非零扰动的保存收敛曲线；97输入中另含1个不绘制迭代的零更新夹具。橙线为左J1 −0.03 rad的原始3次更新。")
    r.table(["组别","收敛 / 调用","迭代与最终残差","资格"],[
        ["固定G*邻域","194/194；97输入×2","0更新2例、2更新144例、3更新48例；最大末值3.6973e−9m / 9.1305e−9rad","真实FR3数学，不是物理抓取"],
        ["原抓取历史诊断","293/293","全部1更新；最大末值2.5227e−11m / 3.1769e−11rad","新投影q，不覆盖原精度FAIL"],
    ],[3.1,3.4,7,3],size=8.5)
    r.table(["更新k","位置残差 / m","旋转残差 / rad"],[
        [0,"1.8247603e−2","3.0000000e−2"],[1,"2.1739847e−4","5.6624837e−5"],
        [2,"1.0937285e−8","8.3760626e−8"],[3,"8.1217671e−16","7.0011104e−16"],
    ],[3,6.7,6.8])
    r.p("194邻域调用计算总计6.707253 ms（均值0.034573 ms），293诊断总计4.717277 ms；它们是函数内计算耗时，不含完整模型加载和序列化，更不是搬运时间。小扰动和确定性重复不能外推为全关节空间投影100%成功，也不能与论文50万随机投影实验直接比较。")

    r.page("10 RBS局部连接与真实碰撞拒绝")
    r.p("已有真实模型实验10个正例覆盖相同构型、邻近历史投影构型的正反向连接、固定数学G*邻域。3个负例分别验证限位、精确端点闭合与原生FCL拒绝。合成接口/终止测试不计入真实碰撞验收分母。[4]")
    r.table(["实验","实际结果","限定结论"],[
        ["原抓取same/邻近连接5例","5/5 CONNECTED","输入为既有合格/投影构型"],
        ["固定G* same/邻域5例","5/5 CONNECTED","数学夹具，不是吸附实测"],
        ["native_bounds_negative","JOINT_LIMIT","原生关节边界拒绝"],
        ["invalid_exact_closure_negative","CLOSURE","不自动修补不合格精确端点"],
        ["real_fcl_box_at_tcp_negative","COLLISION","左TCP处0.12m障碍盒，真实FCL"],
    ],[6.2,4.2,6.1])
    r.table(["已保存指标","值"],[
        ["正例密集点出现次数 / 中点投影","32 / 14（不是32独立构型）"],
        ["原生FCL状态检查","35次（每次包含self/world），非BVH内部次数"],
        ["正例最大闭链残差","1.8427461e−10 m / 2.1963382e−9 rad"],
        ["正例最小关节余量","0.63396393 rad"],
        ["最大相邻L2 / Linf","0.00965945 / 0.00619235 rad，满足原0.01门限"],
        ["最近碰撞距离","未查询原生距离，不可用"],
    ],[7.5,9])
    r.p("真实负例保存12个工具—障碍接触与4个持件Cube—障碍接触的pair、深度及类别。该负例证明本路径接口确实调用真实FCL，不是合成布尔值；但不代表已独立覆盖所有自碰撞/臂间碰撞负例。附加盒只用于私有测试，未改变正式场景或ACM。")
    r.h("10.1 拒绝原因不得互换")
    r.p("投影失败、限位、残差、局部跳变和碰撞是不同拒绝原因。TASK06私有障碍直连恰恰在投影局部性守卫先失败，后续FCL没有在该中点运行；不能凭本节真实碰撞负例把那个案例改写成COLLISION。已核验前缀只是诊断产物，不是到目标的合格完整路径。")

    r.page("11 固定种子双树规划统计")
    r.p("原锁定条件为3场景×seed{7,41,20261012}×2重复，共18调用。简单端点距0.008315945 rad，非平凡与私有障碍端点距0.921690475 rad。私有障碍盒固定中心(0.638889,0,0.335556)m、尺寸(0.02,0.18,0.12)m；不根据规划结果调位置。[5]")
    r.fig("05_planning_statistics","图5 18次原始调用结果与全部计算时间。9对重复共享输入与seed，不能当作18个独立随机样本。")
    r.table(["统计口径","真实保存值"],[
        ["调用结果","14成功 / 4 NODE_LIMIT；9输入中7成功、2预算耗尽"],
        ["确定性","9/9对状态、路径和搜索计数一致；计算时间排除比较"],
        ["候选NR / RBS中点投影","4504 / 46098；候选失败26、中点失败0"],
        ["FCL状态查询 / 碰撞拒绝查询","55118 / 392；碰撞是被拒搜索候选，不在成功路径上"],
        ["成功路径最大闭合误差","9.9176041e−9m / 9.9838552e−9rad"],
        ["成功路径最小关节余量","0.001980096rad；离线限位合法不代表可执行"],
    ],[7.5,9])
    r.p("首轮编译失败没有科学搜索，不并入18调用；节点预算耗尽不是无解证明。简单问题也可能出现长绕行，因为固定协议没有直接返回、goal bias或优化，当前目标是检验真双树与边资格，而不是路径效率最优。")

    r.page("11.1 九组输入的完整原始结果与障碍缺口")
    rows=[]
    cn={"simple":"简单","nontrivial":"非平凡","obstacle":"私有障碍"}
    for pair in data["task06_pairs"]:
        t,u=pair["repeats"]
        length=f"{t['path_metrics']['joint_path_length_rad']:.6f}" if t["status"]=="SUCCESS" else "无解路径"
        rows.append([cn[pair["case"]],pair["seed"],"成功×2" if t["status"]=="SUCCESS" else "NODE_LIMIT×2",
                     t["node_count"],f"{t['path_tree_edges']} / {t['path_points']}",length,
                     f"{t['compute_wall_time_sec']:.6f}\n{u['compute_wall_time_sec']:.6f}"])
    r.table(["场景","seed","两次状态","树节点","解边 / 点","长度rad","两次计算s"],rows,[1.9,2.3,2.5,1.5,2.5,2.5,3.3],size=8)
    r.p("表1 每行对应同一输入的两次调用，完整18行机器可读记录保存在task06_trials.csv。长度仅对成功路径有效，预算失败不能填0伪装为短路径。[5]")
    r.fig("06_tree_search_projection","图6 私有障碍seed20261012成功搜索的真实二维关节投影：553节点、148解树边。二维投影不表示真实障碍或连续安全。",width=6.2)
    r.p("原障碍协议要求“直接TASK05连接因COLLISION拒绝”，实际却为PROJECTION_JUMP：第2检查点（中点depth1）NR已收敛3步，但修正L2=0.0869493378rad，超过原0.020rad；该点collision=null，FCL未执行。因此核心双树软件经审查初步验收，TASK06整体仍部分完成，不能追溯把障碍协议判为通过。")

    r.page("12 既有Task26场景的无执行规划适配")
    r.p("IS-A先遇到历史归档0.606m与旧Task26实际内宽0.628m不一致，来源门禁停止；独立第二阶段采用源一致v2输入，不覆盖旧归档。名义起点为共同抬起后，名义终点为共同短推完成且释放前；两端都属于双臂共同抓持阶段，但它们是离线目标，不是实测14q。[6]")
    r.table(["项目","独立v2值与来源"],[
        ["左右base / rails","(0.65,−0.60,0)/(0.65,+0.60,0)m；rail均0.650m，world_shift=0"],
        ["持件Cube","task26_cube_1，边长0.120m、质量0.800kg；固定旧源码"],
        ["车厢内区间","x∈[0.910,1.160]，y∈[−0.314,0.314]m；墙厚0.020m/高0.150m"],
        ["桌面与world","task26_table、原launch嵌套task24_table、deep/±Y墙、Cube2，共6对象"],
        ["其他Cube政策","Cube2固定原供料点(0.350,0,0.260)m；休眠后批不在原world，事前声明"],
        ["名义START / GOAL","(0.500,0,0.540) → (0.790,0.060,0.280)m；两端identity姿态"],
        ["左右TCP在物体系中的位姿H","平移(0,∓0.061,0)m；四元数xyzw (√½,±√½,0,0)"],
    ],[4.2,12.3],size=8.7)
    r.p("端点调用已有KDL-LMA（epsilon1e−7、max500）及原NR，候选0通过；不开发新IK，不改变抓取或门限。START以HOME仅作数值初猜，GOAL以已合格START作初猜，不能据此宣称机器人实际从HOME安全到达持件起点。")
    r.table(["关节","START 左 / rad","START 右 / rad","GOAL 左 / rad","GOAL 右 / rad"],
            [[f"J{i+1}",*[f"{q[i+o]:.9f}" for q,o in [(data["start_q_rad"],0),(data["start_q_rad"],7),(data["goal_q_rad"],0),(data["goal_q_rad"],7)]]] for i in range(7)],
            [2,3.625,3.625,3.625,3.625],size=8.5)
    r.p("表2 展示完整14q，以上仅四舍五入供阅读；机器可读全精度仍在start_endpoint.json/goal_endpoint.json，不当作历史物理测量。",small=True)

    r.page("12.1 来源一致的实际离线路径")
    r.table(["指标","原存量结果"],[
        ["科学规划次数 / seed","1 / 7，达到第一条合格路径后停止"],
        ["树节点 / 解树边 / 密集点","96（63起树＋33终树） / 94 / 1869"],
        ["规划计算耗时","0.213194502s，不是运动耗时"],
        ["关节L2路径长度","11.940815548rad，不是物体行程m"],
        ["最大 / RMS位置闭链残差","8.62520825e−9 / 1.24358347e−9m"],
        ["最大 / RMS旋转闭链残差","5.30773506e−9 / 6.97732909e−10rad"],
        ["原生限位最小余量","0.004119662522rad＝0.2360392755°"],
        ["FCL状态检查","规划1964＋端点2＋独立1869＝3835"],
        ["最近FCL距离 / 物理时间","未查询 / 未执行，均不可用"],
    ],[6.8,9.7])
    r.p("独立复核1869点全部闭链/限位/FCL合格，94条解树边均能回溯到原TASK05资格。Cube的两端世界姿态已给定，但闭链本身仅使左右预测物体一致，不把全路径世界姿态锁住；现有日志没有正式物理Cube/TCP时间序列，因此本报告不伪造物体轨迹或保持水平曲线。")
    r.h("12.2 离线适配结论不能扩大")
    r.p("源码c88a80c对应实际计算，输入、25个原artifact和620个source记录已在B0核对。run02的PYTHONPATH启动错误没有规划调用，不算科学失败；原来源门禁失败、启动失败和最终成功日志分别保留。原旧两件Cube工程执行成功也不计入本P4物理成功统计。")
    r.p("空间点列具备真实模型离散资格，却没有每点时间、关节速度和加速度。规划耗时与运动时长无关；时间参数化不能消除关节余量风险，也不能自动建立吸附稳定或连续无碰撞保证。")
    r.h("12.3 数值资格与执行资格之间还有哪些门槛")
    r.p("B0发现当前路径先朝随机探索方向明显绕行再折回；到位容差、全程跟踪包络、原生加速度认证与实际持件START对齐尚未满足。第13节图表只分析原路径，不优化、不重规划、不增加seed。下一步是否独立研究路径质量或开始物理实验，由用户审查决定。")

    r.page("13 原1869点的14关节路径")
    r.fig("07_joint_positions","图7 原14关节角度与原生上下限（红虚线）；灰竖线为index860树汇合点。横轴是空间路径点，不是时间。[6、7]",width=6.35)
    r.p("左右各七关节始终在离线原生范围内；小邻接并不代表动态跟踪余量充分。该点列保持原顺序与全精度，没有平滑、捷径、角度wrap或夹紧。",small=True)

    r.page("13.1 闭链残差合格仍可能不具备执行余量")
    r.fig("08_path_residual_and_risk","图8 原1869点闭链残差与左J4余量；位置/旋转分别计量。限位余量在第861点降至0.236039°。[6、7]")
    r.table(["最危险点","保存数值"],[
        ["零基编号 / 关节","860 / left_fr3_joint4（第861点）"],
        ["关节目标 / 上限","−0.121019662522 / −0.1169rad"],
        ["最近上限余量","0.004119662522rad＝0.236039°"],
        ["相邻点最大L2 / Linf","0.009923957 / 0.005315105rad，仍满足原0.010门限"],
    ],[6.3,10.2])
    r.p("执行必要条件为每关节余量m_j > E_tracking,j + E_alignment,j + E_model,j。误差项需要全过程、朝限位最坏方向的可追溯上界，不能用均值/RMS或终点容差替代。旧到位容差0.035/0.060rad分别是此余量8.496/14.564倍，但不等于认证动态包络，因此既不能判定本路径已越界，也不能声称安全可执行。")

    r.page("13.2 全部关节余量曲线")
    r.fig("09_joint_margins","图9 全部14关节的最近原生限位余量，按min(q−lower,upper−q)逐点计算。灰线为树汇合点，原路径未改变。[7]",width=6.35)
    r.p("图中余量始终非负只代表离散目标限位合法，不认证实际运动跟踪误差；最近FCL距离未查询，不能把关节余量当空间净距。",small=True)

    r.page("13.3 全部关节数值表与绕行诊断")
    r.table(["关节","最小余量 °","零基点","累计变化 rad","端点净变化 rad","反转次数"],
        [[("左" if i<7 else "右")+f"J{i%7+1}",f"{j['min_margin_deg']:.6f}",j['point_index'],
          f"{j['cumulative_travel_rad']:.6f}",f"{j['net_displacement_rad']:.6f}",j['raw_direction_reversals']]
         for i,j in enumerate(b0['joints'])],[1.8,3.1,2.1,3.4,3.8,2.3],size=8.5)
    r.fig("10_joint_travel","图10 原14关节累计变化与端点净变化。直线端点差只是诊断参照，不是已经合格的替代路径。[7]",width=6.1)
    r.p("全路径L2长11.940816rad，端点距离1.865774rad，比值6.399925。左J4累计4.152685rad（237.931°），却净变化0.430656rad；11关节在汇合点反向、全链24次逐关节反向，不等于高频抖动或分支断裂。算法只取首个可行解，未优化路径质量。")
    r.p("旧模型速度界[2.62,2.62,2.62,2.62,5.26,4.18,5.26]rad/s有来源，但没有合格加速度、共同时间表或持件起点对齐。B0结论：执行准备条件不满足。不通过降低原容差或放慢速度冒充问题已解决。",small=True)

    r.page("14 原论文结果对应：不能据不同场景宣称性能优势")
    r.p("原文在六类环境上比较RSS/PCS/NR与RRT/Bi-RRT/SBL，Table I对每方法/规划器取500次平均；“—”表示1000s内未找到解。下面摘录全部空间环境Bi-RRT行，与我们采用双向树的路线对应，但不构成相同测试条件下的数值比较。[1，IV-A，Table I，p.6]")
    r.table(["原文环境 / 对象","RSS平均 s","PCS平均 s","NR平均 s"],[
        ["II：双ABB持杆，三柱障碍","7.27","0.22","0.13"],
        ["III：双ABB持杆，窄通道","23.10","1.06","0.92"],
        ["IV：双Comau持杆，多障碍","175.5","27.18","38.84"],
        ["V：双Comau吸盘持盒，工业场景","330.6","3.81","11.76"],
        ["VI：三Comau持环，多闭链","537.2","3.70","36.63"],
    ],[8.1,2.8,2.8,2.8])
    r.p("原文的综合结论是PCS在拥挤场景通常受益于较廉价投影；NR在部分环境更快，且更容易推广到不同关节数。不能将综合结论理解成PCS每一行都快于NR，Table I环境II/III就存在NR更快的情况。")
    r.table(["对应项","原文实验","本项目保存值 / 可比性"],[
        ["方法路线","RSS/PCS/NR三策略","主要NR；RSS/PCS没有数据，无法比较三法"],
        ["Bi-RRT规划时间","500次平均；PQP、i7-6700HQ","IS-A单次0.213194502s；CPU/场景不对齐，不判优劣"],
        ["场景V共同持盒","双6轴Comau、原工业地图","双7轴FR3 + Task26车厢；不同端点/工具/障碍"],
        ["奇异构型通过","原文所选起终点间路径需穿越奇异构型","保守rank/局部性护栏，资格未完成"],
        ["碰撞/连续性","PQP配合原RBS方法","FCL离散点通过；无原生最近距离/连续扫掠证书"],
        ["物理执行与力学","论文是规划比较，力/静力学为后续研究","本P4路径也未物理执行，不把旧Task26算作替代"],
    ],[3.4,5.8,7.3])
    r.p("这里能比较的是方法构成、实验定义和完成程度，而非运行性能优势。我们亚秒规划或极小闭合残差都不支持“优于原论文”的结论：模型维度、地图、端点、门限、碰撞后端、计算机、样本分母和终止规则并不一致。")

    r.page("14.1 可见性、投影代价与节点数的对应关系")
    r.table(["原文Table II，环境II / Bi-RRT","RSS","PCS","NR"],[
        ["可见比例Pv ×100%","0.85","1.31","23.74"],
        ["最优最大扩展距离d_o","0.6","1.0","2.6"],
        ["平均samples","5544","532","151"],
        ["失败samples比例 %","83.64","20.48","16.17"],
        ["sampling占总计算时间 %","3.03","19.05","27.80"],
        ["平均局部连接 ms","4.97","0.32","1.40"],
        ["局部连接检查 / s","102","1639","508"],
        ["局部连接占总时间 %","50.73","53.82","71.58"],
        ["平均树节点","476","408","13"],
        ["平均路径节点","26","29","7"],
    ],[9,2.5,2.5,2.5])
    r.p("可见比例由50,000组起终点局部连接试验得到，不是全路径规划成功率。原文Fig.4图注写环境I，但pp.6–7正文将该组试验说明为环境II；本文遵循Table II及正文并保留这一内部不一致，不自行改写原论文。[1，IV-B，pp.6–7]")
    r.p("原文还记录无关节限位/碰撞的500,000随机配置投影：APC单次较NR约快3–6倍，但APC成功7.92%、NR成功100%[1，II-B，p.2]。本项目194邻域调用是小扰动确定性重复，无FCL；不能据194/194与其随机全空间试验比较成功率。")
    r.h("本项目哪些指标同义、哪些尚缺")
    r.p("原文平均树节点13不能与本项目1869密集点直接相比：本项目96是树总节点、94是解树边、1869是RBS稠密采样点。我们保留计数类型，而不是把点越多或越少说成性能更好。原文可见性曲线、最优d_o profiling、投影分布均匀性及三方法局部连接代价尚没有同条件复现数据。")
    r.p("当前成功14/18与原文Pv没有相同含义；搜索中的392次碰撞查询是候选拒绝，不是成功路径碰撞次数。min_distance未查询，不能用无碰撞推断正距离。所有实验误差门限是已有预声明工程协议，不替代正式Benchmark成功阈值。")

    r.page("15 已完成、未完成及报告怎样使用")
    r.table(["项目","当前证据","尚需补齐但本轮不开展"],[
        ["C/Jc与NR","真实FR3导数和局部投影数学成立","历史原精度仍失败；一般奇异通过未证"],
        ["RBS与双树","真实FCL离散局部边、真双树与重复一致","原障碍直接COLLISION资格仍缺；路径质量不是最优"],
        ["原文三方法","NR路线已组成","RSS/PCS、三策略/三规划器原实验矩阵"],
        ["场景适配","来源一致名义端点与1869点离线路径","持件实际起点、共同赋时、全程误差与执行余量"],
        ["物理测量与Benchmark","旧工程作为来源；正式基准草案","post-step原生时间/步/reset身份认证仍暂缓"],
    ],[3,6.7,6.8])
    r.p("我的阶段结论是：NR路线的真实模型离线构件和场景适配已具备可追溯证据，但P4完整论文复现尚未完成。原TASK06整体部分完成，B0执行准备不满足，Benchmark保持草案。本文整理不会将失败改判，不批准新的算法/物理试验。")
    r.h("15.1 可复用代码、原始结果与报告层")
    r.table(["路径（仓库相对）","用途"],[
        ["baselines/p4_closed_chain/closure.cpp、projection.cpp","TASK03约束/导数与TASK04 NR–SVD"],
        ["baselines/p4_closed_chain/local_connection.cpp、constrained_rrtconnect.cpp","原TASK05 RBS与TASK06双树，不在报告层运行"],
        ["results/20261009_TASK03_closure02、TASK04_projection03、TASK05_local_connection02","真实模型数学/连接原JSON与日志"],
        ["results/20261012_TASK06_constrained_search02","18个完整gzip原始搜索、协议、汇总、确定性证据"],
        ["results/20261012_TASK06_IS_A_nominal_scene03、TASK06_IS_B0_path_review01","完整14q路径、端点、独立资格、原审查"],
        ["reports/p4_report_assets/","10张中文原创图、派生CSV/JSON、来源SHA与只读检查"],
        ["scripts/build_p4_report_assets.py、build_p4_paper_report.py","仅存量出图/Word生成；不import任何科学执行模块"],
    ],[10,6.5],size=8.2)
    r.p("只重建报告的命令：python3 scripts/build_p4_report_assets.py；安装python-docx==1.2.0的Python再运行scripts/build_p4_paper_report.py --private-sources <本地资料目录>。依赖NumPy/Matplotlib/Noto CJK。脚本仅覆盖自己的报告输出，不覆盖results。原实验重跑命令位于各原报告，需新run-id与单独授权，本轮没有执行。",small=True)
    r.p("后续按P2→P3→P5→P1逐篇全文阅读、核心公式、最小真实实验、原数据与独立中文Word报告推进；必须先审查本报告，不自动开始下一篇。")

    r.page("参考依据与数据来源校验")
    refs = [
        "[1] Sintov A, Borum A, Bretl T. Motion Planning of Fully Actuated Closed Kinematic Chains With Revolute Joints: A Comparative Analysis. IEEE RA-L, 3(4):2886–2893, 2018. DOI:10.1109/LRA.2018.2846806。全文为用户提供8页作者手稿；作者公开入口 https://avishaisintov.wordpress.com/wp-content/uploads/2018/06/sintov-pcs.pdf。",
        "[2] reports/TASK03_CLOSURE_CONSTRAINT01.md；baselines/p4_closed_chain/README.md；results/20261009_TASK03_closure02/model_metrics.json。原数据精度FAIL不改；D047审查仅接受离线数学。",
        "[3] reports/TASK04_NR_PROJECTION01.md；PROJECTION.md；results/20261009_TASK04_projection03/projection_metrics.json。",
        "[4] reports/TASK05_LOCAL_CONNECTION01.md；LOCAL_CONNECTION.md；results/20261009_TASK05_local_connection02/local_connection_metrics.json。",
        "[5] reports/TASK06_PRE_TASK_PROTOCOL.md、TASK06_CONSTRAINED_RRTCONNECT01.md、TASK06_REVIEW_CLOSEOUT.md；results/20261012_TASK06_constrained_search02/planning_metrics.json与18个原始gzip。实际测试源码29a9d5c，D054分项审查。",
        "[6] reports/TASK06_IS_A_SCENE_ADAPTER01.md、TASK06_IS_A_NOMINAL_SCENE02.md；独立Task26 v2输入；results/20261012_TASK06_IS_A_nominal_scene03/。测试源码c88a80c。",
        "[7] reports/TASK06_IS_B0_PRE_EXECUTION_REVIEW.md；results/20261012_TASK06_IS_B0_path_review01/。分析源码e1a2e22；交付fe7d3db。",
        "[8] 作者CKCplanning固定提交909bd72fff778aca8682c96f9eb86ae4eb75557b，NR增长逻辑作为参考，不导入代码；https://github.com/avishais/CKCplanning/tree/909bd72fff778aca8682c96f9eb86ae4eb75557b。",
        "[9] 旧工程dual-arm-embodied-palletizing固定631b1f65656d025c1bb2173e874192f3fe4d355a、批准软件修正5ed0c967401e5fedb93b938e1c3904f4a4c87600；只作工程资产/模型来源，不作新P4执行证据。",
        "[10] 用户私有DEFORM_report_example.docx；仅参考两部分报告结构和图表解释方式，不复制其数据、正文或图像。",
    ]
    for value in refs: r.p(value,small=True)
    r.h("六份私有资料SHA256（不上传全文）")
    r.table(["文件","SHA256（分行显示）"],[[name,sha[:32]+"\n"+sha[32:]] for name,sha in verified.items()],[5,11.5],size=8)
    r.p("本次规划基线e32d19d；报告协议266de05。完整图表字段、原source SHA与派生CSV映射见p4_report_assets/report_data.json；报告只读构建检查见build_assets_check.json。",small=True)
    output=ROOT/"reports/P4_论文复现汇总.docx"
    d.save(output)
    # Independent structural roundtrip: core equations remain native OMML, not images.
    import zipfile
    from lxml import etree
    with zipfile.ZipFile(output) as archive:
        root=etree.fromstring(archive.read("word/document.xml"))
        ns={"w":"http://schemas.openxmlformats.org/wordprocessingml/2006/main","m":"http://schemas.openxmlformats.org/officeDocument/2006/math"}
        content="\n".join(root.xpath("//w:t/text() | //m:t/text()",namespaces=ns))
        equations=len(root.xpath("//m:oMath",namespaces=ns))
        figures=len(root.xpath("//w:drawing",namespaces=ns))
        for required in ["第一部分", "第二部分", "RSS", "PCS", "NR", "219/293", "0.236039", "PROJECTION_JUMP", "部分完成", "执行准备"]:
            if required not in content: raise ValueError(f"Missing scope/negative evidence: {required}")
        if equations<10 or figures!=10: raise ValueError("Equation/figure count mismatch")
    check={"private_sources_sha256_verified":verified,"source_hashes_before_after_unchanged":all(
        hashlib.sha256((ROOT/p).read_bytes()).hexdigest()==sha for p,sha in data["source_sha256"].items()),
        "docx_sha256":hashlib.sha256(output.read_bytes()).hexdigest(),"docx_bytes":output.stat().st_size,
        "native_editable_omml_equations":equations,"figures":figures,"negative_scope_assertions":True,
        "report_only":True,"new_scientific_experiments":0,"page_render_review":"PENDING"}
    if not check["source_hashes_before_after_unchanged"]: raise ValueError("Original input modified")
    (ASSETS/"build_word_check.json").write_text(json.dumps(check,ensure_ascii=False,indent=2)+"\n")
    print(json.dumps({"output":str(output),"editable_equations":equations,"figures":figures,"scientific_runs":0},ensure_ascii=False))


if __name__=="__main__": main()
