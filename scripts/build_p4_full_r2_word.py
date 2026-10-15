"""Chinese R2 stage report from immutable real logs; no solver is imported."""
from __future__ import annotations
import argparse
from datetime import datetime, timezone
import json
from pathlib import Path
import sys
ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT/'scripts'))
import build_p4_paper_report as word
from docx.oxml.ns import qn
from docx.shared import Pt
from docx.enum.text import WD_LINE_SPACING


def load(path):return json.loads((ROOT/path).read_text())


def main():
    ap=argparse.ArgumentParser();ap.add_argument('--assets',type=Path,required=True)
    ap.add_argument('--output',type=Path,required=True);args=ap.parse_args()
    if args.output.exists():raise FileExistsError(args.output)
    word.ASSETS=args.assets.resolve();r=word.Report()
    now=datetime.now(timezone.utc)
    r.doc.core_properties.created=now;r.doc.core_properties.modified=now
    r.doc.core_properties.title='P4完整续研：几何一致修复与真实闭合资格'
    r.doc.core_properties.subject='原模型、数值修订、真实失败和未完成实验分开呈现'
    r.doc.core_properties.keywords='P4,APC,NR,RSS,PCS,六环境,闭合资格,部分完成'
    r.doc.sections[0].header.paragraphs[0].text='P4 闭链规划完整续研  |  几何一致修复与真实闭合资格'
    r.doc.sections[0].footer.paragraphs[0].text='2026年10月15日 · 阶段结果 · 全文复现仍部分完成'
    r.doc.add_heading('P4闭链运动规划完整续研\n几何一致修复与真实闭合资格',0)
    r.doc.add_paragraph('2026年10月15日 · R2真实组件实验',style='Subtitle')
    r.p('我在保留六环境原生失败的基础上，实现了数值一致的解析IK、平面循环3R投影和实际闭合约束NR。原机器人的连杆、基座、抓持变换和关节限位保持不变；数学三角恒等式使用真π，原基座和物理编码器的3.1416字面值继续保留。')
    r.p('空间固定资格共2091项，2088项通过、3项失败；平面解析资格233项及平面NR资格56项通过。这里的条目包含FK、分支回代、正控与负控，不是规划次数。Comau IV的三项NR邻域资格因原关节限位而失败，没有放宽判据。')
    r.p('双臂RSS另有700个固定分层目标、5600次解析IK分支请求，获得19个闭合、限位和被动秩合格的目标，共48个分支候选；尚未进行碰撞或跨分支连接。ABB两个环境的同时肘腕奇异模式各0/100，对应两项资格失败。没有新增合格全局路径，正式Table I调用仍为0，完整论文仍部分完成。')
    r.p('原文为Sintov、Borum、Bretl的Motion Planning of Fully-actuated Closed Kinematic Chains with Revolute Joints: A Comparative Analysis，IEEE RA-L，2018，DOI 10.1109/LRA.2018.2846806。全文模型、式1–9、Fig.1–5和Table I–II已在R0/R1逐项审计；本篇承接原生资格报告，不把组件进展写成全文实验完成。')
    r.page('第一部分：原论文方法与实际修订')
    r.h('1 闭链模型与投影')
    r.eq(['C(q)=0,  ',word.sub('C','o'),'=',word.sub('C','cl'),'∩',word.sub('C','free')],'原文闭链与合法空间')
    r.p('APC固定主动角，在非冗余被动链上解析求解。PCS利用不同划分的局部分支可见性，RSS则请求被动Jacobian秩缺陷的人工奇异构型。NR使用全部约束Jacobian伪逆更新，不能用APC回退伪装成NR。RBS递归验证投影后的中点，是之后路径验证的独立关口。')
    r.eq([word.sub('q','k+1'),'=',word.sub('q','k'),'−',word.sup('J','†'),'(',word.sub('q','k'),')C(',word.sub('q','k'),')'],'原论文NR原理')
    r.h('2 解析IK的数值一致修订')
    r.eq(['r=',word.sub('p','5x'),'cos(',word.sub('q','1'),')+',word.sub('p','5y'),'sin(',word.sub('q','1'),')−',word.sub('l','1x')],'选定肩分支的带符号径向')
    r.eq([word.sub('q','2'),'=atan2(r,z)−atan2(B,A),  A=',word.sub('l','2'),'+L cos(',word.sub('q','3'),'+α), B=L sin(',word.sub('q','3'),'+α)'],'独立三角反解')
    r.eq(['W=',word.sup('R','T'),word.sub('03',''),'R=Rx(',word.sub('q','4'),')Ry(',word.sub('q','5'),')Rx(',word.sub('q','6'),')'],'腕部坐标关系')
    r.p('我取消了Comau的FK元素截零以及各模型IK中的小角截零，保留原8个分支编号。零腕角时取q4=0的规范代表，完整连续奇异族并未枚举；肩轴非唯一族和肘边界舍入也不在本轮完整性证明之内。原裁剪目标和同一q产生的连续物理目标分别运行，各自80次/空间环境，不混为一个分母。')
    r.h('3 实际物理闭合的NR')
    r.eq([word.sub('T','j'),'=',word.sub('T','1'),word.sub('Q','1j'),'S,  ',word.sub('e','p'),'=',word.sub('p','j'),'−',word.sub('p','1'),'−',word.sub('R','1'),'d'],'刚抓闭合与位置残差')
    r.eq([word.sub('e','R'),'=Log(',word.sub('R','j'),word.sup('R','T'),word.sub('*',''),')'],'世界坐标姿态残差')
    r.p('原生逆链隐含的精确置换坐标与物理近似基座角不完全一致。我按相同物理变换构建全约束Jacobian，NR上限仍10000次、位置和旋转矩阵门仍各1e−5，不作阻尼、APC兜底或回溯。VI同时求解12×18约束，与公开程序默认两阶段NR不同。当前NumPy参考核用于几何资格，不能与作者C++APC作速度比。')
    r.page('第二部分：真实解析IK与NR实验')
    r.h('4 相同保存输入的IK回代')
    r.fig('01_ik_r1_r2_same_inputs','图1 原生R1、R2原保存目标、R2连续物理目标分别统计；未返回分支继续保留在各80次分母中。',width=6.5)
    r.table(['环境','连续物理目标返回分支/80','最大位置回代差/mm','空间资格通过/失败'],[
        ['II','44','3.41e−13','330/0'],['III','44','3.41e−13','330/0'],
        ['IV','28','1.82e−12','427/3'],['V','24','9.09e−13','430/0'],['VI','20','9.09e−13','571/0']],size=9)
    r.p('500000随机投影、原文APC 7.92%和3–6倍效率尚未复现。此处只说明固定回代反例在新版本中得到修复，不能把44/80或28/80当作原论文随机ambient成功率。固定FK交叉验证、原始端点Jacobian差分和零腕邻域均保存实际输出。')
    r.page('5 NR数值迭代与限位失败')
    r.fig('02_nr_original10_iteration_response','图2 五个空间环境原10个输入的实际NR迭代。横轴为数值迭代次数，分别绘制mm和旋转矩阵误差；不是物体控制时间响应。',width=6.5)
    r.p('空间160次NR包括原50输入和110个额外已闭合组件邻域；154次合法收敛、6次几何闭合后被原限位拒绝。原50输入为47次合法、3次限位拒绝。所有6次拒绝均在IV，臂1的q3低于原下界−1.5708rad。原生KDL的三个false缺少内部返回阶段记录，不能仅因新参考核相同索引越界就把旧原因强行改判为已证实。')
    r.fig('04_iv_three_rejections_q3_bounds','图3 额外IV邻域三项固定资格失败的真实q3及原限位。未把“几何闭合”误当“合法构型”。',width=6.0)
    r.p('IV起点臂1原q3=−1.57rad，距下界仅0.0008rad。无约束伪逆投影可以落到闭合流形上而越过角界；论文的带限位拒绝与无关节界的投影效率实验必须分开。本轮不增加新的限位优化、不改全成功资格，也不删去这三项失败。')
    r.page('6 原始端点与真实奇异条件')
    r.fig('03_raw10_endpoints_separate_units','图4 十个空间原始端点保持字面值；位置误差全部高于1e−5mm。旋转矩阵误差另列，单位不混合。',width=6.4)
    r.p('原始空间端点位置误差约0.00265–0.01486mm。平面原start和goal的位置误差分别为0.0002813与0.0001379个源杆长单位，旋转误差约5.9e−5与4.0e−5。十二个原始端点均未达到当前闭合门。我没有将组件正控的重闭合状态悄悄替换为正式端点；任何精度恢复都需要独立版本、明确分支、逐角差和碰撞资格。')
    r.fig('05_rss_five_fixtures_singular_spectra','图5 固定原肘角、真肘角、腕零及其交集的实际Jacobian谱。秩比较的位置行按预声明除1000，不当作物理混合误差范数。',width=6.3)
    r.page('7 RSS固定闭合构造与平面20R')
    r.eq(['α=atan2(h,k),  ',word.sub('q','3'),'=−α,  ',word.sub('q','5'),'=0'],'肘/腕奇异的几何条件')
    r.p('ABB的(h,k)=(302,70)mm，Comau为(1436.5,210)mm。原ABB肘角−3.1416/2仍为满秩；真肘和腕零在原生连续FK差分与独立解析Jacobian下均失秩。两个奇异条件同时满足不一定独立损失两个秩，不能凭固定两个角宣称余秩2。当前模式没有覆盖肩部全部奇异集合。')
    r.table(['环境/原模式','固定目标','闭合+限位+秩合格目标','仍失败的资格'],[
        ['II 肘腕/腕/肘','各100','0 / 3 / 4','肘腕模式0/100'],
        ['III 肘腕/腕/肘','各100','0 / 3 / 5','肘腕模式0/100'],
        ['V 原作者腕模式','100','4','此构造检查无失败']],size=9)
    r.p('这些700目标来自R1固定q的分层构造，不新增随机数，也不等于原RSS随机规划分布。每个目标遍历8个主动IK分支，5600次请求中48个分支通过数学门，全部拒绝原因保留。19个目标没有碰撞、branch bridge或RBS验证，不能称完整合法RSS目标。环境IV因NR资格失败未进入该构造；VI公开RSS分布未知，重建方案只保留推导。')
    r.p('平面模型保留19根单位杆、ground=(7,4)及物理末角3.1416−q19。一般循环3R解析涵盖全部20被动标签，跨ground段使用真实√65长度；p18对应[19,0,1]、p19对应[18,19,0]。100个FK、80个两端点APC、两点Jacobian及8个解析fixture实际完成，233项检查通过。额外32次平面NR全部合法收敛，56项检查通过。chart合法与179.9°机械限位分别判定。')
    r.h('8 与原论文结果的差距')
    r.p('我尚未执行原六环境RSS/PCS/NR×RRT/Bi-RRT/SBL完整矩阵、Fig.3全部m、500000投影、50000合法点对、Fig.5、无障碍和固定划分对照。这些义务均保留在211入口对照表，没有用本轮组件条目代替。当前没有换链解路径见证，不能比较PCS与NR/RSS的复杂障碍性能，也不能判断原表时间比是否成立。')
    r.p('本轮2核小时/1小时执行额度只用于修复资格；原Table I完整22500次、1000s截断的最坏6250核小时另需资源方案。源码、协议、原始输入输出、限位拒绝和绘图数据分别归档；原P4/P2源码与已接受结论保持。')
    r.page('第三部分：对我当前双机械臂课题的差距与待研究问题')
    r.p('当前实证只支持几何投影组件的有效性，尚不支持把PCS、RSS、NR或RBS的某项性能优势迁移到我的双FR3吸盘搬运与Cube受限装配。六关节解析分支机制与FR3七关节冗余、刚抓与吸盘柔顺、纯运动学路径与接触装配控制之间仍有不同前提。')
    r.p('之后需先完成原论文路径和消融，才决定被动链切换如何与FR3冗余参数衔接、RSS需要哪些真实奇异集合、NR/RBS如何连接现有闭链与碰撞接口。吸盘力/力矩、接触模式和装配过程不属于本篇运动规划的已验证能力。我没有开展Isaac场景适配，也没有把后续应用设想写成已取得的实验结论。')
    r.h('证据索引')
    r.p('协议：r2_geometry_consistent_qualification_v1.json、r2_rss_fixed_construction_addendum_v1.json。运行：20261015_R2_SPATIAL01、20261015_R2_PLANAR01、20261015_R2_RSS_CONSTRUCTION01。原生R1报告和所有失败不覆盖；本次中文报告只读取保存日志，未新增求解调用。')
    for p in r.doc.paragraphs:
        if list(p._p.iter(qn('m:oMath'))):
            p.paragraph_format.line_spacing=Pt(24)
            p.paragraph_format.line_spacing_rule=WD_LINE_SPACING.AT_LEAST
    args.output.parent.mkdir(parents=True,exist_ok=True);r.doc.save(args.output)
    print(args.output)


if __name__=='__main__':main()
