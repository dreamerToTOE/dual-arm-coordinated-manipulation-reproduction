from pathlib import Path
from copy import deepcopy
import argparse, json, hashlib, csv
from docx import Document
from docx.shared import Inches, Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE

ROOT = Path(__file__).resolve().parents[2]
parser=argparse.ArgumentParser(description='Build the D084 evidence report from saved assets only; no scientific execution.')
parser.add_argument('--output',type=Path,default=ROOT/'reports/p4_d084_s5_final/P4_D084_P4证据整合最终报告_20261025_v5.docx')
args=parser.parse_args()
DOCX = args.output.resolve()
if DOCX.exists(): raise SystemExit('Refusing to overwrite an existing report; choose a new versioned --output path.')
OUT = DOCX.parent
OUT.mkdir(parents=True, exist_ok=True)


INPUT_BINDING = Path(__file__).with_name('s5_report_inputs.json')
for relpath, expected in json.loads(INPUT_BINDING.read_text())['sha256'].items():
    actual = hashlib.sha256((ROOT/relpath).read_bytes()).hexdigest()
    if actual != expected: raise SystemExit(f'Frozen report input changed: {relpath}')

# Evidence sources
s1 = json.loads((ROOT/'reports/p4_d084_abb_ii/abb_ii_metrics.json').read_text())
s2 = json.loads((ROOT/'reports/p4_d084_abb_iii/manifest.json').read_text())
s3 = json.loads((ROOT/'reports/p4_d084_fig3_controlled/summary.json').read_text())
s3meta = json.loads((ROOT/'reports/p4_d084_fig3_controlled/run_metadata.json').read_text())
s3verify = json.loads((ROOT/'reports/p4_d084_fig3_controlled/saved_result_verification.json').read_text())
s4 = json.loads((ROOT/'reports/p4_d084_pcs_difficulty/summary.json').read_text())
s4cost = json.loads((ROOT/'reports/p4_d084_pcs_difficulty/source_cost_binding.json').read_text())
s4audit = json.loads((ROOT/'reports/p4_d084_pcs_difficulty/saved_only_self_audit.json').read_text())
original = Document(ROOT/'reports/P4_D082闭链规划算法与完整虚拟仿真研究报告_20261022.docx')

def sha(path): return hashlib.sha256((ROOT/path).read_bytes()).hexdigest()

def add_original_equations(doc, indices):
    # Reuse the native Word math of the preserved D082 report, without changing notation.
    for idx in indices:
        element = deepcopy(original.paragraphs[idx]._p)
        doc._body._body.insert(len(doc._body._body)-1, element)


def set_cell_shading(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr(); shd = tcPr.find(qn('w:shd'))
    if shd is None: shd = OxmlElement('w:shd'); tcPr.append(shd)
    shd.set(qn('w:fill'), fill)

def set_cell_borders(cell, color='D9D9D9', sz='4'):
    tc = cell._tc; tcPr = tc.get_or_add_tcPr(); borders = tcPr.first_child_found_in('w:tcBorders')
    if borders is None: borders = OxmlElement('w:tcBorders'); tcPr.append(borders)
    for edge in ('top','left','bottom','right','insideH','insideV'):
        tag = 'w:'+edge; el = borders.find(qn(tag))
        if el is None: el=OxmlElement(tag); borders.append(el)
        el.set(qn('w:val'),'single'); el.set(qn('w:sz'),sz); el.set(qn('w:space'),'0'); el.set(qn('w:color'),color)

def set_cell_text(cell, text, bold=False, color='000000', size=9.2):
    cell.text = ''
    p = cell.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run(str(text)); r.bold = bold; r.font.size = Pt(size); r.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_borders(cell)

def add_table(doc, headers, rows, widths=None):
    table=doc.add_table(rows=1, cols=len(headers)); table.alignment=WD_TABLE_ALIGNMENT.CENTER; table.style='Table Grid'
    for i,h in enumerate(headers):
        set_cell_text(table.rows[0].cells[i],h,True,'FFFFFF',9); set_cell_shading(table.rows[0].cells[i],'1F4E79')
    repeat=OxmlElement('w:tblHeader'); table.rows[0]._tr.get_or_add_trPr().append(repeat)
    for ridx,row in enumerate(rows):
        cells=table.add_row().cells
        for i,val in enumerate(row):
            set_cell_text(cells[i],val,False,'000000',8.7); set_cell_shading(cells[i], 'F3F6FA' if ridx%2 else 'FFFFFF')
    table.autofit=False
    if widths:
        for i,w in enumerate(widths): table.columns[i].width=Cm(w)
        for row in table.rows:
            for i,w in enumerate(widths): row.cells[i].width=Cm(w)
    for row in table.rows:
        row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
        for cell in row.cells:
            for p in cell.paragraphs:
                p.paragraph_format.space_after=Pt(3); p.paragraph_format.space_before=Pt(3); p.paragraph_format.line_spacing=1.0
    doc.add_paragraph().paragraph_format.space_after=Pt(2)
    return table

def add_heading(doc, text, level=1):
    p=doc.add_heading(text, level=level); p.paragraph_format.keep_with_next=True; return p

def add_para(doc, text, bold_lead=None):
    p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(5); p.paragraph_format.line_spacing=1.12
    if bold_lead and text.startswith(bold_lead):
        r=p.add_run(bold_lead); r.bold=True; p.add_run(text[len(bold_lead):])
    else: p.add_run(text)
    return p

def add_bullet(doc, text):
    p=doc.add_paragraph(style='List Bullet'); p.paragraph_format.space_after=Pt(2); p.paragraph_format.line_spacing=1.08; p.add_run(text); return p

def add_caption(doc, text):
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after=Pt(6)
    r=p.add_run(text); r.italic=True; r.font.size=Pt(8.5); r.font.color.rgb=RGBColor(80,80,80)

def add_image(doc, path, width=6.35, caption=None):
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before=Pt(4); p.paragraph_format.space_after=Pt(3)
    p.paragraph_format.keep_with_next=True
    p.add_run().add_picture(str(path), width=Inches(width))
    if caption: add_caption(doc, caption)

def add_page_break(doc): doc.add_page_break()

def add_field(paragraph, field):
    run = paragraph.add_run(); fldChar1 = OxmlElement('w:fldChar'); fldChar1.set(qn('w:fldCharType'),'begin'); instr = OxmlElement('w:instrText'); instr.set(qn('xml:space'),'preserve'); instr.text=field; fldChar2=OxmlElement('w:fldChar'); fldChar2.set(qn('w:fldCharType'),'end'); run._r.extend([fldChar1,instr,fldChar2])

# Document setup
doc=Document(); sec=doc.sections[0]; sec.top_margin=Cm(1.8); sec.bottom_margin=Cm(1.7); sec.left_margin=Cm(1.9); sec.right_margin=Cm(1.9)
styles=doc.styles
normal=styles['Normal']; normal.font.name='Noto Sans CJK SC'; normal._element.rPr.rFonts.set(qn('w:eastAsia'),'Noto Sans CJK SC'); normal.font.size=Pt(11); normal.paragraph_format.space_after=Pt(5); normal.paragraph_format.line_spacing=1.12
for name,size,color in [('Title',22,'000000'),('Heading 1',15,'000000'),('Heading 2',12,'000000'),('Heading 3',11,'000000')]:
    st=styles[name]; st.font.name='Noto Sans CJK SC'; st._element.rPr.rFonts.set(qn('w:eastAsia'),'Noto Sans CJK SC'); st.font.size=Pt(size); st.font.bold=True; st.font.color.rgb=RGBColor.from_string(color)
    st.paragraph_format.space_before=Pt(10 if name!='Title' else 0); st.paragraph_format.space_after=Pt(5)
    for borders in list(st.element.xpath('./w:pPr/w:pBdr')): borders.getparent().remove(borders)
# Footer
footer=sec.footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.CENTER; footer.add_run('P4 D084 证据整合报告  |  2026-10-25  |  '); add_field(footer,'PAGE')

# Cover
p=doc.add_paragraph(style='Title'); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run('P4 D084 证据整合最终报告')
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=p.add_run('闭链规划 受控分划实验与双 ABB 空间离散证据'); r.font.size=Pt(13); r.font.color.rgb=RGBColor(0,0,0)
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before=Pt(18); r=p.add_run('版本：2026-10-25  |  分支：p4-full-reproduction-pcs-study'); r.font.size=Pt(10)
add_para(doc,'我研究了全驱动转动关节闭链的采样规划，并在重建 20R 机构上完成了带闭链、限位和碰撞检查的端到端虚拟任务。当前保存数据没有显示 PCS 相对 NR 或兼容固定分划的优势。双 ABB II/III 的空间路径及四次角色切换得到离散证据支持，仍不构成原生 TCP 等价或连续安全认证。本文整合 D082 主实验与 D084 新增整理结果；全文复现、动力学和硬件执行仍未完成。')
add_table(doc,['阶段','交付状态','主要证据'],[
 ['S0 安全恢复','完成','本地未提交 D083 资产审计、SHA、进程与预算台账'],
 ['S1 ABB II','完成','4 条保存路径、3D 图、证书绑定与离散限制'],
 ['S2 ABB III','完成','649 状态、14 边、4 次主动臂角色切换'],
 ['S3 受控 Fig.3','完成','96 端点、180 对、2880 候选、900 有限策略调用'],
 ['S4 困难诊断','完成','36 注册调用、35 续跑、1 UNKNOWN、29 独立成功'],
 ['S5 最终报告','本文件','等待用户审核；不自动进入其他研究范围'],
], widths=[3.0,2.3,10.0])
add_page_break(doc)

# Executive summary
add_heading(doc,'一 结论与证据等级',1)
add_para(doc,'核心结论：在当前冻结的重建平面 20R 条件中，PCS 没有显示相对 NR 或兼容固定分划的优势；受控 Fig.3 只描述一个相关、有限的保存端点总体。ABB II 和 ABB III 的空间材料证明了保存离散路径和机制图示，但不证明原生 TCP 等价或连续扫掠安全。S4 的三个条件均未建立真正窄通道或困难障碍，因而不能把有限诊断升级为普遍困难场景结论。')
add_para(doc,'证据层级：D082 主统计和独立离散验证属于已交付的重建证据；S1–S2 是保存路径可视化；S3 是冻结共享总体的受控局部连接实验；S4 是一次有预算界的确定性端点诊断。本项目对原文 Table I 与六环境正式统计的复现仍为 PARTIAL；本项目 D082 重建 RSS 随机完整规划为 0/10、FAILED。这里不对原作者实验作失败判定。')
add_heading(doc,'二 原论文问题与重建模型',1)
add_para(doc,'Sintov、Borum 与 Bretl 的比较研究讨论如何在全驱动转动关节闭链的低维约束流形上采样、投影并连接构型 [1]。所有关节均有执行器；“被动链”只是解析求解时的角色。闭链可行集与避障有效集分开定义，规划路径必须连接指定端点并始终位于有效集内。以下沿用 D082 可编辑 Word 公式，其编号与 D082 一致。')
add_original_equations(doc,[12,13])
add_heading(doc,'平面闭链与 NR 投影',2)
add_para(doc,'重建机构含 19 根等长移动杆和第 20 个零长度转动法兰，固定地杆长 1 m。累积角决定末端位置，总关节角决定闭链姿态。平面约束为三维；空间单闭环通常为六维。')
add_original_equations(doc,[15,16,17,19,20,23])
add_para(doc,'NR 在每次更新中线性化闭链残差，并用 SVD 伪逆求最小范数校正。阈值为 max(10⁻¹², 10⁻¹⁰σ_max)，最多 100 次更新；位置与角度分别按 10⁻⁸ m 和 10⁻⁸ rad 验收。实现无阻尼、无线搜索、无随机重启。角度包裹不是任意工业关节限位截断，也不能保证投影总能收敛。')
add_heading(doc,'APC 解析闭合与 PCS 分划切换',2)
add_para(doc,'APC 固定主动关节，用连续非冗余 3R 被动链的解析 IK 闭合剩余变量。去掉前后链变换后，先求腕点，再由余弦定理得到两条肘部分支；不可达输入按失败保留，不通过放宽工作空间获得成功。')
add_original_equations(doc,[28,30,31,32,33])
add_para(doc,'PCS 对局部边寻找两端共有的分划及非奇异 IK 分支，再在该分划内递归连接。树上不同边可以采用不同分划；只有按真实父子关系或最终路径顺序变化，才记为相应切换。主规划用 p0 与 p16；S3 暴露 17 个连续正长 3R 分划，使用前缀 m=1/2/3/7/16。末端零长度法兰使这些图册不能直接冒充原文 m=20。')
add_para(doc,'原文在独立、等测度等假设下讨论多被动链可见域的上界。下式是带假设的理论表达，不是本次有限样本成功率的拟合或下界；障碍、限位、图册依赖与 APC 不可达性必须另行考虑。')
add_original_equations(doc,[38,39])
add_heading(doc,'RSS 奇异目标与 RBS 局部连接',2)
add_para(doc,'RSS 用概率 γ=0.2 请求人工奇异目标。平面 3R 被动雅可比的行列式与 ℓ₁ℓ₂ sin q₂ 成比例，肘角为 0 或 π 时两肘支可相接；被动雅可比降秩不等于整个闭链雅可比降秩。随机目标还须闭合、限位合法、无碰撞并可由树到达。历史手工 rank-2 夹具只证明局部机制，不能代替随机全局规划成功。')
add_original_equations(doc,[45])
add_para(doc,'RBS 先验证端点，将关节中点用所选投影器送回流形，对每个新中点检查闭链、限位和碰撞，再递归处理左右两半。PCS 保持所选分划及共同分支，NR 使用全链投影，RSS 使用固定被动链。任何投影或有效性失败都拒绝该试探边。')
add_original_equations(doc,[49,50])
add_para(doc,'主任务与 S4 在叶距不超过 0.05 rad 时终止，深度上限 12、每连接 256 次投影；S3 按冻结协议使用深度 16 和 512 次共享预算。有限叶距及离散加密均不是连续 CCD 证明。模型和软件身份分为 ORIGINAL 机制、RECONSTRUCTED 模型与 ENGINEERING 资产整理；具体配置和 SHA 见参数表及附录。')
add_heading(doc,'三 D082 主统计与负结果',1)
add_para(doc,'D082 共有 50 次真实规划调用，E0/E1 两类障碍、五个种子和五种方法各按冻结预算运行。记录为 26 条 SUCCESS、24 条 TIME_LIMIT，13 条不同关节路径在两个环境复用。NR 每个环境 5/5 成功，兼容 fixed-0 每个环境 5/5；PCS full 为每环境 3/5，fixed-16 和 RSS 均为 0/5。D082 的全 PCS 最终解分划切换为 0，RSS 的 2304 次奇异构造全部在 IK_UNREACHABLE 阶段拒绝。')
add_table(doc,['方法','E0 成功/5','E1 成功/5','解释'],[
 ['RSS','0/5','0/5','72 次奇异请求、2304 次构造均未形成合格随机目标'],
 ['PCS full','3/5','3/5','有限重建结果，未显示优于兼容固定或 NR'],
 ['PCS fixed-0','5/5','5/5','端点兼容的固定分划基线'],
 ['PCS fixed-16','0/5','0/5','结构负对照，端点分支不兼容'],
 ['NR','5/5','5/5','重建平面条件下成功率更高'],
], widths=[3,2.5,2.5,7.3])
add_para(doc,'这些计数不是原论文 22500 次正式调用，也不是原文六环境结论。D080 取消的 18000 次计划不重新解释为已完成样本。')


# Fig3
add_heading(doc,'四 D084 S3 受控 Fig3 实验',1).paragraph_format.page_break_before=True
add_para(doc,'S3 从 13 条既有成功轨迹得到 96 个经独立检查合格的保存端点，再固定抽取 180 个无序非自对。端点来自相关轨迹，不能解释为可行流形上的独立同分布样本。所有 m 共享同一端点、配对、方向和距离分箱，避免把不同 m 的构型总体差异误写成因果效果。候选并集是冻结有限 RBS 规则下的存在性估计；有限策略使用全新共享预算，两者严格分开。')
add_table(doc,['m','候选并集成功/180','有限策略成功/180','候选均值/ms','策略均值/ms'],[
 ['1','176','176','14.966','14.977'],['2','180','180','20.793','15.255'],['3','180','180','37.002','15.255'],['7','180','180','108.912','15.355'],['16','180','180','209.505','15.402'],
], widths=[1.5,3.2,3.2,3.0,3.0])
add_image(doc,ROOT/'reports/p4_d084_fig3_controlled/plots_v2/fig3_controlled_visibility_and_local_cost_by_m.png',6.4,'图 1 受控共享总体的可见率与局部检查成本；成本不是原文 Bi-RRT 规划时间。')
add_image(doc,ROOT/'reports/p4_d084_fig3_controlled/plots_v2/fig3_controlled_visibility_by_distance.png',6.4,'图 2 按冻结距离分箱的描述性结果；分母由同一 180 对决定。')
add_para(doc,'实验外层保守成本为 86.29 秒 wall / 90.71 秒进程 CPU，低于 300/300 秒 Stage B 帽。S3 的精确描述率不附总体置信区间，也不能外推原论文均匀可行流形、m=20 或复杂障碍性能。')
add_para(doc,'原始记录共 2880 次候选和 900 次有限策略评估，分别有 2358 和 896 条成功调用。候选并集按 m 汇总为 176/180、180/180、180/180、180/180、180/180；有限策略的四条失败均为 m=1。3254 条接受局部调用记录具有保存的独立资格，涉及 97577 个点引用。失败的首次 prepare 与成功输出均保留；D084 S3 只打包已有结果，没有重跑科学调用。')

# Spatial
add_heading(doc,'五 ABB II 空间保存路径',1)
add_para(doc,'S1 只读取已经合格的保存节点，打包 RRT/PCS 177、CBiRRT/PCS 413、RRT/NR 393、CBiRRT/NR 689 四条 Environment II 路径。图中完整 TCP 轨迹和起点、中点、终点构型来自保存记录；几何表面沿用 aligned CAD-frame 约定。证书中的闭链、限位和离散碰撞结果未被重新升级。')
add_image(doc,ROOT/'reports/p4_d084_abb_ii/fig_S1_abb_ii_four_saved_paths.png',6.4,'图 3 Environment II 四条保存离散路径的起点、中点和终点。')
add_table(doc,['路径','节点/叶','最大叶距/rad','最大位置闭链/mm','连续 CCD'],[
 ['RRT/PCS','177/176','0.0494905','3.41e−13','未证'],['CBiRRT/PCS','413/412','0.0498729','4.55e−13','未证'],['RRT/NR','393/392','0.0478001','9.79e−06','未证'],['CBiRRT/NR','689/688','0.0495190','9.80e−06','未证'],
], widths=[3.0,2.0,3.0,3.2,3.0])
add_para(doc,'S1 不把保存路径图示写成 PCS 优势、native TCP 等价、实体杆连续安全或动态可执行性证书。每条路径的 certificate SHA、execution SHA 和输入样本 SHA 位于 `reports/p4_d084_abb_ii/manifest.json`。')

add_heading(doc,'六 ABB III 角色切换',1)
add_para(doc,'S2 绑定 649 个保存 q 状态、15 个里程点和 14 条接受边。主动臂数组为 [0,0,0,1,0,0,0,1,1,0,0,0,0,0]。四次转移发生在 source node 179、248、392、492，分别为 A1→A2、A2→A1、A1→A2、A2→A1。它们描述接受边的主动/被动角色归属，不是连续物理换手。')
add_image(doc,ROOT/'reports/p4_d084_abb_iii/fig_S2_abb_iii_role_timeline.png',6.4,'图 4 ABB III 接受边角色时间线与四个切换节点。')
add_image(doc,ROOT/'reports/p4_d084_abb_iii/fig_A3_environment_III_four_switch_frames.png',6.4,'图 5 四次切换前后保存构型；仅作离散机制图示。')
add_para(doc,'649/649 aligned TCP 通过而 native TCP 0/649；最大 native 平移差 0.003570536971253091 mm，旋转矩阵差 7.347665530579861e−06。source exact 标记不是独立证书；无连续 CCD、最小净距、遗漏碰撞对或硬件安全证明。')
add_page_break(doc)

# S4
add_heading(doc,'七 S4 有预算界的 PCS 困难诊断',1)
add_para(doc,'S4 使用预注册的 C1–C3 三个确定性端点条件、种子 7/23/41 和 PCS、兼容 fixed、NR、RSS 四种方法。第一条 C1/PCS/seed7 调用实际运行但 JSON 序列化失败，只有 SUCCESS 字面前缀，故保持 UNKNOWN；续跑只执行剩余 35 条。')
add_table(doc,['方法','真实调用','完整可评价','独立成功','超时/UNKNOWN'],[
 ['PCS','9','8','8','1 UNKNOWN'],['兼容 fixed-0','9','9','9','0'],['NR','9','9','9','0'],['RSS','9','9','3','6 超时'],
], widths=[4.0,2.2,2.7,2.5,4.0])
add_image(doc,ROOT/'reports/p4_d084_pcs_difficulty/difficult_diagnostic.png',6.4,'图 6 三条件有限诊断的成功分母与候选 clearance；不是无偏困难基准。')
add_para(doc,'合计外部成本为 17.00 秒 wall / 19.84 秒进程 CPU，低于 180/240 秒 Stage C 帽。29 个完整成功解通过独立保存路径审查，树和最终解相邻边分划切换均为 0。RSS 的 1696 个奇异构造全部 IK_UNREACHABLE，三个 RSS 成功均为普通同分支 p16 边，不能当作奇异桥成功。')

# Parameters and limitations
add_heading(doc,'八 参数与未覆盖范围',1)
add_table(doc,['项目','原论文或待核范围','本项目实际证据'],[
 ['统计规模','论文含六环境；Table I 为 II–VI 五空间环境×3规划器×3方法×500=22500','D082 50 次 pilot；D084 S3/S4 为冻结小样本诊断'],
 ['Fig.3','平面 20R、m=1/2/3/7/20、原文总体','重建 20R proxy，m 最大 16，96 端点/180 对'],
 ['碰撞与几何','原作者模型同一性待核；连续安全属于本项目尚未认证范围','离散 source-policy / aligned 证据；无 CCD'],
 ['RSS','原文随机奇异采样统计','历史 0/10 和 2304 IK_UNREACHABLE 保留；新 S4 仅 3/9 普通同分支成功'],
 ['空间 ABB','原生 TCP 与空间统计','保存路径图示；III native mismatch 649/649'],
 ['硬件/动力学','后续应用需要的执行与动态安全；不冒称原论文已验证这些条件','本项目未执行'],
], widths=[3.0,5.4,7.0])
add_para(doc,'因此，本报告支持“有限重建端到端虚拟任务与离散证据已交付”，不支持“原论文全文已复现”“PCS 在复杂环境普遍优越”或“路径可直接用于真实机器人执行”。')
add_heading(doc,'冻结的重建数值参数',2)
add_table(doc,['范围','参数','值及来源约束'],[
 ['D082 与 S4','机构与碰撞','19 杆长 sin(π/38) m；地杆 1 m；20R 含零长法兰；关节界 [−π,π]；胶囊半径 0.002 m，仅排除真实相邻杆对'],
 ['D082','端点与障碍','a=π/19；起点 q0=q19=π/2−a/2，其余 −a；目标仅 q16..18=(−2a,a,−2a)。E0 圆 (0.45,0.72),r=.045 m；E1 增加 (.50,.18),r=.065 与 (.72,.60),r=.050 m'],
 ['D082 与 S4','投影与 RBS','闭链位置/角度 10⁻⁸ m/rad；NR 100 更新；SVD max(10⁻¹²,10⁻¹⁰σmax)；RBS 0.05 rad/深度12/256投影；扩展 0.6 rad；RSS γ=.2/最多32次构造'],
 ['S3','冻结抽样','E0_sparse；96 端点 seed2026102201；180 无序非自对 seed2026102202；方向按索引递增；m 前缀 0..15，全部可用分划0..16'],
 ['S3','有限局部预算','每候选独立512中点投影；实际策略全部升序图册共享512投影，首成功停止；每连接2.0 s；深度16；叶距.05 rad；分支 ε=10⁻¹⁰，余弦边界舍入10⁻¹²'],
 ['S4','诊断调度','E1_clutter；三个保存中点条件，种子7/23/41；36调用上限，1 s/调用；300树节点、8次connect；兼容fixed从[0,16]首个可用选取，实际均0'],
 ['D082 与 S4','随机与统计','D082 五种子7/23/41/59/73；S4仅7/23/41；同种子不保证算法消耗同一采样向量；没有每方法最优参数调优'],
],widths=[2.2,3.1,10.1])
add_para(doc,'上述几何、端点和阈值属于已冻结重建实验，未声明等同原作者输入。D082 每组 n=5，成功 5/5 的 95% Wilson 区间约 [0.566,1]，0/5 仍约 [0,0.434]。S3 为冻结有限总体的描述性结果，S4 只有三个相关条件和三个种子，PCS 还存在一条 UNKNOWN；均不宜形成总体方法排序或显著性结论。超时不等于无解。')

add_heading(doc,'九 可复现入口与 Git 证据',1)
for t in [
 'S0 恢复提交：a660e14189ded7878685d9388c6241b65166ffc7；报告修订：4b86552829177799f8bc22ca6a6cf113fe8943b0。',
 'S1：613611d89cbc1e39e4ffbb9bd1e73f40c45285ff；S2：882ff171f819d9dd2735edf04ac8f783289f3f7f；S3：0ed66ec30d22a85f6a03088b3bfbfd4b21a50291；S4：bec0f9450a78a115c439368f0ca5ecfdb09b0cfb。',
 '阶段说明与哈希位于 reports/p4_d084_abb_ii、reports/p4_d084_abb_iii、reports/p4_d084_fig3_controlled 和 reports/p4_d084_pcs_difficulty 的 README.md / manifest.json。短交接分别为 reports/P4_D084_S1_POST.md 至 P4_D084_S4_POST.md。',
 'D082 原始 211×257 coverage CSV 未被改写；D082 原 Word/PDF、D083 草稿和未提交 raw archives 均保留原位置。',
]: add_bullet(doc,t)
add_para(doc,'本报告构建脚本 scripts/p4_d084/build_final_report.py 只读取冻结文件和 D082 原生 Word 公式。可用 Codex 内置 Python 指定全新 --output 路径重新排版；PDF 由已有隔离 LibreOffice 渲染器导出。具体环境、校验命令和最终文件 SHA 见 reports/P4_D084_S5_FINAL_HANDOFF.md。该复现入口不重新执行任何规划、IK、RBS、碰撞或资格实验。')

add_heading(doc,'十 面向双机械臂课题的适用性差距与待研究问题',1)
add_para(doc,'刚性共同抓持时，可以用两侧运动学推算的物体位姿相等定义闭链约束，再由投影和局部连接检查几何可行性。这只是 D082 已讨论的概念接口；当前没有 FR3 运动学、抓持标定或实际碰撞体的独立资格。NR 不要求闭式 IK，概念上较易迁移，但收敛、关节限位和奇异邻域仍需单独验证。')
add_para(doc,'冗余 7R 机械臂不能直接套用原文非冗余 6R 被动臂的有限解析分支。PCS 需要可证明的解析子问题与连续分支覆盖；RSS 需要真实降秩集合及可闭合、合法、无碰撞的随机目标。吸盘柔性、滑移、内力、动力学、时间参数化和执行误差超出本轮几何证据。不能通过扩大闭链容差或使用真值绕过这些差距。当前不实现 FR3、Isaac 或硬件迁移，也不据此推出应用收益。')

add_heading(doc,'附录 A D084 有限覆盖增量',1).paragraph_format.page_break_before=True
add_para(doc,'下表只说明本轮证据与原主义务的关联，不修改 D082 的 211 行历史映射或状态单元格。完整历史文件为 reports/P4_D082_ORIGINAL_VS_APPROVED_VS_ACTUAL.csv（211×257），SHA256 '+sha('reports/P4_D082_ORIGINAL_VS_APPROVED_VS_ACTUAL.csv')+'。其 .md 是 30 条主义务摘要；D084 不把任何增量换算为原论文 FULL。')
add_table(doc,['本轮','对应原主义务','新增证据与仍未覆盖'],[
 ['S1','P4-T11 T16 T17 T26','II 四条保存路径的真实图示和来源绑定；没有新规划、完整方法矩阵或连续安全结论'],
 ['S2','P4-T12 T08 T16 T27','III 649 状态、14边、4次角色切换；native mismatch 649/649 保留，角色切换不等于平面分划统计'],
 ['S3','P4-T19 T06 T21 T05 T26','修正共享总体混淆的有限代理实验；m20、匹配 NR、原总体、原 Bi-RRT 成本和约5万点对仍缺失'],
 ['S4','P4-T28 T23 T08 T07 T09 T24','36注册/35完整/1UNKNOWN；没有建立真正困难条件、最终解切换、有效奇异桥或 PCS 优势'],
 ['范围保留','P4-T18 T13 T14 T15 T30','Table I 正式调用0/22500，历史18k取消；IV/V未补完整、VI只读，原投影分布主张未验证'],
 ['S5','P4-T25 T29 及文档交付','整合可追溯报告与概念性应用差距；不新增科学覆盖或工程执行证据'],
],widths=[1.6,4.2,9.6])

add_heading(doc,'附录 B 图与保存路径来源',1)
figures=[
 (1,'reports/p4_d084_fig3_controlled/plots_v2/fig3_controlled_visibility_and_local_cost_by_m.png','S3 summary_by_m.csv 与 pair_by_m_outcomes.csv'),
 (2,'reports/p4_d084_fig3_controlled/plots_v2/fig3_controlled_visibility_by_distance.png','S3 visibility_by_distance.csv；冻结180对'),
 (3,'reports/p4_d084_abb_ii/fig_S1_abb_ii_four_saved_paths.png','S1 四条来源证书与保存 node JSON'),
 (4,'reports/p4_d084_abb_iii/fig_S2_abb_iii_role_timeline.png','S2 role_timeline.csv 和649节点绑定'),
 (5,'reports/p4_d084_abb_iii/fig_A3_environment_III_four_switch_frames.png','S2 manifest 的四切换前后保存状态'),
 (6,'reports/p4_d084_pcs_difficulty/difficult_diagnostic.png','S4 method_summary.csv 与候选诊断'),
]
for num,path,source in figures:
    p=add_para(doc,f'图 {num}  {path}\nSHA256 {sha(path)}\n数据来源：{source}。')
    for r in p.runs:r.font.size=Pt(9)
add_para(doc,'图 3–5 均来自既有保存状态；本轮没有图像插值、轨迹再求解或碰撞重验。图 1–2 和图 6 沿用已冻结图件，不重新绘制。完整输入、输出和约束语义记录在对应阶段 manifest.json。')
for path in json.loads((ROOT/'reports/p4_d084_abb_ii/manifest.json').read_text())['paths']:
    cert=path['certificate']
    p=add_para(doc,f"II {path['label']}  {path['node_count']} 节点\n保存目录：{path['run']}\n证书：{cert}\nSHA256 {sha(cert)}")
    for r in p.runs:r.font.size=Pt(9)
cert=s2['certificate']
p=add_para(doc,f'III RRT/PCS 649 节点\n证书：{cert}\nSHA256 {sha(cert)}\n3253 项保存引用位于该 III 证书的 retained_sha256；原执行与协议 SHA 同在该证书。状态为 PASS_649_DISCRETE_SOURCE_POLICY_PATH，independent_path_qualified=false；source-policy/aligned 离散限制保留。')
for r in p.runs:r.font.size=Pt(9)

add_heading(doc,'附录 C 数值与文献来源',1)
sources=[
 'reports/P4_D082_DELIVERY_README_20261022.md',
 'reports/P4_D082闭链规划算法与完整虚拟仿真研究报告_20261022.docx',
 'reports/P4_D082闭链规划算法与完整虚拟仿真研究报告_20261022.pdf',
 'reports/p4_d084_abb_ii/manifest.json', 'reports/p4_d084_abb_iii/manifest.json',
 'reports/p4_d084_fig3_controlled/manifest.json','reports/p4_d084_pcs_difficulty/manifest.json',
 'configs/p4_d083/fig3_controlled_v1.json','configs/p4_d083/difficult_v1.json',
]
for path in sources:
    p=add_para(doc,path+'\nSHA256 '+sha(path))
    for r in p.runs:r.font.size=Pt(9)
add_para(doc,'S3 原结果位于 results/20261022_D083_fig3_controlled02；raw_pairs 的 180 个 gzip（45389043 bytes）仍仅本地保存，逐文件 SHA 在该结果及交付包的 artifact_sha256.json。S4 原结果位于 results/20261022_D083_difficult02；raw_planner 的 35 个 gzip（8269303 bytes）仍仅本地保存，其逐文件 SHA 同样保留。失败目录 controlled01 与 difficult01、原 D083 草稿和所有渲染尝试均未删除；完整仅本地清单见 S5 交接。')
add_para(doc,'[1] Avishai Sintov, Andy Borum, Timothy Bretl. Motion Planning of Fully Actuated Closed Kinematic Chains With Revolute Joints: A Comparative Analysis. IEEE Robotics and Automation Letters, 3(4): 2886–2893, 2018. DOI 10.1109/LRA.2018.2846806。原作者八页稿 SHA256 193710e5a9d4284cf62b30ff189a1df038a9d5c5ce3f1545525dc2f78bea08c6；本文引用身份沿用 D082，不分发私有论文副本。')
add_para(doc,'[2] 作者 CKCplanning 参考源码快照 909bd72fff778aca8682c96f9eb86ae4eb75557b；本项目仓库 dreamerToTOE/dual-arm-coordinated-manipulation-reproduction，分支 p4-full-reproduction-pcs-study。作者依赖与本项目 Python/GEOS 重建环境不同，绝对时间不拼接计算跨架构加速比。')

# Metadata and save
doc.core_properties.title='P4 D084 证据整合最终报告'
doc.core_properties.subject='闭链规划、受控分划实验与双 ABB 空间离散证据'
doc.core_properties.author='P4 reproduction project'
doc.save(DOCX)
print(DOCX)
