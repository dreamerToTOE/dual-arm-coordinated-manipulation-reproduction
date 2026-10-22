#!/usr/bin/env python3
"""Build a new Chinese report from preserved evidence; never overwrite history.

Run with the bundled artifact Python. Equations remain editable Word math.
"""
from pathlib import Path
import argparse
import json
import math
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[2]
OUT=ROOT/'reports/P4_D082闭链规划算法与完整虚拟仿真研究报告_20261022.docx'
ASSETS=ROOT/'reports/p4_d082_research_assets_20261022'
DOC=Document()

def p(text):
    return DOC.add_paragraph(text)

def h(text,level=1):
    return DOC.add_heading(text,level)

def expr(parent, value):
    if isinstance(value, (list, tuple)):
        for part in value: expr(parent, part)
    elif isinstance(value, str):
        r=OxmlElement('m:r');t=OxmlElement('m:t');t.text=value
        r.append(t);parent.append(r)
    else: parent.append(value)

def sub(base, index):
    node=OxmlElement('m:sSub')
    for tag, value in (('m:e',base),('m:sub',index)):
        child=OxmlElement(tag);expr(child,value);node.append(child)
    return node

def sup(base, exponent):
    node=OxmlElement('m:sSup')
    for tag, value in (('m:e',base),('m:sup',exponent)):
        child=OxmlElement(tag);expr(child,value);node.append(child)
    return node

def frac(numerator, denominator):
    node=OxmlElement('m:f')
    for tag, value in (('m:num',numerator),('m:den',denominator)):
        child=OxmlElement(tag);expr(child,value);node.append(child)
    return node

def large(symbol, lower, upper, body):
    node=OxmlElement('m:nary');pr=OxmlElement('m:naryPr')
    char=OxmlElement('m:chr');char.set(qn('m:val'),symbol);pr.append(char)
    lim=OxmlElement('m:limLoc');lim.set(qn('m:val'),'subSup');pr.append(lim);node.append(pr)
    for tag, value in (('m:sub',lower),('m:sup',upper),('m:e',body)):
        child=OxmlElement(tag);expr(child,value);node.append(child)
    return node

def eq(value, label=None):
    paragraph=DOC.add_paragraph();paragraph.alignment=WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before=Pt(2)
    paragraph.paragraph_format.space_after=Pt(4)
    om=OxmlElement('m:oMath');expr(om,value);paragraph._p.append(om)
    if label: paragraph.add_run('   ('+str(label)+')')
    return paragraph

def page(title):
    DOC.add_page_break();h(title)

def figure(name,caption,width=15.8):
    paragraph=DOC.add_paragraph();paragraph.alignment=WD_ALIGN_PARAGRAPH.CENTER
    paragraph.add_run().add_picture(str(ASSETS/name),width=Cm(width))
    paragraph.paragraph_format.keep_with_next=True
    paragraph=DOC.add_paragraph(caption);paragraph.alignment=WD_ALIGN_PARAGRAPH.CENTER
    for r in paragraph.runs:r.font.size=Pt(9)

def table(headers,rows,widths=None):
    t=DOC.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    for cell,text in zip(t.rows[0].cells,headers):cell.text=str(text)
    for row in rows:
        for cell,text in zip(t.add_row().cells,row):cell.text=str(text)
    if widths:
        widths=[w*16.8/sum(widths) for w in widths]
        for column,w in zip(t.columns,widths):column.width=Cm(w)
        for row in t.rows:
            for cell,w in zip(row.cells,widths):cell.width=Cm(w)
    pr=t._tbl.tblPr;b=OxmlElement('w:tblBorders')
    for side in ('top','left','bottom','right','insideH','insideV'):
        e=OxmlElement('w:'+side);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'5');e.set(qn('w:color'),'D9D9D9');b.append(e)
    pr.append(b)
    repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
    for i,row in enumerate(t.rows):
        row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
        for cell in row.cells:
            cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            margin=OxmlElement('w:tcMar')
            for side in ('top','left','bottom','right'):
                e=OxmlElement('w:'+side);e.set(qn('w:w'),'75');e.set(qn('w:type'),'dxa');margin.append(e)
            cell._tc.get_or_add_tcPr().append(margin)
            if i==0 or i%2==0:
                shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E5EBEF' if i==0 else 'F6F7F8');cell._tc.get_or_add_tcPr().append(shade)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after=Pt(2);paragraph.paragraph_format.line_spacing=1.12
                for run in paragraph.runs:run.font.size=Pt(9);run.bold=i==0
    DOC.add_paragraph().paragraph_format.space_after=Pt(1)
    return t

def setup():
    section=DOC.sections[0];section.page_height=Cm(29.7);section.page_width=Cm(21)
    section.top_margin=Cm(1.8);section.bottom_margin=Cm(1.8);section.left_margin=Cm(2.1);section.right_margin=Cm(2.1)
    for name in ('Normal','Title','Subtitle','Heading 1','Heading 2','Caption'):
        s=DOC.styles[name];s.font.name='Noto Sans CJK SC';s.font.color.rgb=RGBColor(0,0,0)
        s.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Noto Sans CJK SC')
    normal=DOC.styles['Normal'];normal.font.size=Pt(10.5);normal.paragraph_format.line_spacing=1.12;normal.paragraph_format.space_after=Pt(5)
    for name,size in (('Title',21),('Heading 1',15),('Heading 2',12)):
        s=DOC.styles[name];s.font.size=Pt(size);s.paragraph_format.space_before=Pt(8);s.paragraph_format.space_after=Pt(8)
    footer=section.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
    r=footer.add_run();field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');r._r.addnext(field)
    for border in list(DOC.styles.element.iter(qn('w:pBdr'))):
        border.getparent().remove(border)
    DOC.core_properties.title='P4闭链规划算法与完整虚拟仿真研究报告'
    DOC.core_properties.subject='RSS PCS NR RBS及独立离散路径验证'
    DOC.core_properties.author=''

def main():
    global OUT
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--output',type=Path,default=OUT)
    OUT=parser.parse_args().output
    if OUT.exists():raise FileExistsError(OUT)
    # The numerical validation is a required input, not a presumed result.
    audit_path=ROOT/'results/20261022_D082_independent_planar01/summary.json'
    audit=json.loads(audit_path.read_text())
    if audit['status'] != 'PASS_DISCRETE_WITH_STRATEGY_MIDPOINT_REFINEMENT' or audit['independent_qualified_count'] != 26 or not all(audit[k] for k in ('oracle_fixtures_passed','original_manifest_all_match','input_files_unchanged')):
        raise RuntimeError('independent validation did not pass all 26 trajectories')
    setup()
    DOC.add_paragraph('P4闭链规划算法与完整虚拟仿真研究报告','Title')
    p('RSS、PCS、NR 与 RBS 的数学复现和独立数值验证\n2026年10月22日')
    h('研究结果概述')
    p('我研究了全驱动转动关节闭链的采样规划问题，并在带关节限位、自碰撞和圆形障碍的20R重建机构上完成了从起点到终点的规划。程序实际执行解析逆运动学、牛顿投影、递归二分局部连接和双向树搜索，保存完整轨迹及失败记录；独立程序使用另一套几何计算复算成功轨迹。这里的完整性指一套端到端虚拟任务，原论文全部空间环境和统计规模仍只得到部分复现。')
    p('五个预定种子在两类障碍中共形成50次规划。每个环境中，NR成功5/5，PCS全分划成功3/5，兼容的固定分划0成功5/5，固定分划16和RSS均为0/5。当前任务没有显示PCS相对NR或兼容固定分划的优势。E0和E1的成功判定相同，增加障碍未形成足以区分方法的有效困难条件。')
    p('空间环境中复用了已保存的合格路径：环境II有四条早期路径及一条SBL路径，环境III保存649点路径。III最终解中的主动臂角色发生4次切换，但所有649点仍存在原生TCP与重建碰撞位姿差异。空间证据只支持明确标注的aligned离散碰撞策略，不能据此认证原作者原生几何或连续扫掠安全。')
    h('原论文与实现身份',2)
    p('Avishai Sintov、Andy Borum、Timothy Bretl，Motion Planning of Fully Actuated Closed Kinematic Chains With Revolute Joints A Comparative Analysis，IEEE Robotics and Automation Letters，3卷4期，2886–2893，2018。DOI 10.1109/LRA.2018.2846806。原文阅读副本为8页作者稿，SHA256为193710e5a9d4284cf62b30ff189a1df038a9d5c5ce3f1545525dc2f78bea08c6。')
    p('本文将原论文定义与机制标为ORIGINAL，将同类20R机构、胶囊碰撞和有限预算的独立重建标为RECONSTRUCTED，将软件修补标为ENGINEERING。单元奇异桥、统计试验和未来FR3设想分别标明证据层级，不把预设夹具视为随机规划成功。')

    page('1 闭链约束与牛顿投影')
    p('原论文式1–2将规划问题分为运动学闭合与有效性两层。设关节变量q属于n维构型空间，C(q)为闭链残差。平面单环有3个独立约束，空间单环通常有6个。关节限位、机器人自碰撞及外部障碍共同决定自由空间。')
    eq([sub('C','cl'),' = {q ∈ ',sup('R','n'),' : C(q) = 0},   ',sub('C','o'),' = ',sub('C','cl'),' ∩ ',sub('C','free')],1)
    eq(['α : [0,1] → ',sub('C','o'),',   α(0) = ',sub('q','s'),',   α(1) = ',sub('q','g')],2)
    p('对平面机构，令θ_i为前i个相对关节角的累加，ℓ_i为连杆长度。末端位置是每个连杆位移之和，末端角由全部转动关节决定。本模型19个移动杆及第20个零长度转动法兰共同构成闭链；固定地杆连接两个基座。')
    eq([sub('θ','i'),' = ',large('∑','j=1','i',sub('q','j'))],3)
    eq(['x = ',large('∑','i=1','19',[sub('ℓ','i'),' cos ',sub('θ','i')]),',   y = ',large('∑','i=1','19',[sub('ℓ','i'),' sin ',sub('θ','i')])])
    eq(['C(q) = ',sup(['[x − 1, y, wrap(',large('∑','j=1','20',sub('q','j')),')]'],'T')],4)
    p('对第j个关节求偏导，只需累加其后连杆的切向位移。零长度末端关节对位置导数为零，对角度导数仍为1。解析雅可比与独立中心差分的既有四样本最大绝对差为1.708×10⁻¹⁰。')
    eq([sub('J','x,j'),' = −',large('∑','i=j','19',[sub('ℓ','i'),' sin ',sub('θ','i')]),',   ',sub('J','y,j'),' = ',large('∑','i=j','19',[sub('ℓ','i'),' cos ',sub('θ','i')])],5)
    eq([sub('J','θ,j'),' = 1;   j = 1, …, 20'])
    h('1.1 SVD伪逆更新',2)
    p('在迭代点线性化C(q+Δq)≈C(q)+JΔq，取最小范数校正Δq=−J†C。SVD写为J=UΣVᵀ，低于max(10⁻¹²,10⁻¹⁰σ_max)的奇异值不取倒数。每一步保留残差、秩、奇异值和更新量；位置与角度分别以10⁻⁸ m和10⁻⁸ rad验收。')
    eq([sub('q','k+1'),' = wrap(',sub('q','k'),' − V ',sup('Σ','†'),' ',sup('U','T'),' C(',sub('q','k'),'))'],6)
    p('本次使用无阻尼、无线搜索、无随机重启的全约束NR，最多100次更新。包裹角度是周期坐标的表示选择，不能等同于对任意工业机器人关节限位截断。原论文KDL实现、空间毫米单位和数值门限与本次20R实现不同，因此不比较绝对软件速度。')

    page('2 解析被动链与分划切换')
    p('APC将部分关节视为主动变量，使用非冗余被动链的解析IK闭合其余变量。被动只是求解角色，机构的每个关节仍是全驱动。对连续3R被动链，将整链变换写成T_prefix T_p T_suffix=T_target，即可保留全部主动关节并求被动子链。')
    eq([sub('T','p'),' = ',sup(sub('T','prefix'),'−1'),' ',sub('T','target'),' ',sup(sub('T','suffix'),'−1')],7)
    p('若被动末端的所需平面位姿为(x_p,y_p,φ)，先减去第三根杆的位移，得到腕点(u,v)。前两关节满足余弦定理，正负反余弦对应两条肘部分支。')
    eq(['u = ',sub('x','p'),' − ',sub('ℓ','3'),' cos φ,   v = ',sub('y','p'),' − ',sub('ℓ','3'),' sin φ'])
    eq([sub('c','2'),' = ',frac([sup('u','2'),' + ',sup('v','2'),' − ',sup(sub('ℓ','1'),'2'),' − ',sup(sub('ℓ','2'),'2')],['2',sub('ℓ','1'),sub('ℓ','2')]),',   ',sub('q','2'),' = ±acos(',sub('c','2'),')'],8)
    eq([sub('q','1'),' = atan2(v,u) − atan2(',sub('ℓ','2'),' sin ',sub('q','2'),', ',sub('ℓ','1'),' + ',sub('ℓ','2'),' cos ',sub('q','2'),')'],9)
    eq([sub('q','3'),' = φ − ',sub('q','1'),' − ',sub('q','2')])
    p('当|c₂|大于1时，所给主动构型无法由该被动链闭合；这是一种真实采样失败。本实现只容许10⁻¹²浮点边界舍入，随后重新检查闭链残差与关节界，没有通过扩大工作空间制造可达性。')
    h('2.1 PCS如何改变可连接集合',2)
    p('本模型主规划使用零起始索引分划p=0与p=16，分别对应关节0–2和16–18。对候选局部边，程序查找两端共同的非奇异IK分支，然后在该分划内递归投影。两种分划可在不同树边使用；若一条最终解依次使用不同分划，才构成最终解上的换链证据。')
    p('原论文式5–9使用不同被动链的分支覆盖解释可见性。在理想独立且等测度的假设下，m个分划、各ξ_k个IK解给出的可见性上界含有下式；障碍、限位及APC不可达性均会使实际成功率降低。')
    eq([sub('P','v'),' = 1 − ',large('∏','k=1','m',['(1 − ',frac('1',sub('ξ','k')),')'])],10)
    eq(['μ(',sub('V','m'),'(q)) ≤ ',sub('P','v'),' μ(',sub('C','cl'),')'])
    p('该式是带假设的上界，不是本次数值成功率的拟合公式。主规划仅使用两种明确实现的分划；Fig3机制试验另暴露17个连续非零3R分划。由于末端法兰长度为零，未补造原图m=20的数据。')

    page('3 RSS与递归局部连接')
    h('3.1 随机奇异采样的机制',2)
    p('固定被动链的人工奇异性由被动约束雅可比J_p秩下降表征。它不必是整个闭链约束雅可比的秩下降。对平面3R，det J_p与ℓ₁ℓ₂ sin q₂成比例；肘角为0或π时，正负肘支可相接。RSS以概率γ=0.2请求奇异目标，其他目标仍通过普通主动变量采样及固定被动链IK构造，不能用NR代替RSS。')
    eq(['rank(',sub('J','p'),'(',sub('q','sing'),')) < 3,   γ = 0.2'],11)
    p('保存的确定性奇异夹具确有rank 2，最小奇异值1.198×10⁻¹⁷，行列式4.065×10⁻¹⁹。旧单元夹具的两侧RBS连接用于检验分支相接机制，预设夹具不进入随机规划成功统计。随机采样需要同时实现几何闭合、限位、碰撞可行和树可达，构造一个代数奇异点本身不够。')
    h('3.2 RBS的递归条件',2)
    p('RBS先验证端点，再把两端的关节中点投影到约束流形。PCS使用选定的共同分划与分支，NR使用全约束投影，RSS使用其固定被动链。每个中点都执行闭链、限位和真实几何碰撞检查，任一失败则该试探边失败。分段距离降到0.05 rad后形成离散叶；深度最多12、每次连接最多256次投影。')
    eq([sub('q','m'),' = Π(',frac([sub('q','a'),' + ',sub('q','b')],'2'),'),   ',sub('q','m'),' ∈ ',sub('C','o')],12)
    eq(['recurse(',sub('q','a'),',',sub('q','m'),'),   recurse(',sub('q','m'),',',sub('q','b'),')'])
    p('这是重建的完整检查版本。作者旧RSS代码中靠近奇异端点的提前退出、近似π和若干原生接口问题在历史失败中单独保留；本次未把这些版本混成同一个实现。叶长受控仍不是严格连续碰撞检测证书。')
    figure('07_algorithm_flow.png','图1 真实复现程序的计算流程示意',15.8)

    page('4 虚拟机构与实验协议')
    p('20R模型由19根长度ℓ=sin(π/38) m的移动杆和长度1 m的固定地杆构成。移动杆与地杆均以半径0.002 m的胶囊体建模；只排除共享实际转动关节的相邻杆对，保留其余自碰撞检查。圆障碍按线段到圆心距离减去两者半径检查，接触或穿透均判为碰撞。')
    table(['项目','实际设置'],[
        ['构型与约束','20个转动关节；3维SE(2)闭链残差'],['关节限位','每关节[−π,π] rad'],['E0障碍','圆心(0.45,0.72) m，半径0.045 m'],['E1新增障碍','圆心(0.50,0.18)、(0.72,0.60) m；半径0.065、0.050 m'],['共同预算','每次1 s墙钟；120外迭代；最多300树节点'],['局部连接','RBS叶长0.05 rad，深度12，投影次数256'],['配对种子','7、23、41、59、73；NumPy PCG64'],['五种方法','RSS、PCS full、PCS fixed-0、PCS fixed-16、NR']], [3.2,13.2])
    p('设a=π/19。起点关节0与19为π/2−a/2，其余关节为−a。终点只把关节16–18改为(−2a,a,−2a)。两端在同一完整闭链内合法，但p16的肘部分支符号相反，而p0分支兼容。固定0因而是公平可连接的消融基线；固定16是分支受限的诊断负对照，不能只通过它失败来证明PCS优越。')
    p('双向树交替生长，采样后按0.6 rad的共同步长扩展，投影并经RBS后插入树。两棵树只在同一精确构型汇合后回溯；没有隐藏的直接连接捷径、路径平滑或有利种子筛选。墙钟包含搜索、投影、模型检查和终端复核，序列化与文件输出不计入规划时间。')
    p('先前20次探索试验采用种子7和23，后续50次稳定性试验采用五个种子，前两种子重复出现。本报告以50次作为主统计，不把70次调用说成每组7个独立样本。原始配置、源码哈希与完整失败记录一同保存。')

    page('5 完整轨迹及独立验证')
    figure('01_complete_trajectory.png','图2 E1中PCS seed 7的完整保存路径  背景叠绘每个保存构型',14.2)
    p('图2展示从实际搜索得到的路径。蓝、橙、绿分别为起点、中间构型和终点，浅色叠绘全部保存状态。经验证的杆端位置满足闭合，终点改变的是局部肘形；路径为几何关节轨迹，尚未进行速度、加速度或动力学时间参数化。')
    p('50次主试验中共有26条SUCCESS轨迹记录，包含2142个保存状态；其中13对E0/E1记录的完整关节路径逐值相同，即13条不同的关节路径分别在两环境中取得资格；其余24次均为TIME_LIMIT。独立程序从原始gzip读取关节值，重新计算累积刚体变换，并用另一套解析线段距离算法检查胶囊自碰撞和圆障碍，不依赖原程序的valid布尔值或GEOS碰撞函数。它还核对端点身份、树边回溯、每条局部边的分划、分支与叶长，并在局部流形上增密复验。')
    p('独立验证共检查2116条原RBS叶边及6348个新增中点，合计8490状态；最大位置残差9.040×10⁻⁹ m、角度残差2.220×10⁻¹⁵ rad，最小净距0.053965 m、最小关节余量0.346012 rad。11项正负oracle夹具均通过，未发现闭链、限位、碰撞或分支违规。GNUtime记录3.62 s墙钟、4.93 s CPU。增密时，PCS按原边共同分划重新求解析中点，NR按独立约束和雅可比重建投影中点。不能直接在环境关节空间线性插值后要求整段严格闭合，因为该直线通常离开流形。独立中点属于额外的离散复验，既不改写原始路径，也不声称有限采样证明连续安全。详细计数、残差、净距、oracle负例和耗时见独立验证原始summary。')

    page('6 约束误差与空间路径复用')
    figure('02_path_constraints.png','图3 保存路径逐点的闭链误差  关节余量与碰撞净距',15.8)
    p('图3取自原始逐点数值，路径点编号不能解释为秒。位置误差与角度误差分别按其物理单位判断，不能用混合单位范数掩盖某一分量的越界。报告同时保留关节界余量及最小几何净距，使闭合、限位与碰撞三项有效性可以分别复核。')
    table(['空间路径','保存点数','证据边界'],[
        ['II RRT PCS / NR','177 / 393','已保存独立离散证书'],['II CBiRRT PCS / NR','413 / 689','已保存独立离散证书'],['II SBL PCS','387','前缀262与尾部125组合资格'],['III RRT PCS','649','648叶；最终4次主动臂角色切换']], [4.4,2.6,9.4])
    p('III的15个里程点形成14条最终局部边，最大相邻L2距离0.04995649 rad；最大位置闭链分量4.547×10⁻¹³ mm，最大旋转矩阵差4.469×10⁻¹⁵。3253项保留哈希引用与节点、数学、碰撞记录逐项对应。原生TCP差异649/649仍存在，aligned策略合格不等于原生TCP等价。环境V失败前缀未到目标，V5只作监督器工程候选，不列成功路径。')

    page('7 RSS PCS NR的配对规划结果')
    figure('03_paired_comparison.png','图4 主试验成功率及Wilson区间  右轴为含超时的平均规划墙钟',15.8)
    summary=json.loads((ROOT/'results/20261019_D081_stability_pilot/planner_summary.json').read_text())['rows']
    rr=[]
    for method,label in [('RSS','RSS'),('PCS','PCS full'),('PCS_FIXED_0','fixed-0'),('PCS_FIXED_16','fixed-16'),('NR','NR')]:
        a=next(r for r in summary if r['environment']=='E0_sparse' and r['method']==method);b=next(r for r in summary if r['environment']=='E1_clutter' and r['method']==method)
        rr.append([label,f"{a['success']}/5",f"{a['mean_wall_sec']:.4f}",f"{b['success']}/5",f"{b['mean_wall_sec']:.4f}",'—' if b['mean_path_length_rad'] is None else f"{b['mean_path_length_rad']:.4f}"])
    table(['方法','E0成功','E0时间 s','E1成功','E1时间 s','E1成功路径 rad'],rr,[2.8,1.7,2.5,1.7,2.5,5.2])
    p('NR在本任务中最快且成功路径最短；全PCS没有超过兼容的固定0。这些结论限于本机、当前重建场景和1 s预算。成功5/5的95% Wilson区间约为[0.566,1]，0/5仍约为[0,0.434]；样本不足以估计罕见失败或给出广泛方法排序。超时并不证明问题无解。')
    p('E1的平均树节点数分别为RSS 39.6、PCS全分划24.6、固定0为6.2、固定16为46.6、NR为5.0；它们包含超时树，不能当成成功路径节点数。每次路径节点、树节点、投影和局部检查计数均见原始CSV。')

    page('8 分划消融与RSS失败解释')
    h('8.1 搜索树切换与最终路径切换',2)
    p('按树节点的incoming_attempt追踪真实父子关系，而不是把全体尝试按时间相邻排列：10次全PCS调用中只发现1次被接受树边的父子分划变化，出现在E0 seed 73超时搜索树中。6条成功全PCS路径的全部最终边均使用p0，真实解路径切换为0。空间III的4次角色切换是另一种已保存的正例，不与20R结果混为一组。')
    p('固定16的起终点肘支相反，因此它被设计为单图册连通性受限的负对照。固定0兼容两端且10次全部成功，全PCS只成功6/10，说明在短预算下探索其他图册可能增加成本。E1增加障碍后没有改变成功的种子集合，13对两环境均成功的密集路径逐值相同，障碍压力不足，不能从当前试验推出PCS在狭窄空间中的普遍优势。')
    h('8.2 RSS为什么没有完成路径',2)
    p('主试验10次RSS调用均在1 s预算终止。保存记录包含72次奇异目标请求，每次最多32次构造，合计2304次构造均为IK_UNREACHABLE，尚无闭合的随机奇异目标进入rank检查。因此当前首要失败发生在奇异采样的闭合可达性阶段，不能归因为树已经到达奇异点却无法跨越。')
    p('普通APC失败、固定被动分支限制及短时间预算还会进一步影响树扩展。可观测失败来自奇异采样与短闭合链可达性耦合，但尚不能区分具体实现缺陷与采样支持不足，根因归类为UNKNOWN，原论文相同输入不可得的部分属于PAPER RESOURCE；现有证据不支持认定RSS理论或原论文错误。固定夹具rank 2和两侧连接只证明局部数学机制存在。')
    p('历史三次修正及失败已归档，本轮没有重置额度，也没有再调γ、奇异目标分布、障碍或种子。若继续研究，需另行冻结能够提高奇异闭合可达率的假设、原作者对照和预算，保持原实现与修订实现分栏；当前报告将RSS规划结果保留为失败。')
    h('8.3 公平性与可重复性边界',2)
    p('五种方法使用相同的种子列表、起终点、几何规则和墙钟上限，但各算法消耗随机数的次序不同，因此paired seed不意味着每次具有相同的采样向量。计时也会随机器负载改变。重跑应保留新目录和实际终止状态，不要求逐字复制墙钟截断的树节点数量。')

    page('9 投影成本与收敛行为')
    figure('05_projection_and_distance.png','图5 APC与NR的投影成功  成本及邻点距离变化',15.5)
    table(['方法','成功输入','成功率','平均单次墙钟','平均NR更新数'],[
        ['APC','51 / 768','6.6406%','0.2542 ms','—'],['NR','768 / 768','100%','1.8137 ms','7.447']], [2.2,3.2,2.7,4.2,4.1])
    p('768个配对输入来自6档局部扰动尺度，每档128个。APC固定p16并保留两条解析分支请求，当前每输入投影过程的平均墙钟比约为7.14，两实现的工作量及成功集合不同，但大多数主动构型不能闭合。NR实际Newton更新均值为7.447次。旧字段iterations把初始及终止残差记录在内，其8.447是平均trace长度。失败APC也计入单次平均时间，成功率与成功条件下的代价不能混读。')
    p('原论文在两个空间6R机构、无关节限位和无碰撞条件下投影500000个ambient样本，报告APC成功7.92%、NR成功100%、APC快3–6倍。本次机构、采样分布、约束和软件不同，数值接近不能称同条件复现。图5右图另有每方法384次距离变化记录，仅统计成功投影的实际前后邻点距离，保留失败分母；它不是投影误差的理论曲线。')
    figure('06_nr_convergence.png','图6 真实NR记录  分别以10⁻⁸验收  零值以10⁻¹⁸作为对数显示下限',11.0)

    page('10 原论文图表的实际覆盖')
    figure('04_local_visibility.png','图7 Fig3与Fig4对应的真实局部连接数据  分箱宽度0.3 rad',16.3)
    p('Fig4对应试验在同一组240点对上比较RSS、PCS和NR，其中4对为真实抽样得到的零距离平凡点对，连接成功分别为73、105和177次，即30.42%、43.75%和73.75%。距离分箱可能样本很少，且点对包含同一点或由局部生成器得到的邻点，不能将总体解释为原空间环境的全局可见性。')
    p('Fig3对应试验每个m有180点对，m=1、2、3、7、16成功率为68.89%、51.11%、60.00%、94.44%、99.44%。各m重新生成构型总体，不是相同点对上的严格配对消融；非单调性及高m上升都不能单独归因于增加分划。原图m=20、各m规划耗时曲线尚未复现。')
    p('原文Table I环境II的Bi RRT平均规划时间为RSS 7.27 s、PCS 0.22 s、NR 0.13 s；Table II局部可见性为0.85%、1.31%、23.74%，平均局部连接耗时4.97、0.32、1.40 ms。这些是原文报告值，不是本次实验。原模型、原样本与C++依赖不同，不能与当前20R绝对时间拼表计算加速比。')
    table(['原文项目','本次对应证据','可比性限制'],[
        ['Fig3','900次局部连接与真实分箱曲线','总体随m改变；无原m20和耗时曲线'],['Fig4','240配对点对 × 三方法','20R模型；无约50000原样点对'],['Fig5','每方法384条距离变化记录','局部扰动；非原空间采样'],['Table I','50次20R配对规划；空间单路径证书','无500次每组或原22500统计'],['Table II','节点、局部检查、路径和可见性实测','非原环境II Bi RRT同机指标']], [2.4,6.4,7.6])

    page('11 复现资产与研究局限')
    p('完整覆盖CSV继承原211个义务行及所有248个历史列，再追加当前范围、状态、证据和局限。实际主义务编号为P4 T01至T30，旧任务标题所称29项不用于删行。环境VI只读归档；原22500次是论文事实，18000次是已取消历史计划，50次是当前主配对规划调用，三者不得互换。')
    h('11.1 重跑入口',2)
    p('安装requirements-p4-three-methods.txt中的研究依赖后，在仓库根目录执行以下入口。规划脚本使用固定配置且拒绝覆盖已有输出，须指定新的结果目录；绘图与报告脚本拒绝覆盖已有交付，以下命令均另选新输出路径；墙钟预算会使不同机器上的截断轨迹有所差异。图表脚本只读取保存数据，报告构建需要独立验证结果。')
    for text in [
        'python3 -m unittest tests.test_p4_comparison_model tests.test_p4_comparison_strategies',
        'python3 scripts/p4_d081_stability_pilot.py --output results/20261022_D082_replay_example01',
        'python3 scripts/p4_d082/independent_planar_validation.py --output results/20261022_D082_validation_replay01',
        'python3 scripts/p4_d082/audit_spatial_saved_20261022.py --output reports/P4_spatial_replay01.json',
        'python3 scripts/p4_d082/build_research_figures.py --output work/d082_figures_replay01',
        '$ARTIFACT_PY scripts/p4_d082/build_chinese_research_report.py --output reports/P4_report_replay01.docx']:
        paragraph=p(text)
        for run in paragraph.runs:run.font.name='DejaVu Sans Mono';run.font.size=Pt(8)
    p('ARTIFACT_PY须指向Codex内置Python，精确路径见交付README。独立路径验证器始终复核冻结的D081五十调用，输出参数不会将其输入改成上一行的新规划目录。主数据目录为results/20261019_D081_stability_pilot，投影与局部连接数据为results/20261019_D081_numeric_pilot；本轮独立几何验证为results/20261022_D082_independent_planar01，算法保存分析为results/20261022_D082_saved_analysis01。scripts/p4_d082内含独立核验、真实绘图和报告构建入口。精确命令、原始哈希和新增数值计数另见交付README。')
    h('11.2 仍未建立的结论',2)
    p('未认证连续CCD、原作者原生TCP几何等价、原六环境全部组合、原500次统计、m20规划耗时或PCS在窄通道中的普遍优势。RSS随机完整路径仍失败。20R轨迹未建动力学、力矩、速度、抓持滑移或执行误差模型，几何路径不能直接作为机器人可执行指令。')
    p('本次50调用先前实际墙钟40.917 s，探索数值试验90.386 s；它们是历史运行耗时，不算作本轮重跑。独立验证和保存分析的本轮墙钟、CPU由各自summary记录。未启动新的作者求解、VI、Isaac或硬件进程，未扩大正式统计。')

    page('12 面向双FR3协同搬运的潜在应用')
    h('12.1 刚性共同抓持下的闭链规划接口',2)
    p('原论文假设物体与抓手组成刚体，适合共同搬运阶段的几何路径约束。对双FR3关节q_L、q_R，设固定抓持变换G_L、G_R为TCP到物体的标定关系，可用两侧推算的物体位姿相等定义6维残差。')
    eq([sup(sub('T','WO'),'L'),'(',sub('q','L'),') = ',sub('T','WL'),'(',sub('q','L'),') ',sub('G','L')])
    eq([sup(sub('T','WO'),'R'),'(',sub('q','R'),') = ',sub('T','WR'),'(',sub('q','R'),') ',sub('G','R')])
    eq([sub('C','FR3'),'(q) = ',sup(['Log(',sup(['(',sup(sub('T','WO'),'L'),'(',sub('q','L'),'))'],'−1'),' ',sup(sub('T','WO'),'R'),'(',sub('q','R'),'))'],'∨'),' = 0'],13)
    p('将两臂14维关节状态输入运动学模块，输出残差与6×14雅可比，再由NR、局部连接和采样规划器生成关节几何路径。坐标扰动方向、Log的平移旋转定义及单位必须一致；本次平面Jacobian数值资格不能替代FR3的独立差分资格。MoveIt模型可提供FK和实际关节界，FCL接口提供机器人、吸盘、物体及场景障碍的统一碰撞判定。')
    h('12.2 NR与RBS的适用前提',2)
    p('NR对冗余7R臂不要求闭式IK，是优先考虑的可扩展机制，但关节限位和奇异邻域可能使投影漂移或失败。未来需要独立记录投影输入、位移、秩、闭链误差和限位，并在相同FR3任务上比较限步或阻尼版本；这些是PROPOSED方案，当前未实现。')
    p('RBS可检查共享物体运输途中每段局部运动，输入两个闭链构型、投影器和碰撞接口，输出带来源记录的离散连接或具体拒绝原因。执行前还需时间参数化、速度加速度限制、连续碰撞或有保守界的扫掠验证，以及执行中实时闭链误差监测。')
    h('12.3 PCS与RSS的额外困难',2)
    p('FR3单臂7R冗余，原论文6R非冗余被动臂的有限闭式分支不能直接搬用。若固定冗余参数形成6关节子问题，必须证明该子链的IK可得、分支连续且覆盖足够工作空间；只能增加优化初值不等于真正PCS。当前20R最终路径无换链优势，空间III的4次角色切换只能作为实现机制的证据，不能预测双FR3性能提升。')
    p('RSS还需推导所选FR3被动子链的真实秩亏集合，给出可闭合、限位合法、碰撞可行的随机采样器。当前2304次构造失败提醒我们，代数奇异条件远不足以形成可用目标。未完成这些前置条件时，不建议把RSS作为搬运主规划器。')

    page('13 接触阶段边界与后续研究问题')
    h('13.1 吸盘柔性与闭链假设',2)
    p('双侧吸盘与物体之间可能有柔性变形、滑移、吸附失效和内部载荷。式13只表达刚性几何一致性，不能保证抓持力闭合、吸力容量或内部应力受限。未来接口需要同时提供关节状态、标定抓持变换、物体位姿估计、吸附状态及可信的力或接触信息；传感数据来源与控制器读数必须分离于仿真真值。')
    p('若抓持变换随力而变化，可研究显式柔性接触状态或允许受物理约束的小残差集合，但这将改变原论文C(q)=0的模型。应单独推导弹性与内力关系并设版本，不能仅扩大闭链容差来消除执行误差。该问题属于后续研究，当前未给出控制稳定性或载荷安全结论。')
    h('13.2 单臂推进阶段的不同约束',2)
    p('在右臂后侧吸持并向车厢推进、左臂停泊的阶段，双臂共享刚性物体的闭链已不存在。可借鉴NR与RBS处理末端姿态、接触面或导向约束，但约束必须由该阶段实际接触模式重建。接触的单边性、摩擦、卡滞与力位控制不是原P4几何算法已经解决的问题。')
    h('13.3 候选验证顺序',2)
    p('未来可先固定FR3运动学、抓持变换和统一碰撞几何，建立具有独立可行路径的共同搬运任务；然后用相同起终点、环境和预算比较NR、兼容固定被动求解及真实PCS。若没有有限闭式子链，就将PCS列为不适用，不以另一种数值投影冒称PCS。困难障碍需保留可行路径证据，并分开报告树换链、最终解换链和实际执行误差。')
    p('几何规划合格后，再由后续控制研究检验时间参数化、协同跟踪、吸盘载荷、内部力和失效恢复。候选对照应记录成功与失败、CPU和墙钟、投影成本、路径长度、最小净距及闭链误差。只有这些独立证据取得后，才能判断P4机制是否对双FR3系统带来实际收益。本次只完成方法分析，没有实现FR3适配。')
    h('参考文献与数据入口',2)
    p('[1] Sintov A, Borum A, Bretl T. IEEE RA L, 2018, 3(4): 2886–2893. DOI 10.1109/LRA.2018.2846806. 作者稿 https://avishaisintov.wordpress.com/wp-content/uploads/2018/06/sintov-pcs.pdf。')
    p('[2] 作者CKCplanning参考源码快照909bd72fff778aca8682c96f9eb86ae4eb75557b，来源https://github.com/avishais/CKCplanning。许可证边界和本地补丁见既有author_source_manifest；本报告不发布私有论文副本。')
    p('[3] 本研究源码与原始数据 https://github.com/dreamerToTOE/dual-arm-coordinated-manipulation-reproduction ，分支p4-full-reproduction-pcs-study。完整证据路径、版本和哈希见P4_D082_DELIVERY_README_20261022.md。')
    for paragraph in DOC.paragraphs[-3:]:
        paragraph.paragraph_format.line_spacing=1.0
        paragraph.paragraph_format.space_after=Pt(3)
        for run in paragraph.runs:run.font.size=Pt(9)
    for border in list(DOC.element.iter(qn('w:pBdr'))):
        border.getparent().remove(border)
    DOC.save(OUT)
    print(OUT)

if __name__=='__main__':main()
